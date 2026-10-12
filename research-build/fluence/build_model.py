"""Fluence Energy (Nasdaq: FLNC) initiation model. Live formulas throughout.

House style follows the Vertiv model: navy header rows, navy bold titles, blue font for hardcoded inputs,
yellow fill on my own assumptions, green font for cross-sheet links, black for formulas.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from data import *

NAVY = "1F3864"; MUTED = "595959"
F_TTL = Font(bold=True, color=NAVY, size=12)
F_SUB = Font(color=MUTED, size=9)
F_HDR = Font(bold=True, color="FFFFFF", size=10)
F_IN = Font(color="0000FF")
F_INB = Font(color="0000FF", bold=True)
F_LINK = Font(color="008000")
F_B = Font(bold=True)
F_NOTE = Font(color=MUTED, size=9, italic=True)
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_ASSUME = PatternFill("solid", fgColor="FFFF00")
FILL_KEY = PatternFill("solid", fgColor="D9E2F3")
WRAP = Alignment(wrap_text=True, vertical="top")
RIGHT = Alignment(horizontal="right")

USD2 = '"US$"#,##0.00'
M1 = '#,##0.0'
M0 = '#,##0'
M3 = '#,##0.000'
PCT1 = '0.0%'
PCT0 = '0%'
MULT = '0.0"x"'
MULT2 = '0.00"x"'
DLONG = "12 October 2026"
SUB = f"Tommy Lau | Fluence Energy (Nasdaq: FLNC) | {DLONG} | " + ("DRAFT | " if CALL["draft"] else "") + "Personal research, not investment advice."

wb = openpyxl.Workbook()


def sheet(name, title, widths):
    ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet(name)
    ws.title = name
    ws["A1"] = title; ws["A1"].font = F_TTL
    ws["A2"] = SUB; ws["A2"].font = F_SUB
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def head(ws, row, cells, start=1):
    for i, c in enumerate(cells, start):
        x = ws.cell(row=row, column=i, value=c)
        x.font = F_HDR; x.fill = FILL_HDR
        x.alignment = Alignment(horizontal="left" if i == start else "right", vertical="center", wrap_text=True)


def note(ws, row, texts):
    for i, t in enumerate(texts):
        ws.cell(row=row + i, column=1, value=t).font = F_NOTE


# =====================================================================  COVER (filled last)
ws = sheet("Cover", "FLUENCE ENERGY, INC. (NASDAQ: FLNC)  INITIATION MODEL" + (" (DRAFT)" if CALL["draft"] else ""), [36, 24, 90])

# =====================================================================  ASSUMPTIONS
ws = sheet("Assumptions", "ASSUMPTIONS AND INPUTS", [52, 14, 10, 100])
ws["A3"] = ("Blue = hardcoded input. Yellow fill = my own assumption (change these to move the value). "
            "Green = link to another tab. Black = formula. US$ million unless stated. Fiscal years end 30 September.")
ws["A3"].font = F_NOTE
head(ws, 5, ["Input", "Value", "Unit", "Source / note"])
A = {}
MINE = {"g27", "g28", "agm27", "agm28", "opx27", "opx28", "mult", "q4burn"}
rows = [
    ("MARKET", None, None, None, None),
    ("price", "Reference share price", PRICE, "US$", "Nasdaq close, 9 October 2026. Nasdaq.com historical data, retrieved 12 October 2026."),
    ("hiclose", "Highest close, 3 February 2026", HI_CLOSE, "US$", "Nasdaq.com historical data. Intraday high US$33.51 on 30 January 2026."),
    ("sha", "Class A shares outstanding, 31 July 2026", SH_A, "m", "Form 10-Q for the quarter to 30 June 2026, cover page: 143,163,588."),
    ("shb", "Class B-1 shares outstanding, 31 July 2026", SH_B1, "m", "Form 10-Q cover: 41,432,781, held by AES Grid Stability; redeemable for Class A one for one."),
    ("cash", "Cash, cash equivalents and restricted cash, 30 June 2026", CASH_JUN, "US$m", "Form 10-Q balance sheet: 339.328 + 25.630."),
    ("conv", "2.25% convertible senior notes due 2030, principal", CONVERT, "US$m", "Form 10-Q, Note 12. Conversion price about US$21.35; treated as debt."),
    ("q4burn", "Cash used in the September 2026 quarter", Q4_BURN, "US$m", "MY ASSUMPTION. Not yet reported. Close to the minus US$133m implied by the S&P Global full-year free cash flow forecast less the nine-month actual."),
    ("fcf9m", "Free cash flow, nine months to 30 June 2026", FCF_9M, "US$m", "3Q FY26 results release (company definition)."),
    ("fcfcons", "FY26 free cash flow, consensus", FCF_CONS_FY26, "US$m", "stockanalysis.com (S&P Global), retrieved 12 October 2026. Cross-check only."),
    ("cons27", "FY27 revenue, consensus", CONS_REV_27, "US$m", "stockanalysis.com (S&P Global), retrieved 12 October 2026. Cross-check only."),
    ("constp", "Average analyst target", CONS_TP, "US$", "stockanalysis.com, 23 analysts; median US$8.50, low US$3, high US$23. Cross-check only."),
    ("REPORTED (fiscal years to 30 September)", None, None, None, None),
    ("r24", "FY24 revenue", REV_H[0], "US$m", "4Q and FY25 results release (Form 8-K, Exhibit 99.1), 24 November 2025."),
    ("r25", "FY25 revenue", REV_H[1], "US$m", "Same release."),
    ("a24", "FY24 adjusted gross profit", AGP_H[0], "US$m", "Same release, non-GAAP reconciliation."),
    ("a25", "FY25 adjusted gross profit", AGP_H[1], "US$m", "Same release."),
    ("e24", "FY24 adjusted EBITDA", EBITDA_H[0], "US$m", "Same release."),
    ("e25", "FY25 adjusted EBITDA", EBITDA_H[1], "US$m", "Same release."),
    ("f24", "FY24 free cash flow", FCF_H[0], "US$m", "Same release."),
    ("f25", "FY25 free cash flow", FCF_H[1], "US$m", "Same release."),
    ("GUIDANCE FOR FY26 (Form 8-K, 16 September 2026)", None, None, None, None),
    ("g26r", "FY26 revenue guide", G26["rev"], "US$m", "About US$2.4bn. Was US$2.9 to 3.1bn (5 August 2026) and US$3.2 to 3.6bn (24 November 2025)."),
    ("g26e", "FY26 adjusted EBITDA guide", G26["ebitda"], "US$m", "A loss of about US$200m. Was minus US$30m to plus US$10m (5 August) and US$40 to 60m (24 November 2025)."),
    ("g26nov", "FY26 revenue guide midpoint, November 2025", (G26_NOV["rev_lo"] + G26_NOV["rev_hi"]) / 2, "US$m", "4Q and FY25 results release."),
    ("MY ASSUMPTIONS", None, None, None, None),
    ("g27", "FY27 revenue growth", G27, "%", "MY ASSUMPTION. Below the S&P Global consensus of US$3.25bn: US$2.2bn of backlog was scheduled for FY27 at the August call, and Houston must work for the rest."),
    ("g28", "FY28 revenue growth", G28, "%", "MY ASSUMPTION. Backlog of US$6.4bn and record orders support it once delivery runs on time."),
    ("agm27", "FY27 adjusted gross margin", AGM27, "%", "MY ASSUMPTION. Penalties and launch costs fade through the year."),
    ("agm28", "FY28 adjusted gross margin", AGM28, "%", "MY ASSUMPTION. Back to about the 12.9% to 13.7% of FY24 and FY25."),
    ("opx27", "FY27 net operating costs (adjusted gross profit less adjusted EBITDA)", OPX27, "US$m", "MY ASSUMPTION. US$270m in FY24, US$290m in FY25, US$208m in the nine months to June 2026 (Quarters tab)."),
    ("opx28", "FY28 net operating costs", OPX28, "US$m", "MY ASSUMPTION. Grows with volume."),
    ("mult", "EV / FY28E revenue, base", BASE_MULT, "x", "MY ASSUMPTION. About 11x my FY28 EBITDA, below the peer median of 13x trailing EV/EBITDA; see Peers."),
]
r = 6
for key, label, val, unit, src in rows:
    if label is None:
        ws.cell(row=r, column=1, value=key).font = Font(bold=True, color=NAVY)
        r += 1; continue
    ws.cell(row=r, column=1, value=label)
    c = ws.cell(row=r, column=2, value=val)
    c.font = F_INB if key in MINE else F_IN
    if key in MINE: c.fill = FILL_ASSUME
    c.number_format = PCT1 if unit == "%" else (MULT2 if unit == "x" else (M3 if unit == "m" else (USD2 if unit == "US$" else M1)))
    ws.cell(row=r, column=3, value=unit)
    ws.cell(row=r, column=4, value=src).alignment = WRAP
    A[key] = f"Assumptions!$B${r}"
    r += 1

# =====================================================================  FINANCIALS
ws = sheet("Financials", "FINANCIAL SUMMARY AND FORECASTS (US$ million unless stated; fiscal years to 30 September)", [48, 13, 13, 13, 13, 13])
head(ws, 4, ["", "FY24A", "FY25A", "FY26 guide", "FY27E*", "FY28E*"])
L = lambda k: "=" + A[k]
EVR = "Valuation!$B$16"
lines = [
    ("Revenue", [L("r24"), L("r25"), L("g26r"), "=D5*(1+" + A["g27"] + ")", "=E5*(1+" + A["g28"] + ")"], M1),       # 5
    ("Revenue growth", [None, "=C5/B5-1", "=D5/C5-1", "=E5/D5-1", "=F5/E5-1"], PCT1),                                   # 6
    ("Adjusted gross margin", ["=B8/B5", "=C8/C5", None, L("agm27"), L("agm28")], PCT1),                                # 7
    ("Adjusted gross profit", [L("a24"), L("a25"), None, "=E5*E7", "=F5*F7"], M1),                                      # 8
    ("Net operating costs (adjusted gross profit less adjusted EBITDA)", ["=B8-B10", "=C8-C10", None, L("opx27"), L("opx28")], M1),  # 9
    ("Adjusted EBITDA", [L("e24"), L("e25"), L("g26e"), "=E8-E9", "=F8-F9"], M1),                                       # 10
    ("Adjusted EBITDA margin", ["=B10/B5", "=C10/C5", "=D10/D5", "=E10/E5", "=F10/F5"], PCT1),                          # 11
    ("Free cash flow", [L("f24"), L("f25"), None, None, None], M1),                                                     # 12
    ("EV / revenue at reference price", [None, None, f"={EVR}/D5", f"={EVR}/E5", f"={EVR}/F5"], MULT2),                 # 13
    ("EV / adjusted EBITDA at reference price", [None, None, None, f"={EVR}/E10", f"={EVR}/F10"], MULT),               # 14
]
r = 5
for label, vals, fmt in lines:
    ws.cell(row=r, column=1, value=label).font = F_B if label in ("Revenue", "Adjusted EBITDA") else Font()
    for j, v in enumerate(vals):
        c = ws.cell(row=r, column=2 + j, value=v if v is not None else "")
        c.number_format = fmt; c.alignment = RIGHT
        if isinstance(v, str) and v.startswith("=Assumptions"):
            c.font = F_LINK
    r += 1
note(ws, 17, [
    "*FY27E and FY28E are my own estimates, not Fluence's. FY26 is Fluence's own guide of 16 September 2026; it gives no gross margin guide.",
    "Adjusted gross profit and adjusted EBITDA are Fluence's non-GAAP measures. Net operating costs is my derived line: adjusted gross profit less adjusted EBITDA.",
    "EV uses my estimate of net debt at 30 September 2026 (Valuation tab). Fluence has given no guidance beyond FY26; its FY27 plan is due with the FY26 results.",
])

# =====================================================================  VALUATION
ws = sheet("Valuation", "VALUATION: EV / FY28E REVENUE, TWELVE MONTHS OUT", [60, 16, 76])
head(ws, 4, ["Line", "Value", "Working / note"])
val = [
    ("rev28", "FY28E revenue, US$m", "=Financials!F5", M1, "Financials tab. By October 2027 the market will be pricing the year to September 2028."),   # 5
    ("mult", "Multiple, x revenue", "=" + A["mult"], MULT2, "Assumptions tab."),                                                                    # 6
    ("ev28", "Enterprise value at the base multiple, US$m", "=B5*B6", M1, "Revenue x multiple."),                                                   # 7
    ("nd", "Less net debt, estimated at 30 September 2026, US$m", "=" + A["conv"] + "-" + A["cash"] + "+" + A["q4burn"], M1, "Convertible principal less June cash, plus my estimate of September quarter cash use. Held flat: Fluence aims for neutral to positive operating cash flow in FY27."),  # 8
    ("sh", "Shares, Class A plus Class B-1, m", "=" + A["sha"] + "+" + A["shb"], M3, "Economic shares. Convertible far out of the money; no dilution assumed."),  # 9
    ("base", "BASE VALUE PER SHARE, US$", "=(B7-B8)/B9", USD2, "Equity value / shares."),                                                          # 10
    ("px", "Reference price, US$", "=" + A["price"], USD2, "Nasdaq close, 9 October 2026."),                                                         # 11
    ("up", "Base value against price", "=B10/B11-1", PCT1, "Band: LONG at +15% or more, SHORT at -25% or worse, NO CALL between."),                  # 12
    ("rule", "VIEW ON THE BAND", '=IF(B12>=0.15,"LONG",IF(B12<=-0.25,"SHORT","NO CALL"))', None, ("DRAFT view. The final call is Tommy Lau's." if CALL["draft"] else "Band rule on the base value. Tommy Lau's call: " + CALL["direction"] + ".")),       # 13
    (None, "", None, None, ""),                                                                                                                      # 14
    ("mcap", "Market value, US$m", "=B9*B11", M1, "Price x shares."),                                                                               # 15
    ("ev", "Enterprise value, US$m", "=B15+B8", M1, "Market value plus estimated net debt."),                                                         # 16
    ("need", "FY28 revenue the price needs at the base multiple, US$m", "=B16/B6", M1, "Compare with my FY28E (B5) and FY27 consensus."),             # 17
    ("needm", "That figure against my FY28E", "=B17/B5", PCT1, ""),                                                                                 # 18
    ("evs26", "EV / FY26 revenue guide", "=B16/Financials!D5", MULT2, ""),                                                                           # 19
    ("evs27", "EV / FY27 consensus revenue", "=B16/" + A["cons27"], MULT2, "Cross-check only."),                                                     # 20
    ("evs28", "EV / FY28E revenue", "=B16/B5", MULT2, ""),                                                                                          # 21
    ("eve28", "EV / FY28E adjusted EBITDA", "=B16/Financials!F10", MULT, ""),                                                                       # 22
    ("basee", "Base case EV / FY28E adjusted EBITDA", "=B7/Financials!F10", MULT, "The base revenue multiple expressed on EBITDA; compare Peers."),    # 23
    ("revl", "Price at which the base value is 15% above it, US$", "=B10/1.15", USD2, "Revisit level for a LONG."),                                 # 24
    ("revs", "Price at which the base value is 25% below it, US$", "=B10/0.75", USD2, "Revisit level for a SHORT."),                               # 25
    ("offhi", "Price against highest close", "=B11/" + A["hiclose"] + "-1", PCT1, "Highest close US$32.23, 3 February 2026."),                       # 26
    ("ndj", "Net debt at 30 June 2026, US$m", "=" + A["conv"] + "-" + A["cash"], M1, "Reported basis."),                                             # 27
    ("q4c", "Consensus-implied free cash flow, September 2026 quarter, US$m", "=" + A["fcfcons"] + "-" + A["fcf9m"], M1, "Cross-check for my September quarter cash use."),  # 28
    ("cut", "FY26 revenue guide against the November 2025 midpoint", "=Financials!D5/" + A["g26nov"] + "-1", PCT1, ""),                               # 29
]
V = {}
r = 5
for key, label, f, fmt, nt in val:
    ws.cell(row=r, column=1, value=label)
    if f is not None:
        c = ws.cell(row=r, column=2, value=f)
        if fmt: c.number_format = fmt
        if key in ("base", "rule"):
            ws.cell(row=r, column=1).font = F_B; c.font = F_B; c.fill = FILL_KEY
        if isinstance(f, str) and f.startswith("=Assumptions") and "+" not in f and "-" not in f:
            c.font = F_LINK
    ws.cell(row=r, column=3, value=nt).alignment = WRAP
    if key: V[key] = f"Valuation!$B${r}"
    r += 1

# =====================================================================  SCENARIOS
ws = sheet("Scenarios", "THREE CASES, TWELVE MONTHS OUT", [10, 11, 11, 13, 11, 12, 12, 12, 12, 12, 10])
head(ws, 4, ["Case", "FY27 growth", "FY28 growth", "FY28 adj. gross margin", "EV/sales", "FY27 revenue", "FY28 revenue",
             "FY28 adj. EBITDA", "EV/EBITDA", "Value, US$", "vs price"])
r = 5
for name, g1, g2, gm, mult, p in SCEN:
    ws.cell(row=r, column=1, value=name).font = F_B
    if name == "Base":
        for j, k in enumerate(["g27", "g28", "agm28"]):
            c = ws.cell(row=r, column=2 + j, value="=" + A[k]); c.font = F_LINK; c.number_format = PCT1
        c = ws.cell(row=r, column=5, value="=" + A["mult"]); c.font = F_LINK
    else:
        for j, v in enumerate([g1, g2, gm]):
            c = ws.cell(row=r, column=2 + j, value=v); c.font = F_INB; c.fill = FILL_ASSUME; c.number_format = PCT1
        c = ws.cell(row=r, column=5, value=mult); c.font = F_INB; c.fill = FILL_ASSUME
    ws.cell(row=r, column=5).number_format = MULT2
    ws.cell(row=r, column=6, value=f"=Financials!$D$5*(1+B{r})").number_format = M1
    ws.cell(row=r, column=7, value=f"=F{r}*(1+C{r})").number_format = M1
    ws.cell(row=r, column=8, value=f"=G{r}*D{r}-{A['opx28']}").number_format = M1
    ws.cell(row=r, column=9, value=f"=G{r}*E{r}/H{r}").number_format = MULT
    ws.cell(row=r, column=10, value=f"=(G{r}*E{r}-Valuation!$B$8)/Valuation!$B$9").number_format = USD2
    ws.cell(row=r, column=11, value=f"=J{r}/{A['price']}-1").number_format = PCT0
    r += 1
ws["A9"] = "Probability"; ws["A9"].font = F_B
for j, s in enumerate(SCEN):
    ws.cell(row=8, column=2 + j, value=s[0]).font = F_B
    c = ws.cell(row=9, column=2 + j, value=s[5]); c.font = F_INB; c.fill = FILL_ASSUME; c.number_format = PCT0
ws["E9"] = "=SUM(B9:D9)"; ws["E9"].number_format = PCT0; ws["F9"] = "must be 100%"; ws["F9"].font = F_NOTE
ws["A11"] = "Probability-weighted value, US$"; ws["A11"].font = F_B
ws["E11"] = "=B9*J5+C9*J6+D9*J7"; ws["E11"].number_format = USD2; ws["E11"].font = F_B; ws["E11"].fill = FILL_KEY
ws["F11"] = f"=E11/{A['price']}-1"; ws["F11"].number_format = PCT1
note(ws, 13, [
    "All cases start from the FY26 revenue guide of about US$2.4bn, use my FY28 net operating costs (Assumptions) for the EBITDA cross-check,",
    "and take off the same estimated net debt over the same shares. Bear: Houston stays slow into 2027 and the FY27 plan is cut again; the bear multiple",
    "prices a company with almost no earnings and a cash question, so its EV/EBITDA is not meaningful. Bull: Houston fixed early in 2027 and data centre",
    "awards become orders; the bull multiple is about 10x FY28 EBITDA.",
])

# =====================================================================  SENSITIVITY
ws = sheet("Sensitivity", "VALUE PER SHARE: FY28E REVENUE AGAINST EV / REVENUE MULTIPLE (US$)", [28, 14, 14, 14])
head(ws, 4, ["FY28 revenue (US$m) \\ multiple", "", "", ""])
for j, (m, link) in enumerate(zip(SENS_MULT, [False, True, False])):
    c = ws.cell(row=4, column=2 + j, value=("=" + A["mult"]) if link else m)
    c.number_format = MULT2; c.font = Font(bold=True, color="FFFFFF"); c.fill = FILL_HDR; c.alignment = RIGHT
for i, (e, link) in enumerate(zip(SENS_REV, [False, True, False])):
    rr = 5 + i
    c = ws.cell(row=rr, column=1, value="=Financials!F5" if link else e)
    c.number_format = M0; c.font = F_LINK if link else F_IN
    for j in range(3):
        col = get_column_letter(2 + j)
        ws.cell(row=rr, column=2 + j, value=f"=($A{rr}*{col}$4-Valuation!$B$8)/Valuation!$B$9").number_format = USD2
ws["C6"].fill = FILL_KEY
ws["A9"] = "Cells above the reference price"; ws["A9"].font = F_B
ws["B9"] = f'=COUNTIF(B5:D7,">"&{A["price"]})'; ws["C9"] = "of 9"
ws["A10"] = "Cells 15% or more above the price"; ws["A10"].font = F_B
ws["B10"] = f'=COUNTIF(B5:D7,">="&{A["price"]}*1.15)'; ws["C10"] = "of 9"
note(ws, 12, ["Middle row links to the model's FY28E revenue and the middle column to the base multiple, so the grid moves with the assumptions."])

# =====================================================================  QUARTERS
ws = sheet("Quarters", "QUARTERLY RESULTS, BACKLOG AND THE FY26 CUTS (US$ million unless stated)", [52] + [12] * 7)
head(ws, 4, ["Quarter (fiscal year to 30 September)"] + Q)
cols = [get_column_letter(2 + j) for j in range(7)]


def qrow(r, label, vals, fmt, inp=True, bold=False):
    ws.cell(row=r, column=1, value=label).font = F_B if bold else Font()
    for j, v in enumerate(vals):
        c = ws.cell(row=r, column=2 + j, value=v)
        c.number_format = fmt; c.alignment = RIGHT
        if inp and not (isinstance(v, str) and v.startswith("=")): c.font = F_IN


qrow(5, "Revenue", Q_REV, M1, bold=True)
qrow(6, "GAAP gross profit", Q_GP, M1)
qrow(7, "GAAP gross margin (derived)", [f"={c}6/{c}5" for c in cols], PCT1, inp=False)
qrow(8, "Adjusted gross margin, as reported", [m / 100 for m in Q_AGM], PCT1)
qrow(9, "Adjusted EBITDA", Q_EBITDA, M1)
qrow(10, "Backlog at quarter end, US$bn", BACKLOG, "0.0")
ws["A12"] = "Nine months to 30 June 2026"; ws["A12"].font = Font(bold=True, color=NAVY)
ws["A13"] = "Revenue"; ws["B13"] = "=SUM(F5:H5)"
ws["A14"] = "Adjusted gross profit (reported)"; ws["B14"] = AGP_9M; ws["B14"].font = F_IN
ws["A15"] = "Adjusted EBITDA"; ws["B15"] = "=SUM(F9:H9)"
ws["A16"] = "Net operating costs (derived)"; ws["B16"] = "=B14-B15"
ws["A17"] = "Revenue from customers in the Americas"; ws["B17"] = AMERICAS_9M; ws["B17"].font = F_IN
ws["A18"] = "Americas share of revenue"; ws["B18"] = "=B17/B13"; ws["B18"].number_format = PCT1
for c in ("B13", "B14", "B15", "B16", "B17"): ws[c].number_format = M1
ws["A20"] = "September 2026 quarter implied by the FY26 guide"; ws["A20"].font = Font(bold=True, color=NAVY)
ws["A21"] = "Revenue"; ws["B21"] = "=" + A["g26r"] + "-B13"
ws["A22"] = "Adjusted EBITDA"; ws["B22"] = "=" + A["g26e"] + "-B15"
for c in ("B21", "B22"): ws[c].number_format = M1
ws["A24"] = "Demand indicators (Form 10-Q)"; ws["A24"].font = Font(bold=True, color=NAVY)
head(ws, 25, ["Metric", "Latest", "Year earlier", "Change"])
for i, (lab, a, b) in enumerate([("Storage contracted, nine months, GW", GW_ORD_9M, GW_ORD_9M_PY),
                                  ("Storage pipeline, GWh (30 Jun 2026 vs 30 Sep 2025)", PIPE_GWH, PIPE_GWH_PY)]):
    rr = 26 + i
    ws.cell(row=rr, column=1, value=lab)
    for j, v in enumerate([a, b]):
        c = ws.cell(row=rr, column=2 + j, value=v); c.font = F_IN; c.number_format = "0.0"
    ws.cell(row=rr, column=4, value=f"=B{rr}/C{rr}-1").number_format = PCT0
note(ws, 29, [
    "Sources: Fluence quarterly results releases, 10 February 2025 to 5 August 2026; Form 10-Q for the quarter to 30 June 2026 (revenue by region,",
    "key operating metrics). Adjusted gross margin as first reported in each release. Order intake in US$: 778.0 (Dec 24), 200.0 (Mar 25), 508.8 (Jun 25),",
    "over 1,400 (Sep 25), over 750 (Dec 25), over 1,440 (Jun 26); the March 2026 quarter was not given on its own.",
])

# =====================================================================  PEERS
ws = sheet("Peers", "PEER MULTIPLES (stockanalysis.com, RETRIEVED 12 OCTOBER 2026; TRAILING)", [22, 16, 12, 14, 70])
head(ws, 4, ["Company", "Ticker", "EV/sales", "EV/EBITDA", "What they make"])
for i, (co, tk, evs, eve, what) in enumerate(PEERS):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=co).font = F_B if co == "Fluence" else Font()
    ws.cell(row=rr, column=2, value=tk)
    for j, v in enumerate([evs, eve]):
        c = ws.cell(row=rr, column=3 + j, value=v if v is not None else "n.m."); c.number_format = MULT2 if j == 0 else MULT; c.font = F_IN
        c.alignment = RIGHT
    ws.cell(row=rr, column=5, value=what).alignment = WRAP
ws["A12"] = "Peer median (four peers with positive EBITDA; excludes Fluence and Eos Energy)"; ws["A12"].font = F_B
ws["C12"] = "=MEDIAN(C6:C9)"; ws["D12"] = "=MEDIAN(D6:D9)"
ws["C12"].number_format = MULT2; ws["D12"].number_format = MULT
ws["A13"] = "Fluence base case on FY28E EBITDA"; ws["D13"] = "=Valuation!B23"; ws["D13"].number_format = MULT
note(ws, 15, ["Multiples only, to set the range. No view on the peers' shares is expressed. n.m. = not meaningful (negative EBITDA).",
              "Fluence's own trailing EV/sales as computed by stockanalysis.com; the model's EV uses my net debt estimate."])

# =====================================================================  SOURCES
ws = sheet("Sources", "SOURCES", [6, 140])
head(ws, 4, ["No.", "Source and what it supports"])
srcs = [
    "Fluence quarterly results releases, Form 8-K Exhibit 99.1: 10 February 2025, 7 May 2025, 11 August 2025, 24 November 2025 (FY25 and first FY26 guide), 4 February 2026, 6 May 2026, 5 August 2026 (first FY26 cut). Revenue, gross and adjusted gross profit, adjusted EBITDA, free cash flow, orders, backlog, liquidity.",
    "Fluence Form 8-K of 16 September 2026, Exhibit 99.1 (revised FY26 guide; Houston named as the primary reason; plans to report US production, actual and forecast; FY27 plan with FY26 results) and Exhibit 99.2 (Bernerd Da Santos as chief operating officer).",
    "Fluence Form 10-Q for the quarter to 30 June 2026, filed 5 August 2026: balance sheet, shares, convertible notes (Note 12), backlog conversion, key operating metrics, revenue by region, liquidated damages, credit agreement covenants, shelf registration, secondary offering.",
    "Fluence Form 8-K of 15 May 2026 (AES redemption) and Form 8-K of 9 October 2026 (separation of the chief product officer).",
    "Fluence third quarter fiscal 2026 earnings call, 6 August 2026, transcript published by The Motley Fool on 12 August 2026: Houston capacity and timing, generators, US$2.2bn of backlog scheduled for FY27, data centre pipeline and awards.",
    "Latitude Media (21 September 2026), mgrid.org (18 September 2026) and the Quartr summary of the 16 September 2026 call: Houston unit rates, the American share of the cut, penalties, quarterly US production reporting. Not verified against the replay.",
    "Nasdaq.com historical data, retrieved 12 October 2026: daily and month-end closes, highs and lows.",
    "stockanalysis.com (S&P Global consensus), retrieved 12 October 2026: consensus revenue, free cash flow, targets, ratings, short interest, peer multiples.",
    "The Physical Layer, 'Fluence has plenty of batteries. What it is short of is American-made ones', 12 October 2026: https://thephysicallayer.fyi/journal/fluence-short-of-american-made/",
    "Estimates for FY27 and FY28 and the September 2026 quarter cash use are the author's own. Personal research, not investment advice.",
]
for i, s in enumerate(srcs):
    ws.cell(row=5 + i, column=1, value=i + 1)
    ws.cell(row=5 + i, column=2, value=s).alignment = WRAP

# =====================================================================  COVER
ws = wb["Cover"]
head(ws, 4, ["Item", "Value", "Note"])
cover = [
    (("Draft view (the call is Tommy Lau's)" if CALL["draft"] else "Call (Tommy Lau)"), ("=" + V["rule"]) if CALL["draft"] else CALL["direction"], None, "Applied on the band to the base value: LONG at +15% or more, SHORT at -25% or worse, NO CALL between." + (" DRAFT." if CALL["draft"] else "")),
    ("Base value per share, US$", "=" + V["base"], USD2, f"{BASE_MULT:.2f}x FY28E revenue less estimated net debt, over Class A plus Class B-1 shares."),
    ("Reference price, US$", "=" + A["price"], USD2, "Nasdaq close, 9 October 2026."),
    ("Base value against price", "=" + V["up"], PCT1, ""),
    ("Probability-weighted value, US$", "=Scenarios!E11", USD2, "25 / 50 / 25 bear / base / bull."),
    (("Conviction (draft)" if CALL["draft"] else "Conviction"), CALL["conviction"], None, "Judgement, not formula."),
    ("Horizon", "12 months, to October 2027", None, ""),
    ("Revisit if", "See note", None, REVISIT_IF),
    ("Wrong if", "See note", None, WRONG_IF),
]
for i, (lab, v, fmt, nt) in enumerate(cover):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=lab).font = F_B
    c = ws.cell(row=rr, column=2, value=v)
    if fmt: c.number_format = fmt
    c.alignment = RIGHT
    c.font = F_LINK if isinstance(v, str) and v.startswith("=") else F_IN
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
ws.row_dimensions[12].height = 44; ws.row_dimensions[13].height = 32
ws["A16"] = "Tabs"; ws["A16"].font = Font(bold=True, color=NAVY)
for i, (t, d) in enumerate([
    ("Assumptions", "Every input, with its source. Change the yellow cells to move the value."),
    ("Financials", "FY24A to FY28E: revenue, adjusted gross profit, net operating costs, adjusted EBITDA, EV multiples."),
    ("Valuation", "Base value, the band, what the price needs, cross-checks, net debt."),
    ("Scenarios", "Bear / base / bull and the probability-weighted value."),
    ("Sensitivity", "FY28E revenue against EV/revenue multiple."),
    ("Quarters", "Quarterly results, backlog, the nine months to June 2026, the implied September quarter, demand indicators."),
    ("Peers", "EV/sales and EV/EBITDA used to set the multiple."),
    ("Sources", "Every source used."),
]):
    ws.cell(row=17 + i, column=1, value=t)
    ws.cell(row=17 + i, column=2, value=d)
ws["A26"] = "Colour key: blue = hardcoded input; yellow fill = my own assumption; green = link to another tab; black = formula."
ws["A26"].font = F_NOTE
ws["A27"] = f"Companion note: {DATE}_{COMPANY}_Initiation.pdf. Personal research, not investment advice."
ws["A27"].font = F_NOTE

for w in wb.worksheets:
    w.freeze_panes = "A5"
wb.active = 0
wb.save(XLSX)
print("model written", XLSX)
