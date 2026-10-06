"""Bloom Energy initiation model, 6 Oct 2026. Live formulas throughout.

House style follows the Vertiv, TSMC and Texas Instruments models: navy header rows, blue font for hardcoded inputs,
yellow fill on my own assumptions, green font for cross-sheet links, black for formulas. Non-GAAP basis.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from data import *

NAVY = "1F3864"; MUTED = "595959"
F_TTL = Font(bold=True, color=NAVY, size=12)
F_SUB = Font(color=MUTED, size=9)
F_HDR = Font(bold=True, color="FFFFFF", size=10)
F_IN = Font(color="0000FF")
F_INB = Font(color="0000FF", bold=True)
F_LINK = Font(color="008000")
F_B = Font(bold=True)
F_NOTE = Font(color=MUTED, size=9, italic=True)
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_ASSUME = PatternFill("solid", fgColor="FFFF00")
FILL_KEY = PatternFill("solid", fgColor="D9E2F3")
WRAP = Alignment(wrap_text=True, vertical="top")
RIGHT = Alignment(horizontal="right")

USD2 = '"US$"#,##0.00'; USD0 = '"US$"#,##0'; BN2 = '#,##0.00'; BN3 = '#,##0.000'; BN1 = '#,##0.0'
PCT1 = '0.0%'; PCT0 = '0%'; MULT = '0.0"x"'; MULT0 = '0"x"'
SUB = f"Tommy Lau | Bloom Energy (NYSE: BE) | {DATE_LONG} | Non-GAAP basis | Personal research, not investment advice."

wb = openpyxl.Workbook()


def sheet(name, title, widths):
    ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet(name)
    ws.title = name
    ws["A1"] = title; ws["A1"].font = F_TTL
    ws["A2"] = SUB; ws["A2"].font = F_SUB
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def head(ws, row, cells, start=1):
    for i, c in enumerate(cells, start):
        x = ws.cell(row=row, column=i, value=c)
        x.font = F_HDR; x.fill = FILL_HDR
        x.alignment = Alignment(horizontal="left" if i == start else "right", vertical="center", wrap_text=True)


def note(ws, row, texts):
    for i, t in enumerate(texts):
        ws.cell(row=row + i, column=1, value=t).font = F_NOTE


sheet("Cover", "BLOOM ENERGY (NYSE: BE)  INITIATION MODEL", [38, 24, 80])

# ===================================================================== ASSUMPTIONS
ws = sheet("Assumptions", "ASSUMPTIONS AND INPUTS", [56, 14, 10, 92])
ws["A3"] = ("Blue = hardcoded input. Yellow fill = my own assumption (change these to move the value). Green = link to "
            "another tab. Black = formula. US$ billion unless stated; per-share figures in US$. Non-GAAP basis unless marked GAAP.")
ws["A3"].font = F_NOTE
head(ws, 5, ["Input", "Value", "Unit", "Source / note"])
A = {}
MINE = {"q3rev", "q4rev", "q3gm", "q4gm", "q3opex", "q4opex", "otherq", "othery", "tax26", "tax27", "tax28", "sbcdil",
        "prodshare", "pe"}
rows = [
    ("MARKET (retrieved 6 October 2026)", None, None, None, None),
    ("price", "Share price (NYSE: BE)", PRICE, "US$", "NYSE close, 5 October 2026. Yahoo Finance."),
    ("hi52", "52-week high (intraday, 25 June 2026)", HI52, "US$", "Yahoo Finance."),
    ("hiclose", "Highest close (22 June 2026)", HI_CLOSE, "US$", "Yahoo Finance."),
    ("pxdec25", "Month-end close, 31 Dec 2025", PRICE_DEC25, "US$", "Yahoo Finance."),
    ("shout", "Shares outstanding, 22 Jul 2026", SHARES_OUT, "bn", "10-Q cover: 294,527,346."),
    ("cash", "Cash and equivalents, 30 Jun 2026", CASH, "US$bn", "Q2 2026 release: 2,666.9m (excludes restricted cash)."),
    ("debt", "Debt principal, 30 Jun 2026", DEBT, "US$bn", "10-Q: 0% 2030 notes 2,500.0m; Green 2029 27.0m; Green 2028 0.8m (redeemed 10 Jul); Korea loan 2.6m."),
    ("finob", "Financing obligations, 30 Jun 2026 (not in net cash)", FIN_OBLIG, "US$bn", "Old sale-leaseback structures: 62.0m current + 144.4m long-term."),
    ("cons26", "2026 consensus EPS", CONS_EPS_26, "US$", "stockanalysis.com, 25 analysts, retrieved 6 October 2026."),
    ("cons27", "2027 consensus EPS", CONS_EPS_27, "US$", "Nasdaq.com (Zacks), 6 estimates."),
    ("cons28", "2028 consensus EPS", CONS_EPS_28, "US$", "Nasdaq.com (Zacks), 2 estimates."),
    ("consq3", "Q3 2026 consensus EPS", CONS_Q3, "US$", "Nasdaq.com (Zacks), 4 estimates."),
    ("fwdpe", "Forward P/E shown by stockanalysis.com", FWD_PE_SA, "x", "stockanalysis.com, 6 October 2026."),
    ("constp", "Average analyst target", CONS_TP, "US$", "stockanalysis.com, 29 analysts; median US$303, range US$97 to 390."),
    ("REPORTED (Bloom releases, Form 10-K, Form 10-Q)", None, None, None, None),
    ("h1rev", "Revenue, H1 2026", H1_26_REV, "US$bn", "Q2 2026 release: 751.1 + 1,065.4."),
    ("h1gp", "Non-GAAP gross profit, H1 2026", H1_26_GP, "US$bn", "Q2 2026 release: 236.3 + 365.4."),
    ("h1opex", "Non-GAAP operating expenses, H1 2026", H1_26_OPEX, "US$bn", "Q2 2026 release: 106.6 + 125.7."),
    ("h1opinc", "Non-GAAP operating income, H1 2026", H1_26_OPINC, "US$bn", "Q2 2026 release: 129.7 + 239.6."),
    ("h1ni", "Adjusted net profit, H1 2026", H1_26_NI, "US$bn", "Q2 2026 release: 138.1 + 248.2."),
    ("h1eps", "Non-GAAP diluted EPS, H1 2026", H1_26_EPS, "US$", "Q2 2026 release: 0.44 + 0.78."),
    ("grevlo", "2026 revenue guide, low", GUIDE_26["rev"][0], "US$bn", "Release of 28 July 2026."),
    ("grevhi", "2026 revenue guide, high", GUIDE_26["rev"][1], "US$bn", "Same."),
    ("gepslo", "2026 non-GAAP EPS guide, low", GUIDE_26["eps"][0], "US$", "Same."),
    ("gepshi", "2026 non-GAAP EPS guide, high", GUIDE_26["eps"][1], "US$", "Same."),
    ("gopinclo", "2026 non-GAAP operating income guide, low", GUIDE_26["opinc"][0], "US$bn", "Same."),
    ("gopinchi", "2026 non-GAAP operating income guide, high", GUIDE_26["opinc"][1], "US$bn", "Same."),
    ("sbc", "Stock-based compensation in operating costs, Q2 2026", SBC_Q2, "US$bn", "Q2 2026 reconciliation: 56.4m."),
    ("opincq2", "Non-GAAP operating income, Q2 2026", 0.2396, "US$bn", "Q2 2026 release."),
    ("aepusd", "AEP order value", AEP_DEAL[0], "US$bn", "Exercise of most of AEP's 900 MW option, reported 8 January 2026."),
    ("aepgw", "AEP option size", AEP_DEAL[1], "GW", "900 MW balance of AEP's 1 GW agreement (Nov 2024)."),
    ("capgw", "Fremont annual capacity, end-2026 plan", CAPACITY_GW[1], "GW", "Form 10-K for 2025: doubling from 1 GW to 2 GW by end of 2026."),
    ("MY ASSUMPTIONS", None, None, None, None),
    ("q3rev", "Revenue, Q3 2026", Q3E_REV, "US$bn", "MINE. Guide implies H2 of US$2.08 to 2.38bn."),
    ("q4rev", "Revenue, Q4 2026", Q4E_REV, "US$bn", "MINE."),
    ("q3gm", "Non-GAAP gross margin, Q3 2026", Q3E_GM, "%", "MINE. Guide about 34% for the year; H1 33.1%."),
    ("q4gm", "Non-GAAP gross margin, Q4 2026", Q4E_GM, "%", "MINE."),
    ("q3opex", "Non-GAAP operating expenses, Q3 2026", Q3E_OPEX, "US$bn", "MINE. Q2: 0.126."),
    ("q4opex", "Non-GAAP operating expenses, Q4 2026", Q4E_OPEX, "US$bn", "MINE."),
    ("otherq", "Adjusted net profit less non-GAAP operating income, per quarter, before tax", OTHER_Q, "US$bn", "MINE. Q1 and Q2 2026: about 0.008 (net interest income)."),
    ("othery", "Same, per year, 2027 and 2028", OTHER_Y, "US$bn", "MINE."),
    ("tax26", "Non-GAAP tax rate, H2 2026", TAX26, "%", "MINE. H1 2026 effective rate 0.7% (valuation allowance)."),
    ("tax27", "Non-GAAP tax rate, 2027", TAX27, "%", "MINE. Losses carried forward run down."),
    ("tax28", "Non-GAAP tax rate, 2028", TAX28, "%", "MINE."),
    ("sbcdil", "Net new shares a year from staff awards", SBC_DIL, "%", "MINE."),
    ("prodshare", "Product share of revenue, for megawatt conversion", PROD_SHARE, "%", "MINE. H1 2026: 87.5%."),
    ("pe", "Target multiple on 2028E EPS (base)", BASE_PE, "x", "MINE. Below GE Vernova, above Vertiv (Peers tab)."),
]
r = 6
for key, label, val, unit, src in rows:
    if label is None:
        ws.cell(row=r, column=1, value=key).font = Font(bold=True, color=NAVY); r += 1; continue
    ws.cell(row=r, column=1, value=label)
    c = ws.cell(row=r, column=2, value=val)
    mine = key in MINE
    c.font = F_INB if mine else F_IN
    if mine: c.fill = FILL_ASSUME
    c.number_format = PCT1 if unit == "%" else (MULT0 if unit == "x" else (BN3 if unit in ("bn", "GW") else "#,##0.000"))
    ws.cell(row=r, column=3, value=unit)
    ws.cell(row=r, column=4, value=src).alignment = WRAP
    A[key] = f"Assumptions!$B${r}"
    r += 1

# ===================================================================== DILUTION
ws = sheet("Dilution", "SHARE COUNT, COUNTED FULLY DILUTED", [66, 16, 70])
head(ws, 4, ["Line", "Value", "Working / note"])
dil = [
    ("Shares outstanding, 22 Jul 2026, bn", "=" + A["shout"], BN3, "10-Q cover."),
    ("0% notes due Nov 2030, principal, US$bn", CONV_2030, BN3, "Issued 4 Nov 2025."),
    ("Conversion rate, shares per US$1,000", CONV_2030_RATE, '0.0000', "Conversion price US$194.97."),
    ("Shares on conversion, bn", "=B6*B7/1000", BN3, "Up to 19.554m after a make-whole fundamental change."),
    ("3.0% Green notes due Jun 2029, principal outstanding, US$bn", CONV_2029, BN3, "After US$147m of Green notes converted in H1 2026."),
    ("Conversion rate, shares per US$1,000", CONV_2029_RATE, '0.0000', "Conversion price US$20.84."),
    ("Shares on conversion, bn", "=B9*B10/1000", BN3, ""),
    ("Stock options and awards in the Q2 2026 diluted count, bn", OPTIONS_Q2, BN3, "10-Q Note 15, treasury stock method."),
    ("Fully diluted shares, bn", "=B5+B8+B11+B12", BN3, "Used for H2 2026 EPS; Q2 2026 reported diluted: 323.3m."),
    ("Note conversions against shares outstanding", "=(B8+B11)/B5", PCT1, ""),
    (None, None, None, ""),
    ("Early-conversion trigger price (130% of US$194.97), US$", f"=1.3*{CONV_2030_PX}", USD2, "20 of the last 30 trading days of a quarter."),
    ("Q3 2026 trading days of the last 30 closing above it", CONV_TRIGGER_DAYS, "0", "My count of Yahoo Finance closes; 20 needed. Not verified against the trustee."),
    ("Stock-based compensation, Q2 2026, US$bn", "=" + A["sbc"], BN3, "Excluded from non-GAAP figures."),
    ("Against non-GAAP operating income, Q2 2026", f"=B18/{A['opincq2']}", PCT1, ""),
    ("Oracle warrant: shares issued on exercise, m", ORCL_WARRANT[4], '0.000', "1,905,433 on net exercise plus 248,798 inducement shares, 1 May 2026."),
    ("Oracle warrant: value to come off revenue, US$bn", 0.3244, BN3, "Amortised as Oracle's systems are delivered; US$17.9m so far."),
]
for i, (label, f, fmt, nt) in enumerate(dil):
    rr = 5 + i
    if label is None: continue
    ws.cell(row=rr, column=1, value=label)
    c = ws.cell(row=rr, column=2, value=f); c.number_format = fmt
    if not (isinstance(f, str) and f.startswith("=")): c.font = F_IN
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
SHFD = "Dilution!$B$13"

# ===================================================================== QUARTERLY
ws = sheet("Quarterly", "QUARTERLY HISTORY, Q1 2024 TO Q2 2026, AND MY Q3 / Q4 2026", [50] + [10] * 12)
QH = Q + ["Q3 26E", "Q4 26E"]
head(ws, 4, ["US$ billion unless stated"] + QH)
C = [get_column_letter(2 + j) for j in range(12)]


def qrow(row, label, vals, fmt, bold=False):
    ws.cell(row=row, column=1, value=label).font = F_B if bold else Font()
    for j, v in enumerate(vals):
        c = ws.cell(row=row, column=2 + j, value=v if v is not None else "")
        c.number_format = fmt; c.alignment = RIGHT
        if v is not None and not (isinstance(v, str) and v.startswith("=")):
            c.font = F_IN


bn = lambda xs: [x / 1000 for x in xs]
pct = lambda xs: [x / 100 for x in xs]
E2 = [None, None]
qrow(5, "Revenue", bn(Q_REV) + ["=" + A["q3rev"], "=" + A["q4rev"]], BN3, True)
qrow(6, "Product revenue", bn(Q_PROD) + E2, BN3)
qrow(7, "Product cost of revenue", bn(Q_PROD_COST) + E2, BN3)
qrow(8, "Product gross margin, GAAP (derived)", [f"=({c}6-{c}7)/{c}6" for c in C[:10]] + E2, PCT1, True)
qrow(9, "Service revenue", bn(Q_SERV) + E2, BN3)
qrow(10, "Service cost of revenue", bn(Q_SERV_COST) + E2, BN3)
qrow(11, "Service gross margin, GAAP (derived)", [f"=({c}9-{c}10)/{c}9" for c in C[:10]] + E2, PCT1)
qrow(12, "Gross margin, GAAP (reported)", pct(Q_GM_GAAP) + E2, PCT1)
qrow(13, "Gross margin, non-GAAP (reported; Q3 / Q4 26 mine)", pct(Q_GM) + ["=" + A["q3gm"], "=" + A["q4gm"]], PCT1, True)
qrow(14, "Non-GAAP operating income (Q3 / Q4 26 mine)", [None] * 10 + [f"=L5*L13-{A['q3opex']}", f"=M5*M13-{A['q4opex']}"], BN3)
qrow(15, "Non-GAAP diluted EPS, US$ (reported; Q3 / Q4 26 mine)", Q_EPS + [
    f"=(L14+{A['otherq']})*(1-{A['tax26']})/{SHFD}", f"=(M14+{A['otherq']})*(1-{A['tax26']})/{SHFD}"], USD2, True)
qrow(16, "Revenue to related parties", bn(Q_RELATED) + E2, BN3)
qrow(17, "Related-party share of revenue (derived)", [f"={c}16/{c}5" for c in C[:10]] + E2, PCT1)
ws["A19"] = "Checks"; ws["A19"].font = Font(bold=True, color=NAVY)
ws["A20"] = "Trailing four quarters non-GAAP EPS to Q2 26"; ws["B20"] = "=SUM(H15:K15)"; ws["B20"].number_format = USD2
ws["A21"] = "Product gross margin, H1 2026"; ws["B21"] = "=(J6+K6-J7-K7)/(J6+K6)"; ws["B21"].number_format = PCT1
ws["A22"] = "Revenue, H1 2026 against H1 2024 (times)"; ws["B22"] = "=(J5+K5)/(B5+C5)"; ws["B22"].number_format = '0.00"x"'
ws["A23"] = "Revenue growth, Q2 2026 on Q2 2025"; ws["B23"] = "=K5/G5-1"; ws["B23"].number_format = PCT0
ws["A24"] = "H2 2026E revenue against the guide's implied H2 (low / high)"; ws["B24"] = f"=L5+M5"; ws["C24"] = f"={A['grevlo']}-{A['h1rev']}"; ws["D24"] = f"={A['grevhi']}-{A['h1rev']}"
for c_ in ("B24", "C24", "D24"): ws[c_].number_format = BN2
note(ws, 26, ["Source: Bloom results releases (Exhibit 99.1 to Form 8-K): Q2 2024 (with Q1 2024), Q2 2025 (with Q1 2025), Q3 2025 (with Q3 2024), Q4 2025 (with Q4 2024), Q2 2026 (with Q1 2026).",
              "Q3 26E and Q4 26E are my estimates: EPS = (revenue x gross margin - operating costs + other income) x (1 - tax) / fully diluted shares (Dilution tab)."])

# ===================================================================== FINANCIALS
ws = sheet("Financials", "FINANCIAL SUMMARY AND FORECASTS (US$ billion unless stated, non-GAAP basis)", [52, 12, 12, 12, 12, 12])
head(ws, 4, ["", "2024A", "2025A", "2026E*", "2027E*", "2028E*"])
fin = [
    ("Revenue", [REV_H[0], REV_H[1], f"={A['h1rev']}+Quarterly!L5+Quarterly!M5", "=D5*(1+Scenarios!B6)", "=E5*(1+Scenarios!C6)"], BN2, True),
    ("Revenue growth", [None, "=C5/B5-1", "=D5/C5-1", "=E5/D5-1", "=F5/E5-1"], PCT1, False),
    ("Product revenue", [PROD_H[0], PROD_H[1], "=Quarterly!J6+Quarterly!K6", None, None], BN2, False),
    ("Product cost of revenue", [PROD_COST_H[0], PROD_COST_H[1], "=Quarterly!J7+Quarterly!K7", None, None], BN2, False),
    ("Product gross margin, GAAP (2026: H1)", ["=(B7-B8)/B7", "=(C7-C8)/C7", "=(D7-D8)/D7", None, None], PCT1, True),
    ("Gross profit", ["=B5*0.287", "=C5*0.303", f"={A['h1gp']}+Quarterly!L5*Quarterly!L13+Quarterly!M5*Quarterly!M13", "=E5*Scenarios!D6", "=F5*Scenarios!E6"], BN2, False),
    ("Gross margin", ["=B10/B5", "=C10/C5", "=D10/D5", "=E10/E5", "=F10/F5"], PCT1, True),
    ("Operating expenses", ["=B10-B13", "=C10-C13", f"={A['h1opex']}+{A['q3opex']}+{A['q4opex']}", "=Scenarios!F6", "=Scenarios!G6"], BN2, False),
    ("Operating income", [OPINC_H[0], OPINC_H[1], "=D10-D12", "=E10-E12", "=F10-F12"], BN2, True),
    ("Operating margin", ["=B13/B5", "=C13/C5", "=D13/D5", "=E13/E5", "=F13/F5"], PCT1, False),
    ("Other income less operating income gap (2026: H1 actual plus H2)", [None, None, f"=({A['h1ni']}-{A['h1opinc']})+2*{A['otherq']}", "=" + A["othery"], "=" + A["othery"]], BN3, False),
    ("Tax rate (2026: H2)", [None, None, "=" + A["tax26"], "=" + A["tax27"], "=" + A["tax28"]], PCT0, False),
    ("Adjusted net profit", [None, None, f"={A['h1ni']}+(Quarterly!L14+{A['otherq']})*(1-D16)+(Quarterly!M14+{A['otherq']})*(1-D16)", "=(E13+E15)*(1-E16)", "=(F13+F15)*(1-F16)"], BN3, True),
    ("Fully diluted shares, bn", [None, None, "=" + SHFD, f"=D18*(1+{A['sbcdil']})", f"=E18*(1+{A['sbcdil']})"], BN3, False),
    ("Diluted EPS, US$", [EPS_H[0], EPS_H[1], f"={A['h1eps']}+Quarterly!L15+Quarterly!M15", "=E17/E18", "=F17/F18"], USD2, True),
    ("EPS growth", [None, "=C19/B19-1", "=D19/C19-1", "=E19/D19-1", "=F19/E19-1"], PCT1, False),
    ("P/E at reference price", [f"={A['price']}/B19", f"={A['price']}/C19", f"={A['price']}/D19", f"={A['price']}/E19", f"={A['price']}/F19"], MULT, False),
    ("Consensus EPS, US$", [None, None, "=" + A["cons26"], "=" + A["cons27"], "=" + A["cons28"]], USD2, False),
    ("Mine against consensus", [None, None, "=D19/D22-1", "=E19/E22-1", "=F19/F22-1"], PCT1, False),
    ("Product megawatt-equivalents, GW (derived)", [None, None, "=D5*Capacity!$B$7/Capacity!$B$6", "=E5*Capacity!$B$7/Capacity!$B$6", "=F5*Capacity!$B$7/Capacity!$B$6"], BN2, False),
]
for i, (label, vals, fmt, bold) in enumerate(fin):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=label).font = F_B if bold else Font()
    for j, v in enumerate(vals):
        c = ws.cell(row=rr, column=2 + j, value=v if v is not None else "")
        c.number_format = fmt; c.alignment = RIGHT
        if v is not None and not isinstance(v, str): c.font = F_IN
        elif isinstance(v, str) and v.startswith("=Assumptions"): c.font = F_LINK
note(ws, 26, [
    "*2026E to 2028E are my own estimates. 2026E = reported H1 plus my Q3 and Q4 (Quarterly tab). 2027E and 2028E use the base case on the Scenarios tab (row 6).",
    "2024A and 2025A gross profit use the reported non-GAAP gross margins (28.7%, 30.3%); operating income and EPS as reported in the Q4 2025 release.",
    f"Note check: 2026E revenue {REV26:.2f}, EPS {EPS26:.2f}; 2027E {REV27:.2f}, {EPS27:.2f}; 2028E {REV28:.2f}, {EPS28:.2f}.",
])

# ===================================================================== SCENARIOS
ws = sheet("Scenarios", "THREE CASES, TWELVE MONTHS OUT", [30] + [11] * 15)
HDR = ["Case", "Growth 27", "Growth 28", "GM 27", "GM 28", "Opex 27", "Opex 28", "2027 revenue", "2028 revenue",
       "2028 op. income", "2028 net profit", "2028 EPS", "Multiple", "Value, US$", "vs price", "2028 GW"]
head(ws, 4, HDR)
R26 = "Financials!$D$5"


def scen_row(rr, name, vals, mult):
    ws.cell(row=rr, column=1, value=name).font = F_B
    for j, v in enumerate(vals):
        c = ws.cell(row=rr, column=2 + j, value=v)
        if not isinstance(v, str):
            c.font = F_INB; c.fill = FILL_ASSUME
        c.number_format = PCT1 if j < 4 else BN2
    ws.cell(row=rr, column=8, value=f"={R26}*(1+B{rr})").number_format = BN2
    ws.cell(row=rr, column=9, value=f"=H{rr}*(1+C{rr})").number_format = BN2
    ws.cell(row=rr, column=10, value=f"=I{rr}*E{rr}-G{rr}").number_format = BN2
    ws.cell(row=rr, column=11, value=f"=(J{rr}+{A['othery']})*(1-{A['tax28']})").number_format = BN3
    ws.cell(row=rr, column=12, value=f"=K{rr}/Financials!$F$18").number_format = USD2
    c = ws.cell(row=rr, column=13, value=("=" + A["pe"]) if mult == "pe" else mult)
    if mult == "pe": c.font = F_LINK
    else: c.font = F_INB; c.fill = FILL_ASSUME
    c.number_format = MULT0
    ws.cell(row=rr, column=14, value=f"=L{rr}*M{rr}").number_format = USD2
    ws.cell(row=rr, column=15, value=f"=N{rr}/{A['price']}-1").number_format = PCT0
    ws.cell(row=rr, column=16, value=f"=I{rr}*Capacity!$B$7/Capacity!$B$6").number_format = BN2


for i, s in enumerate(SCEN):
    scen_row(5 + i, s[0], list(s[1:7]), "pe" if s[0] == "Base" else s[7])
ws["A10"] = "Probability"; ws["A10"].font = F_B
for j, s in enumerate(SCEN):
    ws.cell(row=9, column=2 + j, value=s[0]).font = F_B
    c = ws.cell(row=10, column=2 + j, value=s[8]); c.font = F_INB; c.fill = FILL_ASSUME; c.number_format = PCT0
ws["E10"] = "=SUM(B10:D10)"; ws["E10"].number_format = PCT0; ws["F10"] = "must be 100%"; ws["F10"].font = F_NOTE
ws["A12"] = "Probability-weighted value, US$"; ws["A12"].font = F_B
ws["N12"] = "=B10*N5+C10*N6+D10*N7"; ws["N12"].number_format = USD2; ws["N12"].font = F_B; ws["N12"].fill = FILL_KEY
ws["O12"] = f"=N12/{A['price']}-1"; ws["O12"].number_format = PCT1
note(ws, 14, ["All cases start from my 2026E and use 2028 fully diluted shares from the Financials tab. Gross margin is non-GAAP; Opex is non-GAAP operating costs, US$bn.",
              "Bear: faster grid connections and more turbine supply narrow the window Bloom sells into. Bull: Oracle's 2.8 GW and Brookfield's US$25bn fill, Fremont expands fast.",
              "The base row (6) drives the Financials tab for 2027E and 2028E."])

# ===================================================================== CAPACITY
ws = sheet("Capacity", "THE FACTORY IS THE CEILING: REVENUE CONVERTED TO GIGAWATTS (DERIVED)", [66, 16, 70])
head(ws, 4, ["Line", "Value", "Working / note"])
cap = [
    ("AEP order, US$bn", "=" + A["aepusd"], BN2, "Reported 8 January 2026."),
    ("Price per megawatt, US$m", f"={A['aepusd']}/{A['aepgw']}", BN2, "My unit for converting revenue to megawatts. Bloom does not report MW shipped."),
    ("Product share of revenue", "=" + A["prodshare"], PCT1, "H1 2026."),
    ("2026E, GW", "=Financials!D24", BN2, ""),
    ("2027E, GW", "=Financials!E24", BN2, ""),
    ("2028E base, GW", "=Financials!F24", BN2, ""),
    ("Fremont annual capacity by end 2026, GW", "=" + A["capgw"], BN1, "Site room for about 5 GW (10-K)."),
    ("Revenue a 2 GW factory supports at that price, US$bn", "=B11*B6/B7", BN2, ""),
    ("2028 GW the price needs (Valuation tab)", "=Valuation!B29", BN2, ""),
]
for i, (label, f, fmt, nt) in enumerate(cap):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=label)
    c = ws.cell(row=rr, column=2, value=f); c.number_format = fmt
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP

# ===================================================================== VALUATION
ws = sheet("Valuation", "VALUATION: 40x 2028E EPS, TWELVE MONTHS OUT", [62, 16, 70])
head(ws, 4, ["Line", "Value", "Working / note"])
val = [
    ("eps28", "2028E diluted EPS, US$", "=Financials!F19", USD2, "The market will be pricing 2028 by October 2027."),
    ("mult", "Base multiple, x", "=" + A["pe"], MULT0, "Assumptions tab."),
    ("base", "Base-case value per share, US$", "=B5*B6", USD2, "EPS x multiple."),
    ("px", "Reference price, US$", "=" + A["price"], USD2, "NYSE close, 5 October 2026."),
    ("upb", "Base value against the price", "=B7/B8-1", PCT1, ""),
    ("wtd", "Probability-weighted value, US$", "=Scenarios!N12", USD2, "Scenarios tab."),
    ("rule", "Rule on value alone (15% either side)", '=IF(B9>0.15,"LONG",IF(B9<-0.15,"SHORT","NO CALL"))', None, "Mechanical only. The call below is judgement."),
    ("call", "CALL", CALL["direction"], None, "No call, leaning short: the base is only just outside the band and the stock moves 25% on one print."),
    ("mcap", "Market value on shares outstanding, US$bn", f"={A['price']}*{A['shout']}", BN1, ""),
    ("nc", "Net cash, US$bn", f"={A['cash']}-{A['debt']}", BN3, "30 June 2026; excludes US$0.21bn of financing obligations."),
    ("ev", "Enterprise value, US$bn", "=B13-B14", BN1, ""),
    ("rise", "Change since 31 Dec 2025", f"={A['price']}/{A['pxdec25']}-1", PCT0, ""),
    ("hi", "Price below the highest close", f"=1-{A['price']}/{A['hiclose']}", PCT1, ""),
    ("pe26", "P/E on my 2026E", "=Financials!D21", MULT, ""),
    ("pe27", "P/E on my 2027E", "=Financials!E21", MULT, ""),
    ("pe28", "P/E on my 2028E", "=Financials!F21", MULT, ""),
    ("peg", "P/E on the 2026 EPS guide midpoint", f"={A['price']}/(({A['gepslo']}+{A['gepshi']})/2)", MULT, ""),
    ("pec28", "P/E on the 2028 consensus", f"={A['price']}/{A['cons28']}", MULT, "Nasdaq.com (Zacks), 2 estimates."),
    ("pettm", "P/E on trailing four quarters EPS", f"={A['price']}/Quarterly!B20", MULT, ""),
    (None, "", None, None, ""),
    ("need", "2028 EPS the price needs at the base multiple, US$", f"={A['price']}/{A['pe']}", USD2, "What the market prices in."),
    ("needoi", "2028 operating income that needs, US$bn", f"=B25*Financials!F18/(1-{A['tax28']})-{A['othery']}", BN2, ""),
    ("needrev", "2028 revenue that needs at base margin and costs, US$bn", "=(B26+Scenarios!G6)/Scenarios!E6", BN2, "Against my base US$8.70bn."),
    ("needx", "That against my 2028E base revenue", "=B27/Financials!F5", '0.00"x"', ""),
    ("needgw", "2028 GW that needs (derived)", "=B27*Capacity!B7/Capacity!B6", BN2, "Against 2 GW of planned factory capacity."),
    ("revl", "Price that puts the base case 15% above (revisit as long)", "=B7/1.15", USD2, ""),
    ("revs", "Price that puts the base case 25% below (revisit as short)", "=B7/0.75", USD2, ""),
]
V = {}
for i, (key, label, f, fmt, nt) in enumerate(val):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=label)
    if f is not None:
        c = ws.cell(row=rr, column=2, value=f)
        if fmt: c.number_format = fmt
        if key in ("base", "call"):
            ws.cell(row=rr, column=1).font = F_B; c.font = F_B; c.fill = FILL_KEY
        if key == "call": c.font = F_INB
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
    if key: V[key] = f"Valuation!$B${rr}"

# ===================================================================== SENSITIVITY
ws = sheet("Sensitivity", "VALUE PER SHARE: 2028E EPS AGAINST MULTIPLE (US$)", [24, 14, 14, 14])
head(ws, 4, ["2028 EPS \\ multiple", "", "", ""])
for j, (pe, link) in enumerate(zip(SENS_PE, [False, True, False])):
    c = ws.cell(row=4, column=2 + j, value=("=" + A["pe"]) if link else pe)
    c.number_format = MULT0; c.font = Font(bold=True, color="FFFFFF"); c.fill = FILL_HDR; c.alignment = RIGHT
for i, (e, link) in enumerate(zip(SENS_EPS, [False, True, False])):
    rr = 5 + i
    c = ws.cell(row=rr, column=1, value="=Financials!F19" if link else e)
    c.number_format = USD2; c.font = F_LINK if link else F_IN
    for j in range(3):
        col = get_column_letter(2 + j)
        ws.cell(row=rr, column=2 + j, value=f"=$A{rr}*{col}$4").number_format = USD0
ws["C6"].fill = FILL_KEY
ws["A9"] = "Cells above the reference price"; ws["A9"].font = F_B
ws["B9"] = f'=COUNTIF(B5:D7,">"&{A["price"]})'; ws["C9"] = "of 9"
note(ws, 11, ["Middle row links to the model's 2028E EPS and the middle column to the base multiple."])

# ===================================================================== PEERS
ws = sheet("Peers", "PEER MULTIPLES (FORWARD P/E, stockanalysis.com, 6 OCTOBER 2026)", [26, 16, 14, 80])
head(ws, 4, ["Company", "Ticker", "Forward P/E", "What they make"])
for i, (co, tk, pe, what) in enumerate(PEERS):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=co).font = F_B if co == NAME else Font()
    ws.cell(row=rr, column=2, value=tk)
    c = ws.cell(row=rr, column=3, value=("=" + A["fwdpe"]) if co == NAME else (pe if pe else "n/m"))
    c.number_format = MULT; c.font = F_LINK if co == NAME else F_IN
    ws.cell(row=rr, column=4, value=what).alignment = WRAP
ws["A13"] = "Median, power equipment peers (GE Vernova to Cummins)"; ws["C13"] = "=MEDIAN(C6:C10)"; ws["C13"].number_format = MULT
note(ws, 15, ["Siemens Energy on its Frankfurt listing. FuelCell Energy is loss-making. Multiples only, to set the base multiple.",
              "No view on the peers' shares is expressed."])

# ===================================================================== SOURCES
ws = sheet("Sources", "SOURCES", [6, 130])
head(ws, 4, ["No.", "Source and what it supports"])
srcs = [
    "Bloom results releases (Exhibit 99.1 to Form 8-K): Q2 2024, Q2 2025, Q3 2025, Q4 2025 (5 February 2026) and Q2 2026 (28 July 2026), with their comparative quarters: revenue and cost by line, GAAP and non-GAAP margins, operating income, EPS, related-party revenue, guidance, backlog.",
    "Bloom Form 10-K for 2025: backlog definitions; service contracts of 5 to 20 years, terminable for convenience each year; 1 GW to 2 GW capacity plan; grid queues and the SPP 90-day study route approved by FERC in January 2026.",
    "Bloom Form 10-Q for Q2 2026: debt table and note terms, Green note conversions, Oracle warrant (3,531,073 at US$113.28; exercised 1 May 2026), EPS reconciliation, Brookfield joint ventures, customer concentration, the 8 July 2026 short report.",
    "Bloom announcements: Oracle, up to 2.8 GW with 1.2 GW contracted (13 April 2026); Brookfield framework to US$25bn (30 June 2026); Power Connect (19 August 2026); Q3 2025 results date notice (9 October 2025).",
    "AEP exercise of most of its 900 MW option, about US$2.65bn, as reported on 8 January 2026.",
    "Yahoo Finance chart API, retrieved 6 October 2026: month-end and daily closes, 52-week range.",
    "stockanalysis.com, retrieved 6 October 2026: 2026 consensus EPS and revenue, average target, forward P/E for Bloom and peers. Nasdaq.com (Zacks): 2027 and 2028 consensus EPS, Q3 2026 consensus.",
    "The Physical Layer, 'Bloom Energy is not selling electricity. It is selling time.', 24 September 2026: https://thephysicallayer.fyi/journal/the-product-is-time/",
    "Estimates for 2026 to 2028 are the author's own. Personal research, not investment advice. I hold no position in Bloom Energy.",
]
for i, s in enumerate(srcs):
    ws.cell(row=5 + i, column=1, value=i + 1)
    ws.cell(row=5 + i, column=2, value=s).alignment = WRAP

# ===================================================================== COVER
ws = wb["Cover"]
head(ws, 4, ["Item", "Value", "Note"])
cover = [
    ("Call", "=" + V["call"], None, "Tommy Lau's call, published 6 Oct 2026."),
    ("Conviction", CALL["conviction"], None, "Judgement, not formula."),
    ("Reference price, US$", "=" + A["price"], USD2, "NYSE close, 5 October 2026."),
    ("Base value, US$", "=" + V["base"], USD2, "40x 2028E EPS."),
    ("Base value against the price", "=" + V["upb"], PCT1, ""),
    ("Probability-weighted value, US$", "=" + V["wtd"], USD2, "25 / 50 / 25 bear / base / bull."),
    ("2026E / 2027E / 2028E EPS, US$", '=TEXT(Financials!D19,"0.00")&" / "&TEXT(Financials!E19,"0.00")&" / "&TEXT(Financials!F19,"0.00")', None, "Own estimates, non-GAAP, fully diluted."),
    ("Horizon", "12 months, to October 2027", None, ""),
    ("Wrong if", "See note", None, "Bloom's 2027 revenue of US$7bn or more with non-GAAP gross margin of 36% or more; or the shares close above US$360 or below US$200 by October 2027."),
    ("Revisit if", "See note", None, "A price above about US$322 (consider short) or below about US$210 (consider long); a 2027 guide below 40% growth; or product gross margin of 38% or more in two quarters running."),
]
for i, (lab, v, fmt, nt) in enumerate(cover):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=lab).font = F_B
    c = ws.cell(row=rr, column=2, value=v)
    if fmt: c.number_format = fmt
    c.alignment = RIGHT
    c.font = F_LINK if isinstance(v, str) and v.startswith("=") else F_IN
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
ws["A17"] = "Tabs"; ws["A17"].font = Font(bold=True, color=NAVY)
for i, (t, d) in enumerate([
    ("Assumptions", "Every input, with its source. Change the yellow cells to move the value."),
    ("Dilution", "Shares outstanding, convertible notes, staff awards, the Oracle warrant; the fully diluted count."),
    ("Quarterly", "Q1 2024 to Q2 2026 reported, my Q3 and Q4 2026; product and service margins, related-party sales."),
    ("Financials", "2024A to 2028E: revenue, margins, operating income, net profit, EPS, P/E, megawatt-equivalents."),
    ("Scenarios", "Bear / base / bull and the probability-weighted value."),
    ("Capacity", "Revenue converted to gigawatts at AEP's price, against the factory."),
    ("Valuation", "Base value, market value, multiples, what the price needs."),
    ("Sensitivity", "2028E EPS against multiple."),
    ("Peers", "Forward multiples used to set the base multiple."),
    ("Sources", "Every source used."),
]):
    ws.cell(row=18 + i, column=1, value=t); ws.cell(row=18 + i, column=2, value=d)
ws["A29"] = "Colour key: blue = hardcoded input; yellow fill = my own assumption; green = link to another tab; black = formula."
ws["A29"].font = F_NOTE
ws["A30"] = f"Companion note: {DATE}_{COMPANY}_Initiation.pdf. Personal research, not investment advice."; ws["A30"].font = F_NOTE

order = ["Cover", "Assumptions", "Dilution", "Quarterly", "Financials", "Scenarios", "Capacity", "Valuation", "Sensitivity", "Peers", "Sources"]
wb._sheets = [wb[n] for n in order]
for w in wb.worksheets:
    w.freeze_panes = "A5"
wb.active = 0
wb.save(XLSX)
print("model written", XLSX)
