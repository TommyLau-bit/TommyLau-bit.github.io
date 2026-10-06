"""Lumentum initiation model, 6 Oct 2026. Live formulas throughout.

House style follows the Vertiv, TSMC, Texas Instruments and Bloom Energy models: navy header rows, blue font for
hardcoded inputs, yellow fill on my own assumptions, green font for cross-sheet links, black for formulas.
Non-GAAP basis unless marked GAAP. Fiscal years end late June (FY2026 = year to 27 June 2026).
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from data import *

NAVY = "1F3864"; MUTED = "595959"
F_TTL = Font(bold=True, color=NAVY, size=12)
F_SUB = Font(color=MUTED, size=9)
F_HDR = Font(bold=True, color="FFFFFF", size=10)
F_IN = Font(color="0000FF")
F_INB = Font(color="0000FF", bold=True)
F_LINK = Font(color="008000")
F_B = Font(bold=True)
F_NOTE = Font(color=MUTED, size=9, italic=True)
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_ASSUME = PatternFill("solid", fgColor="FFFF00")
FILL_KEY = PatternFill("solid", fgColor="D9E2F3")
WRAP = Alignment(wrap_text=True, vertical="top")
RIGHT = Alignment(horizontal="right")

USD2 = '"US$"#,##0.00'; USD0 = '"US$"#,##0'; BN2 = '#,##0.00'; BN3 = '#,##0.000'; BN4 = '#,##0.0000'; BN1 = '#,##0.0'
PCT1 = '0.0%'; PCT0 = '0%'; MULT = '0.0"x"'; MULT0 = '0"x"'
SUB = f"Tommy Lau | Lumentum (Nasdaq: LITE) | {DATE_LONG} | Non-GAAP basis; fiscal years to late June | Personal research, not investment advice."

wb = openpyxl.Workbook()


def sheet(name, title, widths):
    ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet(name)
    ws.title = name
    ws["A1"] = title; ws["A1"].font = F_TTL
    ws["A2"] = SUB; ws["A2"].font = F_SUB
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def head(ws, row, cells, start=1):
    for i, c in enumerate(cells, start):
        x = ws.cell(row=row, column=i, value=c)
        x.font = F_HDR; x.fill = FILL_HDR
        x.alignment = Alignment(horizontal="left" if i == start else "right", vertical="center", wrap_text=True)


def note(ws, row, texts):
    for i, t in enumerate(texts):
        ws.cell(row=row + i, column=1, value=t).font = F_NOTE


sheet("Cover", "LUMENTUM (NASDAQ: LITE)  INITIATION MODEL", [38, 26, 80])

# ===================================================================== ASSUMPTIONS
ws = sheet("Assumptions", "ASSUMPTIONS AND INPUTS", [58, 14, 10, 92])
ws["A3"] = ("Blue = hardcoded input. Yellow fill = my own assumption (change these to move the value). Green = link to "
            "another tab. Black = formula. US$ billion unless stated; per-share figures in US$; shares in billions.")
ws["A3"].font = F_NOTE
head(ws, 5, ["Input", "Value", "Unit", "Source / note"])
A = {}
MINE = {"q1rev", "q2rev", "q3rev", "q4rev", "q1om", "q2om", "q3om", "q4om", "otherq", "othery", "shdil", "pe"}
rows = [
    ("MARKET (retrieved 6 October 2026)", None, None, None, None),
    ("price", "Share price (Nasdaq: LITE)", PRICE, "US$", "Nasdaq close, 5 October 2026. Nasdaq.com historical data (Yahoo Finance refused requests on 6 Oct)."),
    ("hi52", "52-week high (intraday, 5 October 2026)", HI52, "US$", "Nasdaq.com."),
    ("hiclose", "Highest close (5 October 2026)", HI_CLOSE, "US$", "Nasdaq.com."),
    ("lo52", "52-week low (intraday, 10 October 2025)", LO52, "US$", "Nasdaq.com."),
    ("loclose", "Lowest close (10 October 2025)", LO_CLOSE, "US$", "Nasdaq.com."),
    ("pxdec25", "Close, 31 Dec 2025", PRICE_DEC25, "US$", "Nasdaq.com."),
    ("common", "Common shares outstanding, 14 Aug 2026", COMMON_OUT, "bn", "10-K cover: 89.7 million."),
    ("pref", "Series A convertible preferred (Nvidia), converts 1 for 1", PREF, "bn", "8-K of 2 March 2026: 2,876,415 shares."),
    ("prefpx", "Price Nvidia paid per preferred share", PREF_PX, "US$", "8-K of 2 March 2026: US$2.0bn in total."),
    ("cash", "Cash, equivalents and short-term investments, 27 Jun 2026", CASH, "US$bn", "Q4 FY26 release: 2,043.5 + 694.9."),
    ("japan", "Japan term loans (SMBC, Mizuho), 27 Jun 2026", JAPAN_LOANS, "US$bn", "10-K."),
    ("cons27", "FY27 consensus EPS", CONS_EPS_27, "US$", "stockanalysis.com, 23 analysts. Nasdaq.com (Zacks): 19.78, 8 estimates."),
    ("cons27z", "FY27 consensus EPS, Zacks", CONS_EPS_27_Z, "US$", "Nasdaq.com (Zacks), 8 estimates."),
    ("cons28", "FY28 consensus EPS", CONS_EPS_28, "US$", "Nasdaq.com (Zacks), 6 estimates; range 23.24 to 36.80."),
    ("consq1", "Q1 FY27 consensus EPS", CONS_Q1, "US$", "Nasdaq.com (Zacks), 6 estimates; below the company's own guide, so likely stale."),
    ("fwdpe", "Forward P/E shown by stockanalysis.com", FWD_PE_SA, "x", "stockanalysis.com, 6 October 2026, intraday."),
    ("constp", "Average analyst target", CONS_TP, "US$", "stockanalysis.com, 26 analysts; median 1,168, range 820 to 1,400."),
    ("REPORTED (Lumentum releases, Form 10-K, Q4 FY26 presentation)", None, None, None, None),
    ("grevlo", "Q1 FY27 revenue guide, low", GUIDE_Q1["rev"][0], "US$bn", "Release of 11 August 2026."),
    ("grevhi", "Q1 FY27 revenue guide, high", GUIDE_Q1["rev"][1], "US$bn", "Same."),
    ("gomlo", "Q1 FY27 non-GAAP operating margin guide, low", GUIDE_Q1["om"][0], "%", "Same."),
    ("gomhi", "Q1 FY27 non-GAAP operating margin guide, high", GUIDE_Q1["om"][1], "%", "Same."),
    ("gepslo", "Q1 FY27 non-GAAP EPS guide, low", GUIDE_Q1["eps"][0], "US$", "Same."),
    ("gepshi", "Q1 FY27 non-GAAP EPS guide, high", GUIDE_Q1["eps"][1], "US$", "Same."),
    ("shguide", "Q1 FY27 diluted share guide", SH_GUIDE, "bn", "Q4 FY26 presentation: 102.0 million."),
    ("tax", "Non-GAAP tax rate", TAX_GUIDE, "%", "Q4 FY26 presentation and release: 16.5%."),
    ("sbc26", "Stock-based pay and related payroll tax, FY26 (excluded from non-GAAP)", SBC_26, "US$bn", "Q4 FY26 reconciliation."),
    ("capex26", "Capital spending, FY26", CAPEX_26, "US$bn", "10-K cash flow."),
    ("ocf26", "Operating cash flow, FY26", OCF_26, "US$bn", "10-K cash flow."),
    ("fy23rev", "Revenue, FY23", FY23_REV, "US$bn", "Q4 FY24 release."),
    ("MY ASSUMPTIONS", None, None, None, None),
    ("q1rev", "Revenue, Q1 FY27 (to Sep 2026)", Q27_REV[0], "US$bn", "MINE. Guide 1.225 to 1.275."),
    ("q2rev", "Revenue, Q2 FY27 (to Dec 2026)", Q27_REV[1], "US$bn", "MINE."),
    ("q3rev", "Revenue, Q3 FY27 (to Mar 2027)", Q27_REV[2], "US$bn", "MINE."),
    ("q4rev", "Revenue, Q4 FY27 (to Jun 2027)", Q27_REV[3], "US$bn", "MINE."),
    ("q1om", "Non-GAAP operating margin, Q1 FY27", Q27_OM[0], "%", "MINE. Guide 39.5 to 40.5%."),
    ("q2om", "Non-GAAP operating margin, Q2 FY27", Q27_OM[1], "%", "MINE."),
    ("q3om", "Non-GAAP operating margin, Q3 FY27", Q27_OM[2], "%", "MINE."),
    ("q4om", "Non-GAAP operating margin, Q4 FY27", Q27_OM[3], "%", "MINE."),
    ("otherq", "Non-GAAP other income per quarter, FY27", OTHER_Q, "US$bn", "MINE. Q4 FY26: 0.022; cash falls as notes convert."),
    ("othery", "Non-GAAP other income, FY28", OTHER_Y, "US$bn", "MINE."),
    ("shdil", "Net new diluted shares, FY28 on FY27", SH_DIL, "%", "MINE. Awards and note conversions."),
    ("pe", "Target multiple on FY28E EPS (base)", BASE_PE, "x", "MINE. Median forward P/E of five optical peers (Peers tab)."),
]
r = 6
for key, label, val, unit, src in rows:
    if label is None:
        ws.cell(row=r, column=1, value=key).font = Font(bold=True, color=NAVY); r += 1; continue
    ws.cell(row=r, column=1, value=label)
    c = ws.cell(row=r, column=2, value=val)
    mine = key in MINE
    c.font = F_INB if mine else F_IN
    if mine: c.fill = FILL_ASSUME
    c.number_format = PCT1 if unit == "%" else (MULT0 if unit == "x" else (BN4 if unit == "bn" else ("#,##0.00" if unit == "US$" else BN3)))
    ws.cell(row=r, column=3, value=unit)
    ws.cell(row=r, column=4, value=src).alignment = WRAP
    A[key] = f"Assumptions!$B${r}"
    r += 1

# ===================================================================== SHARES
ws = sheet("Shares", "SHARE COUNT AND BALANCE SHEET", [62, 14, 14, 16, 60])
head(ws, 4, ["Convertible notes, 27 Jun 2026", "Principal, US$bn", "Conversion price, US$", "Net shares at price, bn", "Working / note"])
for i, (nm, p, cp) in enumerate(NOTES):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=nm)
    c = ws.cell(row=rr, column=2, value=p); c.font = F_IN; c.number_format = BN4
    c = ws.cell(row=rr, column=3, value=cp); c.font = F_IN; c.number_format = USD2
    ws.cell(row=rr, column=4, value=f"=MAX(0,B{rr}/C{rr}*(1-C{rr}/{A['price']}))").number_format = BN4
    ws.cell(row=rr, column=5, value="Principal paid in cash; value above the conversion price in shares (10-K).")
ws["A9"] = "Total"; ws["A9"].font = F_B
ws["B9"] = "=SUM(B5:B8)"; ws["B9"].number_format = BN4
ws["D9"] = "=SUM(D5:D8)"; ws["D9"].number_format = BN4
lines = [
    (11, "2032 capped call cap price, US$", CAP_PRICE, USD2, "10-K. Offsets 2032 note dilution up to this price."),
    (12, "Shares the capped call offsets, bn", "=B5/C5*(1-C5/B11)", BN4, "At or above the cap."),
    (13, "Common shares outstanding, 14 Aug 2026, bn", "=" + A["common"], BN4, ""),
    (14, "Preferred shares (Nvidia), bn", "=" + A["pref"], BN4, ""),
    (15, "Staff awards in Q4 FY26 diluted count, bn", AWARDS, BN4, "Q4 FY26 release, note 8: about 3.1 million (treasury method)."),
    (16, "My diluted count at the price, bn", "=B13+B14+D9-B12+B15", BN4, "Against the company's 102.0m guide for Q1 FY27."),
    (17, "Company Q1 FY27 diluted share guide, bn", "=" + A["shguide"], BN4, "Used for FY27E."),
    (19, "Debt: notes principal plus Japan loans, US$bn", f"=B9+{A['japan']}", BN3, ""),
    (20, "Net cash, US$bn", f"={A['cash']}-B19", BN3, "27 June 2026."),
    (21, "Early conversion requests by 14 Aug 2026, principal, US$bn", EARLY_CONV, BN3, "10-K. Principal settles in cash."),
    (22, "Nvidia preferred stake at the price, US$bn", f"=B14*{A['price']}", BN3, ""),
    (23, "Nvidia stake against the price it paid", f"={A['price']}/{A['prefpx']}-1", PCT1, ""),
    (24, "Preferred as a share of common plus preferred", "=B14/(B13+B14)", PCT1, ""),
    (25, "Stock-based pay against FY26 non-GAAP operating income", f"={A['sbc26']}/Financials!D9", PCT1, ""),
]
for rr, label, f, fmt, nt in lines:
    ws.cell(row=rr, column=1, value=label)
    c = ws.cell(row=rr, column=2, value=f); c.number_format = fmt
    if not (isinstance(f, str) and f.startswith("=")): c.font = F_IN
    ws.cell(row=rr, column=5, value=nt).alignment = WRAP

# ===================================================================== QUARTERLY
ws = sheet("Quarterly", "QUARTERLY HISTORY, Q1 FY25 TO Q4 FY26, AND MY FY27 BY QUARTER", [50] + [10] * 12)
QH = Q + ["Q1 27E", "Q2 27E", "Q3 27E", "Q4 27E"]
head(ws, 4, ["US$ billion unless stated"] + QH)
C = [get_column_letter(2 + j) for j in range(12)]


def qrow(row, label, vals, fmt, bold=False):
    ws.cell(row=row, column=1, value=label).font = F_B if bold else Font()
    for j, v in enumerate(vals):
        c = ws.cell(row=row, column=2 + j, value=v if v is not None else "")
        c.number_format = fmt; c.alignment = RIGHT
        if v is not None and not (isinstance(v, str) and v.startswith("=")):
            c.font = F_IN
        elif isinstance(v, str) and v.startswith("=Assumptions"):
            c.font = F_LINK


bn = lambda xs: [x / 1000 for x in xs]
pct = lambda xs: [x / 100 for x in xs]
E4 = [None] * 4
qrow(5, "Revenue", bn(Q_REV) + ["=" + A[k] for k in ("q1rev", "q2rev", "q3rev", "q4rev")], BN3, True)
qrow(6, "Components", bn(Q_COMP) + E4, BN3)
qrow(7, "Systems", bn(Q_SYS) + E4, BN3)
qrow(8, "Revenue growth on a year earlier", [None] * 4 + [f"={C[j]}5/{C[j - 4]}5-1" for j in range(4, 12)], PCT1)
qrow(9, "Gross margin, GAAP", pct(Q_GM_GAAP) + E4, PCT1, True)
qrow(10, "Gross margin, non-GAAP", pct(Q_GM) + E4, PCT1)
qrow(11, "Operating margin, non-GAAP (FY27 mine)", pct(Q_OM) + ["=" + A[k] for k in ("q1om", "q2om", "q3om", "q4om")], PCT1, True)
qrow(12, "Operating income, non-GAAP", [f"={c}5*{c}11" for c in C], BN3)
qrow(13, "Net profit, non-GAAP (FY27 mine)", [None] * 8 + [f"=({c}12+{A['otherq']})*(1-{A['tax']})" for c in C[8:]], BN3)
qrow(14, "Diluted shares, bn (FY27: company guide)", [x / 1000 for x in Q_SH] + ["=" + A["shguide"]] * 4, BN4)
qrow(15, "Diluted EPS, US$ (reported; FY27 mine)", Q_EPS + [f"={c}13/{c}14" for c in C[8:]], USD2, True)
ws["A17"] = "Checks"; ws["A17"].font = Font(bold=True, color=NAVY)
chk = [
    (18, "Trailing four quarters EPS to Q4 FY26", "=SUM(F15:I15)", USD2),
    (19, "Revenue growth, Q4 FY26 on Q4 FY25", "=I5/E5-1", PCT1),
    (20, "Components growth, Q4 FY26", "=I6/E6-1", PCT1),
    (21, "Systems growth, Q4 FY26", "=I7/E7-1", PCT1),
    (22, "GAAP gross margin rise, Q4 FY25 to Q4 FY26, points", "=(I9-E9)*100", BN1),
    (23, "Q1 FY27E EPS against the guide (low / high)", "=J15", USD2),
    (24, "Q1 FY27E operating margin", "=J11", PCT1),
]
for rr, lab, f, fmt in chk:
    ws.cell(row=rr, column=1, value=lab); c = ws.cell(row=rr, column=2, value=f); c.number_format = fmt
ws["C23"] = "=" + A["gepslo"]; ws["D23"] = "=" + A["gepshi"]
for c_ in ("C23", "D23"): ws[c_].number_format = USD2
note(ws, 26, ["Source: Lumentum results releases (Exhibit 99.1 to Form 8-K), 7 November 2024 to 11 August 2026. Non-GAAP figures as recast from FY25.",
              "Historical operating income is revenue x reported non-GAAP operating margin (Q4 FY26 reported: 0.3688).",
              "FY27E quarters are my estimates: EPS = (revenue x operating margin + other income) x (1 - 16.5%) / 102.0m shares."])

# ===================================================================== FINANCIALS
ws = sheet("Financials", "FINANCIAL SUMMARY AND FORECASTS (US$ billion unless stated, non-GAAP basis)", [52, 12, 12, 12, 12, 12])
head(ws, 4, ["", "FY24A", "FY25A", "FY26A", "FY27E*", "FY28E*"])
fin = [
    ("Revenue", [REV_H[0], REV_H[1], REV_H[2], "=SUM(Quarterly!J5:M5)", "=E5*(1+Scenarios!B6)"], BN2, True),
    ("Revenue growth", [f"=B5/{A['fy23rev']}-1", "=C5/B5-1", "=D5/C5-1", "=E5/D5-1", "=F5/E5-1"], PCT1, False),
    ("Gross margin, GAAP", [GM_GAAP_H[0], GM_GAAP_H[1], GM_GAAP_H[2], None, None], PCT1, False),
    ("Gross margin, non-GAAP", [GM_H[0], GM_H[1], GM_H[2], None, None], PCT1, False),
    ("Operating income", [OPINC_H[0], OPINC_H[1], OPINC_H[2], "=SUM(Quarterly!J12:M12)", "=F5*Scenarios!C6"], BN3, True),
    ("Operating margin", ["=B9/B5", "=C9/C5", "=D9/D5", "=E9/E5", "=F9/F5"], PCT1, False),
    ("Other income", [None, None, 0.0399, f"=4*{A['otherq']}", "=" + A["othery"]], BN3, False),
    ("Tax rate", [None, None, 0.165, "=" + A["tax"], "=" + A["tax"]], PCT1, False),
    ("Net profit", [None, None, 0.7823, "=SUM(Quarterly!J13:M13)", "=(F9+F11)*(1-F12)"], BN3, True),
    ("Diluted shares, bn", [None, SH_H[1], SH_H[2], "=" + A["shguide"], f"=E14*(1+{A['shdil']})"], BN4, False),
    ("Diluted EPS, US$", [EPS_H[0], EPS_H[1], EPS_H[2], "=SUM(Quarterly!J15:M15)", "=F13/F14"], USD2, True),
    ("EPS growth", [None, "=C15/B15-1", "=D15/C15-1", "=E15/D15-1", "=F15/E15-1"], PCT1, False),
    ("P/E at reference price", [None, None, f"={A['price']}/D15", f"={A['price']}/E15", f"={A['price']}/F15"], MULT, False),
    ("Consensus EPS, US$", [None, None, None, "=" + A["cons27"], "=" + A["cons28"]], USD2, False),
    ("Mine against consensus", [None, None, None, "=E15/E18-1", "=F15/F18-1"], PCT1, False),
    ("P/E on consensus", [None, None, None, f"={A['price']}/E18", f"={A['price']}/F18"], MULT, False),
]
for i, (label, vals, fmt, bold) in enumerate(fin):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=label).font = F_B if bold else Font()
    for j, v in enumerate(vals):
        c = ws.cell(row=rr, column=2 + j, value=v if v is not None else "")
        c.number_format = fmt; c.alignment = RIGHT
        if v is not None and not isinstance(v, str): c.font = F_IN
        elif isinstance(v, str) and v.startswith("=Assumptions"): c.font = F_LINK
note(ws, 23, [
    "*FY27E and FY28E are my own estimates. FY27E sums the Quarterly tab. FY28E uses the base case on the Scenarios tab (row 6).",
    "FY24A to FY26A as reported (FY24 non-GAAP recast in the Q4 FY25 release). FY26 net profit and EPS exclude a US$7.76bn non-cash loss on debt extinguishment.",
    "FY27 consensus from stockanalysis.com (23 analysts); FY28 from Nasdaq.com (Zacks, 6 estimates).",
    f"Note check: FY27E revenue {REV27:.2f}, EPS {EPS27:.2f}; FY28E revenue {REV28:.2f}, EPS {EPS28:.2f}.",
])

# ===================================================================== SCENARIOS
ws = sheet("Scenarios", "THREE CASES, TWELVE MONTHS OUT (VALUED ON FY28E)", [30] + [12] * 9)
head(ws, 4, ["Case", "Growth FY28", "Op. margin FY28", "FY28 revenue", "FY28 op. income", "FY28 net profit", "FY28 EPS",
             "Multiple", "Value, US$", "vs price"])
R27 = "Financials!$E$5"
for i, s in enumerate(SCEN):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=s[0]).font = F_B
    for j, v in enumerate(s[1:3]):
        c = ws.cell(row=rr, column=2 + j, value=v); c.font = F_INB; c.fill = FILL_ASSUME; c.number_format = PCT1
    ws.cell(row=rr, column=4, value=f"={R27}*(1+B{rr})").number_format = BN2
    ws.cell(row=rr, column=5, value=f"=D{rr}*C{rr}").number_format = BN3
    ws.cell(row=rr, column=6, value=f"=(E{rr}+{A['othery']})*(1-{A['tax']})").number_format = BN3
    ws.cell(row=rr, column=7, value=f"=F{rr}/Financials!$F$14").number_format = USD2
    if s[0] == "Base":
        c = ws.cell(row=rr, column=8, value="=" + A["pe"]); c.font = F_LINK
    else:
        c = ws.cell(row=rr, column=8, value=s[3]); c.font = F_INB; c.fill = FILL_ASSUME
    c.number_format = MULT0
    ws.cell(row=rr, column=9, value=f"=G{rr}*H{rr}").number_format = USD2
    ws.cell(row=rr, column=10, value=f"=I{rr}/{A['price']}-1").number_format = PCT0
ws["A10"] = "Probability"; ws["A10"].font = F_B
for j, s in enumerate(SCEN):
    ws.cell(row=9, column=2 + j, value=s[0]).font = F_B
    c = ws.cell(row=10, column=2 + j, value=s[4]); c.font = F_INB; c.fill = FILL_ASSUME; c.number_format = PCT0
ws["E10"] = "=SUM(B10:D10)"; ws["E10"].number_format = PCT0; ws["F10"] = "must be 100%"; ws["F10"].font = F_NOTE
ws["A12"] = "Probability-weighted value, US$"; ws["A12"].font = F_B
ws["I12"] = "=B10*I5+C10*I6+D10*I7"; ws["I12"].number_format = USD2; ws["I12"].font = F_B; ws["I12"].fill = FILL_KEY
ws["J12"] = f"=I12/{A['price']}-1"; ws["J12"].number_format = PCT1
note(ws, 14, ["All cases start from my FY27E revenue (Financials tab) and use FY28 diluted shares. Operating margin is non-GAAP.",
              "Bear: new capacity (Japan, Greensboro, Coherent's six-inch wafers) meets an inventory pause, as in FY24. Bull: the shortage lasts through calendar 2028.",
              "The base row (6) drives the Financials tab for FY28E."])

# ===================================================================== VALUATION
ws = sheet("Valuation", "VALUATION: 35x FY28E EPS, TWELVE MONTHS OUT", [64, 16, 70])
head(ws, 4, ["Line", "Value", "Working / note"])
val = [
    ("eps28", "FY28E diluted EPS, US$", "=Financials!F15", USD2, "By October 2027 the market will be pricing FY28 (year to June 2028)."),
    ("mult", "Base multiple, x", "=" + A["pe"], MULT0, "Assumptions tab."),
    ("base", "Base-case value per share, US$", "=B5*B6", USD2, "EPS x multiple."),
    ("px", "Reference price, US$", "=" + A["price"], USD2, "Nasdaq close, 5 October 2026."),
    ("upb", "Base value against the price", "=B7/B8-1", PCT1, ""),
    ("wtd", "Probability-weighted value, US$", "=Scenarios!I12", USD2, "Scenarios tab."),
    ("rule", "Rule on value alone (long at 15% above, short at 25% below)", '=IF(B9>0.15,"LONG",IF(B9<-0.25,"SHORT","NO CALL"))', None, "Mechanical only; wider on the short side because a short's losses are open-ended. The call below is judgement."),
    ("call", "DRAFT VIEW" if CALL["draft"] else "CALL", CALL["direction"], None, "No call, leaning short. Tommy Lau's call, published 7 Oct 2026."),
    ("mcap", "Market value on common shares, US$bn", f"={A['price']}*{A['common']}", BN1, ""),
    ("mcapfd", "Market value with the preferred, US$bn", f"={A['price']}*({A['common']}+{A['pref']})", BN1, ""),
    ("nc", "Net cash, US$bn", "=Shares!B20", BN3, "27 June 2026."),
    ("ev", "Enterprise value, US$bn", "=B14-B15", BN1, ""),
    ("riselo", "Change since the lowest close (10 Oct 2025)", f"={A['price']}/{A['loclose']}-1", PCT0, ""),
    ("rise", "Change since 31 Dec 2025", f"={A['price']}/{A['pxdec25']}-1", PCT0, ""),
    ("pe27", "P/E on my FY27E", "=Financials!E17", MULT, ""),
    ("pe28", "P/E on my FY28E", "=Financials!F17", MULT, ""),
    ("pec27", "P/E on FY27 consensus", "=Financials!E20", MULT, "stockanalysis.com."),
    ("pec28", "P/E on FY28 consensus", "=Financials!F20", MULT, "Nasdaq.com (Zacks)."),
    ("pettm", "P/E on trailing four quarters EPS", f"={A['price']}/Quarterly!B18", MULT, "Sum of quarterly EPS (US$8.37) differs from FY26 annual EPS (US$8.67) because share counts differ by quarter. The note uses the annual P/E (Financials!D17)."),
    ("perr", "P/E on Q1 FY27 guided EPS midpoint x 4", f"={A['price']}/(4*({A['gepslo']}+{A['gepshi']})/2)", MULT, ""),
    (None, "", None, None, ""),
    ("need", "FY28 EPS the price needs at the base multiple, US$", f"={A['price']}/{A['pe']}", USD2, "What the market prices in."),
    ("needoi", "FY28 operating income that needs, US$bn", f"=B26*Financials!F14/(1-{A['tax']})-{A['othery']}", BN3, ""),
    ("needrev", "FY28 revenue that needs at the base margin, US$bn", "=B27/Scenarios!C6", BN2, "Against my base FY28E."),
    ("needg", "That as growth on my FY27E", "=B28/Financials!E5-1", PCT1, ""),
    ("needx", "That against FY26 revenue (times)", "=B28/Financials!D5", '0.00"x"', ""),
    ("consrev", "FY28 revenue the Zacks consensus EPS needs at the base margin, US$bn", f"=({A['cons28']}*Financials!F14/(1-{A['tax']})-{A['othery']})/Scenarios!C6", BN2, ""),
    ("revl", "Price that puts the base case 15% above (revisit as long)", "=B7/1.15", USD2, ""),
    ("revs", "Price that puts the base case 25% below (revisit as short)", "=B7/0.75", USD2, ""),
]
V = {}
for i, (key, label, f, fmt, nt) in enumerate(val):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=label)
    if f is not None:
        c = ws.cell(row=rr, column=2, value=f)
        if fmt: c.number_format = fmt
        if key in ("base", "call"):
            ws.cell(row=rr, column=1).font = F_B; c.font = F_B; c.fill = FILL_KEY
        if key == "call": c.font = F_INB
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
    if key: V[key] = f"Valuation!$B${rr}"

# ===================================================================== SENSITIVITY
ws = sheet("Sensitivity", "VALUE PER SHARE: FY28E EPS AGAINST MULTIPLE (US$)", [24, 14, 14, 14])
head(ws, 4, ["FY28 EPS \\ multiple", "", "", ""])
for j, (pe, link) in enumerate(zip(SENS_PE, [False, True, False])):
    c = ws.cell(row=4, column=2 + j, value=("=" + A["pe"]) if link else pe)
    c.number_format = MULT0; c.font = Font(bold=True, color="FFFFFF"); c.fill = FILL_HDR; c.alignment = RIGHT
for i, (e, link) in enumerate(zip(SENS_EPS, [False, True, False])):
    rr = 5 + i
    c = ws.cell(row=rr, column=1, value="=Financials!F15" if link else e)
    c.number_format = USD2; c.font = F_LINK if link else F_IN
    for j in range(3):
        col = get_column_letter(2 + j)
        ws.cell(row=rr, column=2 + j, value=f"=$A{rr}*{col}$4").number_format = USD0
ws["C6"].fill = FILL_KEY
ws["A9"] = "Cells above the reference price"; ws["A9"].font = F_B
ws["B9"] = f'=COUNTIF(B5:D7,">"&{A["price"]})'; ws["C9"] = "of 9"
note(ws, 11, ["Middle row links to the model's FY28E EPS and the middle column to the base multiple."])

# ===================================================================== PEERS
ws = sheet("Peers", "PEER MULTIPLES (FORWARD P/E, stockanalysis.com, 6 OCTOBER 2026)", [26, 16, 14, 80])
head(ws, 4, ["Company", "Ticker", "Forward P/E", "What they make"])
for i, (co, tk, pe, what) in enumerate(PEERS):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=co).font = F_B if co == NAME else Font()
    ws.cell(row=rr, column=2, value=tk)
    c = ws.cell(row=rr, column=3, value=("=" + A["fwdpe"]) if co == NAME else pe)
    c.number_format = MULT; c.font = F_LINK if co == NAME else F_IN
    ws.cell(row=rr, column=4, value=what).alignment = WRAP
ws["A12"] = "Median, five optical peers (excluding Lumentum)"; ws["C12"] = "=MEDIAN(C6:C10)"; ws["C12"].number_format = MULT
note(ws, 14, ["Forward P/E as shown on stockanalysis.com during US trading on 6 October 2026. Multiples only, to set the base multiple.",
              "No view on the peers' shares is expressed."])

# ===================================================================== SOURCES
ws = sheet("Sources", "SOURCES", [6, 130])
head(ws, 4, ["No.", "Source and what it supports"])
srcs = [
    "Lumentum results releases (Exhibit 99.1 to Form 8-K): Q4 FY24 (14 August 2024), Q1 FY25 (7 November 2024) to Q4 FY26 (11 August 2026), with comparatives: revenue, components and systems, GAAP and non-GAAP margins, EPS, diluted shares, guidance.",
    "Lumentum Form 10-K for the year to 27 June 2026 (filed 17 August 2026): shares outstanding (89.7m, 14 August 2026), preferred stock, convertible notes and conversion prices, capped call, early conversion requests, Japan loans, customer concentration, capital spending, cash flow, Greensboro (17 March 2026, US$38.0m), Sagamihara.",
    "Lumentum Form 8-K of 2 March 2026 (Items 3.02, 5.03, 7.01) and joint release with Nvidia: 2,876,415 Series A convertible preferred shares at US$695.31; non-exclusive multibillion purchase commitment and capacity rights.",
    "Coherent Form 8-K, Exhibit 99.1, 2 March 2026: Nvidia's US$2bn investment in Coherent with a purchase commitment and capacity rights. Coherent release of 25 March 2024 on six-inch indium phosphide wafers.",
    "Lumentum Q4 FY26 earnings presentation and call, 11 August 2026: 200G EMLs above 25% of EML revenue; Q1 FY27 guide incl. 102.0m shares and 16.5% tax; supply behind demand; Japanese fabs; Greensboro timing; repricing.",
    "Lumentum investor relations event listing: Q1 FY27 results call, 5 November 2026, 5pm ET.",
    "Nasdaq.com historical prices, retrieved 6 October 2026: daily and month-end closes, 52-week range. (Yahoo Finance chart API refused requests on 6 October 2026.)",
    "stockanalysis.com, retrieved 6 October 2026: FY27 consensus EPS and revenue, average target, forward P/E for Lumentum and peers. Nasdaq.com (Zacks): FY27, FY28 and Q1 FY27 consensus EPS.",
    "The Physical Layer, 'Nvidia paid Lumentum $2 billion for lasers, because silicon cannot make light', 5 October 2026: https://thephysicallayer.fyi/journal/lumentum-makes-the-light/",
    "Estimates for FY27 and FY28 are the author's own. Personal research, not investment advice. I hold no position in Lumentum.",
]
for i, s in enumerate(srcs):
    ws.cell(row=5 + i, column=1, value=i + 1)
    ws.cell(row=5 + i, column=2, value=s).alignment = WRAP

# ===================================================================== COVER
ws = wb["Cover"]
head(ws, 4, ["Item", "Value", "Note"])
cover = [
    ("Draft view" if CALL["draft"] else "Call", "=" + V["call"], None, "Draft for Tommy Lau's decision; not yet a call." if CALL["draft"] else "Tommy Lau's call."),
    ("Conviction", CALL["conviction"], None, "Judgement, not formula."),
    ("Reference price, US$", "=" + A["price"], USD2, "Nasdaq close, 5 October 2026."),
    ("Base value, US$", "=" + V["base"], USD2, "35x FY28E EPS."),
    ("Base value against the price", "=" + V["upb"], PCT1, ""),
    ("Probability-weighted value, US$", "=" + V["wtd"], USD2, "25 / 50 / 25 bear / base / bull."),
    ("FY27E / FY28E EPS, US$", '=TEXT(Financials!E15,"0.00")&" / "&TEXT(Financials!F15,"0.00")', None, "Own estimates, non-GAAP, diluted."),
    ("Horizon", "12 months, to October 2027", None, ""),
    ("Wrong if", "See note", None, "FY27 revenue of US$6.8bn or more with GAAP gross margin of 50% or more in its fourth quarter; or the shares close above US$1,450 or below US$750 by October 2027."),
    ("Revisit if", "See note", None, "A price below about US$860 (consider long) or above about US$1,320 (consider short); GAAP gross margin below 44% in any quarter while revenue grows; or December 2026 quarter revenue above US$1.6bn with GAAP gross margin above 50%."),
]
for i, (lab, v, fmt, nt) in enumerate(cover):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=lab).font = F_B
    c = ws.cell(row=rr, column=2, value=v)
    if fmt: c.number_format = fmt
    c.alignment = RIGHT
    c.font = F_LINK if isinstance(v, str) and v.startswith("=") else F_IN
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
ws["A17"] = "Tabs"; ws["A17"].font = Font(bold=True, color=NAVY)
for i, (t, d) in enumerate([
    ("Assumptions", "Every input, with its source. Change the yellow cells to move the value."),
    ("Shares", "Common, preferred, convertible notes, capped call, awards; net cash; Nvidia's stake."),
    ("Quarterly", "Q1 FY25 to Q4 FY26 reported, my FY27 by quarter."),
    ("Financials", "FY24A to FY28E: revenue, margins, operating income, net profit, EPS, P/E, consensus."),
    ("Scenarios", "Bear / base / bull and the probability-weighted value."),
    ("Valuation", "Base value, market value, multiples, what the price needs."),
    ("Sensitivity", "FY28E EPS against multiple."),
    ("Peers", "Forward multiples used to set the base multiple."),
    ("Sources", "Every source used."),
]):
    ws.cell(row=18 + i, column=1, value=t); ws.cell(row=18 + i, column=2, value=d)
ws["A28"] = "Colour key: blue = hardcoded input; yellow fill = my own assumption; green = link to another tab; black = formula."
ws["A28"].font = F_NOTE
ws["A29"] = f"Companion note: {DATE}_{COMPANY}_Initiation.pdf. Personal research, not investment advice."; ws["A29"].font = F_NOTE

order = ["Cover", "Assumptions", "Shares", "Quarterly", "Financials", "Scenarios", "Valuation", "Sensitivity", "Peers", "Sources"]
wb._sheets = [wb[n] for n in order]
for w in wb.worksheets:
    w.freeze_panes = "A5"
wb.active = 0
wb.save(XLSX)
print("model written", XLSX)
