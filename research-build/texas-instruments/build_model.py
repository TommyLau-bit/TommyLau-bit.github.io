"""Texas Instruments initiation model, 6 Oct 2026. Live formulas throughout.

House style follows the Vertiv and TSMC models: navy header rows, blue font for hardcoded inputs, yellow fill on my
own assumptions, green font for cross-sheet links, black for formulas.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from data import *

NAVY = "1F3864"; MUTED = "595959"
F_TTL = Font(bold=True, color=NAVY, size=12)
F_SUB = Font(color=MUTED, size=9)
F_HDR = Font(bold=True, color="FFFFFF", size=10)
F_IN = Font(color="0000FF")
F_INB = Font(color="0000FF", bold=True)
F_LINK = Font(color="008000")
F_B = Font(bold=True)
F_NOTE = Font(color=MUTED, size=9, italic=True)
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_ASSUME = PatternFill("solid", fgColor="FFFF00")
FILL_KEY = PatternFill("solid", fgColor="D9E2F3")
WRAP = Alignment(wrap_text=True, vertical="top")
RIGHT = Alignment(horizontal="right")

USD2 = '"US$"#,##0.00'; USD0 = '"US$"#,##0'; BN2 = '#,##0.00'; BN3 = '#,##0.000'; BN1 = '#,##0.0'
PCT1 = '0.0%'; PCT0 = '0%'; MULT = '0.0"x"'; MULT0 = '0"x"'
SUB = f"Tommy Lau | Texas Instruments (Nasdaq: TXN) | {DATE_LONG} | Personal research, not investment advice."

wb = openpyxl.Workbook()


def sheet(name, title, widths):
    ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet(name)
    ws.title = name
    ws["A1"] = title; ws["A1"].font = F_TTL
    ws["A2"] = SUB; ws["A2"].font = F_SUB
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def head(ws, row, cells, start=1):
    for i, c in enumerate(cells, start):
        x = ws.cell(row=row, column=i, value=c)
        x.font = F_HDR; x.fill = FILL_HDR
        x.alignment = Alignment(horizontal="left" if i == start else "right", vertical="center", wrap_text=True)


def note(ws, row, texts):
    for i, t in enumerate(texts):
        ws.cell(row=row + i, column=1, value=t).font = F_NOTE


sheet("Cover", "TEXAS INSTRUMENTS (Nasdaq: TXN)  INITIATION MODEL", [36, 24, 80])

# ===================================================================== ASSUMPTIONS
ws = sheet("Assumptions", "ASSUMPTIONS AND INPUTS", [50, 14, 10, 92])
ws["A3"] = ("Blue = hardcoded input. Yellow fill = my own assumption (change these to move the value). Green = link to "
            "another tab. Black = formula. US$ billion unless stated; per-share figures in US$.")
ws["A3"].font = F_NOTE
head(ws, 5, ["Input", "Value", "Unit", "Source / note"])
A = {}
MINE = {"nonopq", "tax", "q3rev", "q3om", "q4rev", "q4om", "dcq3e", "dcq4e", "fall", "ddep27", "ddep28",
        "dcg27", "dcg28", "reg27", "reg28", "pe"}
rows = [
    ("MARKET (retrieved 6 October 2026)", None, None, None, None),
    ("price", "Share price (Nasdaq: TXN)", PRICE, "US$", "Nasdaq close, 5 October 2026. Yahoo Finance."),
    ("hi52", "52-week high", HI52, "US$", "Yahoo Finance."),
    ("px25", "Month-end close, 30 Sep 2025", PRICE_SEP25, "US$", "Yahoo Finance."),
    ("diluted", "Diluted shares", DILUTED, "bn", "Q2 2026 release: 920 million average diluted shares."),
    ("cash", "Cash and short-term investments, 30 Jun 2026", CASH, "US$bn", "Q2 2026 release: 3.660 + 3.341."),
    ("debt", "Total debt, 30 Jun 2026", DEBT, "US$bn", "Q2 2026 release: long-term 12.903 + current portion 1.149."),
    ("dps", "Quarterly dividend from November 2026", DPS_NEW, "US$", "TI 8-K of 17 September 2026: raised 7% from US$1.42."),
    ("cons26", "2026 consensus EPS", CONS_EPS_26, "US$", "stockanalysis.com, retrieved 6 October 2026."),
    ("cons27", "2027 consensus EPS", CONS_EPS_27, "US$", "stockanalysis.com, retrieved 6 October 2026."),
    ("consrev26", "2026 consensus revenue", CONS_REV_26, "US$bn", "stockanalysis.com, retrieved 6 October 2026."),
    ("fwdpe", "Forward P/E shown by stockanalysis.com", FWD_PE_SA, "x", "stockanalysis.com, 6 October 2026."),
    ("constp", "Average analyst target", CONS_TP, "US$", "stockanalysis.com, 36 analysts, 6 October 2026."),
    ("REPORTED (TI releases, Form 10-K, calls)", None, None, None, None),
    ("dc25", "Data centre revenue, 2025", DC_25, "US$bn", "TI call, 27 January 2026: US$1.5bn, up 64%, 9% of revenue."),
    ("dcq4", "Data centre revenue, Q4 2025", DC_Q4_25, "US$bn", "TI call, 27 January 2026: left 2025 at about US$450m a quarter."),
    ("dcq4qq", "Data centre, Q4 2025 growth on Q3", DC_Q4_QQ, "%", "TI call, 27 January 2026: mid-single digits sequentially."),
    ("dcq1yy", "Data centre, Q1 2026 growth on a year earlier", DC_Q1_YY, "%", "TI call, 22 April 2026, as reported: about 90%."),
    ("dcq2yy", "Data centre, Q2 2026 growth on a year earlier", DC_Q2_YY, "%", "TI call, 22 July 2026: doubled."),
    ("dcq2qq", "Data centre, Q2 2026 growth on Q1", DC_Q2_QQ, "%", "TI call, 22 July 2026: around 20% sequentially."),
    ("capexlo", "2026 capital spending guide, low", CAPEX_GUIDE_26[0], "US$bn", "2025 Form 10-K and Q2 2026 10-Q."),
    ("capexhi", "2026 capital spending guide, high", CAPEX_GUIDE_26[1], "US$bn", "Same."),
    ("deplo", "2026 depreciation guide, low", DEP_GUIDE_26[0], "US$bn", "TI call, 27 January 2026."),
    ("dephi", "2026 depreciation guide, high", DEP_GUIDE_26[1], "US$bn", "Same."),
    ("g3lo", "Q3 2026 revenue guide, low", G3_REV[0], "US$bn", "Release of 22 July 2026."),
    ("g3hi", "Q3 2026 revenue guide, high", G3_REV[1], "US$bn", "Same."),
    ("g3elo", "Q3 2026 EPS guide, low", G3_EPS[0], "US$", "Same."),
    ("g3ehi", "Q3 2026 EPS guide, high", G3_EPS[1], "US$", "Same."),
    ("MY ASSUMPTIONS", None, None, None, None),
    ("nonopq", "Other income less interest, per quarter", NONOP_Q, "US$bn", "MINE. Q2 2026: other income 0.069 less interest 0.141. Excludes any Silicon Labs debt."),
    ("tax", "Tax rate", TAX, "%", "MINE. TI guides about 13% for Q3 2026 and 13 to 14% for 2026."),
    ("q3rev", "Q3 2026 revenue", Q3E_REV, "US$bn", "MINE. Guide 5.65 to 6.15; loadings rising (Q2 call)."),
    ("q3om", "Q3 2026 operating margin", Q3E_OM, "%", "MINE. Gives EPS at the US$2.40 guide midpoint."),
    ("q4rev", "Q4 2026 revenue", Q4E_REV, "US$bn", "MINE. Fourth quarter seasonally softer (10-K)."),
    ("q4om", "Q4 2026 operating margin", Q4E_OM, "%", "MINE."),
    ("dcq3e", "Data centre revenue, Q3 2026", 0.76, "US$bn", "MINE. About 15% on Q2."),
    ("dcq4e", "Data centre revenue, Q4 2026", 0.80, "US$bn", "MINE."),
    ("fall", "Fall-through before depreciation", FALL, "%", "MINE. TI gives 70 to 85% (Q2 2026 call); I use 75%, in both directions."),
    ("ddep27", "Extra depreciation in 2027", DDEP27, "US$bn", "MINE. TI: depreciation rises in 2027 at a slower rate (27 Jan 2026 call)."),
    ("ddep28", "Extra depreciation in 2028", DDEP28, "US$bn", "MINE."),
    ("dcg27", "Data centre growth, 2027", DC_G27, "%", "MINE. 2026E about 85%."),
    ("dcg28", "Data centre growth, 2028", DC_G28, "%", "MINE."),
    ("reg27", "Rest of TI growth, 2027", RE_G27, "%", "MINE. Third year of the analog upturn."),
    ("reg28", "Rest of TI growth, 2028", RE_G28, "%", "MINE. Upturn maturing."),
    ("pe", "Target multiple on 2028E EPS (base)", BASE_PE, "x", "MINE. Between the analog peer median (Peers tab) and TI's 30.5x today."),
]
r = 6
for key, label, val, unit, src in rows:
    if label is None:
        ws.cell(row=r, column=1, value=key).font = Font(bold=True, color=NAVY); r += 1; continue
    ws.cell(row=r, column=1, value=label)
    c = ws.cell(row=r, column=2, value=val)
    mine = key in MINE
    c.font = F_INB if mine else F_IN
    if mine: c.fill = FILL_ASSUME
    c.number_format = PCT1 if unit == "%" else (MULT0 if unit == "x" else (BN3 if unit in ("bn",) else "#,##0.00"))
    ws.cell(row=r, column=3, value=unit)
    ws.cell(row=r, column=4, value=src).alignment = WRAP
    A[key] = f"Assumptions!$B${r}"
    r += 1

# ===================================================================== QUARTERLY
ws = sheet("Quarterly", "QUARTERLY HISTORY, Q1 2024 TO Q2 2026, AND MY Q3 / Q4 2026", [44] + [10] * 12)
QH = Q + ["Q3 26E", "Q4 26E"]
head(ws, 4, ["US$ billion unless stated"] + QH)
C = [get_column_letter(2 + j) for j in range(12)]


def qrow(row, label, vals, fmt, bold=False):
    ws.cell(row=row, column=1, value=label).font = F_B if bold else Font()
    for j, v in enumerate(vals):
        c = ws.cell(row=row, column=2 + j, value=v if v is not None else "")
        c.number_format = fmt; c.alignment = RIGHT
        if v is not None and not (isinstance(v, str) and v.startswith("=")):
            c.font = F_IN
        elif isinstance(v, str) and "Assumptions" in v and "*" not in v:
            c.font = F_LINK


bn = lambda xs: [x / 1000 for x in xs]
qrow(5, "Revenue", bn(Q_REV) + ["=" + A["q3rev"], "=" + A["q4rev"]], BN3, True)
qrow(6, "Revenue growth y/y (derived)", [None] * 4 + [f"={C[j]}5/{C[j-4]}5-1" for j in range(4, 12)], PCT1)
qrow(7, "Gross profit", bn(Q_GP) + [None, None], BN3)
qrow(8, "Gross margin (derived)", [f"={c}7/{c}5" for c in C[:10]] + [None, None], PCT1, True)
qrow(9, "Operating profit", bn(Q_OP) + [f"={C[10]}5*{A['q3om']}", f"={C[11]}5*{A['q4om']}"], BN3)
qrow(10, "Operating margin (derived)", [f"={c}9/{c}5" for c in C], PCT1)
qrow(11, "Diluted EPS, US$ (reported; Q3 / Q4 26 mine)", Q_EPS + [
    f"=({C[10]}9+{A['nonopq']})*(1-{A['tax']})/{A['diluted']}",
    f"=({C[11]}9+{A['nonopq']})*(1-{A['tax']})/{A['diluted']}"], USD2, True)
qrow(12, "Depreciation", bn(Q_DEP) + [None, None], BN3)
qrow(13, "Capital expenditures", bn(Q_CAPEX) + [None, None], BN3)
qrow(14, "Cash flow from operations", bn(Q_CFO) + [None, None], BN3)
qrow(15, "Proceeds from CHIPS Act incentives", bn(Q_CHIPS) + [None, None], BN3)
qrow(16, "Free cash flow (TI definition, derived)", [f"={c}14-{c}13+{c}15" for c in C[:10]] + [None, None], BN3, True)
qrow(17, "Analog segment revenue", bn(Q_ANALOG) + [None, None], BN3)
r = 19
ws.cell(row=r, column=1, value="Checks").font = Font(bold=True, color=NAVY)
ws.cell(row=20, column=1, value="Trailing four quarters free cash flow to Q2 26 (TI: US$6.534bn)")
ws["B20"] = "=SUM(H16:K16)"; ws["B20"].number_format = BN3
ws.cell(row=21, column=1, value="Q2 26 depreciation, annualised")
ws["B21"] = "=K12*4"; ws["B21"].number_format = BN3
ws.cell(row=22, column=1, value="Q3 26 EPS guide midpoint")
ws["B22"] = f"=({A['g3elo']}+{A['g3ehi']})/2"; ws["B22"].number_format = USD2
note(ws, 24, ["Source: TI quarterly earnings releases (Exhibit 99 to Form 8-K), 23 April 2024 to 22 July 2026.",
              "Q3 26E and Q4 26E are my estimates; EPS = (operating profit + other income less interest) x (1 - tax) / diluted shares."])

# ===================================================================== DATA CENTRE
ws = sheet("DataCentre", "DATA CENTRE: TI'S STATED GROWTH RATES AND MY DERIVATION OF ITS SIZE", [52, 14, 80])
head(ws, 4, ["Line", "Value", "Working / note"])
dc = [
    ("q4", "Q4 2025, US$bn", "=" + A["dcq4"], BN3, "TI: about US$450m a quarter exiting 2025."),
    ("q3", "Q3 2025, US$bn (derived)", f"=B5/(1+{A['dcq4qq']})", BN3, "Q4 grew mid-single digits on Q3."),
    ("k", "Ratio of Q2 2025 to Q1 2025 (derived)", f"=(1+{A['dcq2qq']})*(1+{A['dcq1yy']})/(1+{A['dcq2yy']})", '0.000', "Q2 26 = 2 x Q2 25 = 1.2 x Q1 26 = 1.2 x 1.9 x Q1 25."),
    ("q1", "Q1 2025, US$bn (derived)", f"=({A['dc25']}-B5-B6)/(1+B7)", BN3, "The four quarters of 2025 sum to US$1.5bn."),
    ("q2", "Q2 2025, US$bn (derived)", "=B7*B8", BN3, ""),
    ("q126", "Q1 2026, US$bn (derived)", f"=B8*(1+{A['dcq1yy']})", BN3, ""),
    ("q226", "Q2 2026, US$bn (derived)", f"=B9*(1+{A['dcq2yy']})", BN3, "My estimate; TI does not disclose the figure."),
    ("sh", "Q2 2026 share of TI revenue", "=B11/Quarterly!K5", PCT1, "Against 9% for 2025 (10-K)."),
    ("gsh", "Share of TI's y/y revenue growth in Q2 2026", "=(B11-B9)/(Quarterly!K5-Quarterly!G5)", PCT1, "Data centre's part of the US$1.0bn rise."),
    ("q4sh", "Q4 2025 share of TI revenue", "=B5/Quarterly!I5", PCT1, ""),
    ("dc26", "2026E data centre revenue, US$bn", f"=B10+B11+{A['dcq3e']}+{A['dcq4e']}", BN3, "MINE for the second half."),
    ("dc26g", "2026E data centre growth", f"=B15/{A['dc25']}-1", PCT1, ""),
]
DCX = {}
for i, (key, lab, f, fmt, nt) in enumerate(dc):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=lab)
    c = ws.cell(row=rr, column=2, value=f); c.number_format = fmt
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
    DCX[key] = f"DataCentre!$B${rr}"
note(ws, 18, ["TI reports data centre growth rates on its calls but not quarterly dollar sizes. The quarterly figures here are my derivation from those rates and are approximate.",
              "Sources: TI earnings calls of 27 January, 22 April and 22 July 2026; 2025 Form 10-K for the 9% share."])

# ===================================================================== FINANCIALS
ws = sheet("Financials", "FINANCIAL SUMMARY AND FORECASTS (US$ billion unless stated)", [50, 12, 12, 12, 12, 12, 12])
head(ws, 4, ["", "2022A", "2024A", "2025A", "2026E*", "2027E*", "2028E*"])
L = lambda k: "=" + A[k]
fin = [
    ("Data centre revenue", [None, None, L("dc25"), "=" + DCX["dc26"], f"=E5*(1+{A['dcg27']})", f"=F5*(1+{A['dcg28']})"], BN2, False),
    ("Rest of TI revenue", [None, None, "=D7-D5", "=E7-E5", f"=E6*(1+{A['reg27']})", f"=F6*(1+{A['reg28']})"], BN2, False),
    ("Revenue", [REV_H[3], REV_H[5], REV_H[6], "=SUM(Quarterly!J5:M5)", "=F5+F6", "=G5+G6"], BN2, True),
    ("Revenue growth", [None, None, "=D7/C7-1", "=E7/D7-1", "=F7/E7-1", "=G7/F7-1"], PCT1, False),
    ("Data centre share of revenue", [None, None, "=D5/D7", "=E5/E7", "=F5/F7", "=G5/G7"], PCT1, False),
    ("Gross margin (reported; not modelled ahead)", [GP_H[3] / REV_H[3], GP_H[5] / REV_H[5], GP_H[6] / REV_H[6], None, None, None], PCT1, False),
    ("Operating profit", [OP_H[3], OP_H[5], OP_H[6], "=SUM(Quarterly!J9:M9)", f"=E11+{A['fall']}*(F7-E7)-{A['ddep27']}", f"=F11+{A['fall']}*(G7-F7)-{A['ddep28']}"], BN2, True),
    ("Operating margin", ["=B11/B7", "=C11/C7", "=D11/D7", "=E11/E7", "=F11/F7", "=G11/G7"], PCT1, False),
    ("Other income less interest", [None, None, None, f"=4*{A['nonopq']}", f"=4*{A['nonopq']}", f"=4*{A['nonopq']}"], BN2, False),
    ("Net income", [None, None, None, f"=(E11+E13)*(1-{A['tax']})", f"=(F11+F13)*(1-{A['tax']})", f"=(G11+G13)*(1-{A['tax']})"], BN2, False),
    ("Diluted EPS, US$", [EPS_H[3], EPS_H[5], EPS_H[6], "=SUM(Quarterly!J11:M11)", f"=F14/{A['diluted']}", f"=G14/{A['diluted']}"], USD2, True),
    ("EPS growth", [None, None, "=D15/C15-1", "=E15/D15-1", "=F15/E15-1", "=G15/F15-1"], PCT1, False),
    ("P/E at reference price", [f"={A['price']}/{c}15" for c in "BCDEFG"], MULT, False),
    ("Depreciation", [DEP_H[3], DEP_H[5], DEP_H[6], f"=({A['deplo']}+{A['dephi']})/2", f"=E18+{A['ddep27']}", f"=F18+{A['ddep28']}"], BN2, False),
    ("Capital expenditures", [CAPEX_H[3], CAPEX_H[5], CAPEX_H[6], f"=({A['capexlo']}+{A['capexhi']})/2", None, None], BN2, False),
    ("Cash flow from operations", [CFO_H[3], CFO_H[5], CFO_H[6], None, None, None], BN2, False),
    ("Proceeds from CHIPS Act incentives", [CHIPS_H[3], CHIPS_H[5], CHIPS_H[6], None, None, None], BN2, False),
    ("Free cash flow (TI definition)", ["=B20-B19+B21", "=C20-C19+C21", "=D20-D19+D21", None, None, None], BN2, True),
    ("Dividends declared per share, US$", [DPS_H[3], DPS_H[5], DPS_H[6], None, None, None], USD2, False),
]
for i, (label, vals, fmt, bold) in enumerate(fin):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=label).font = F_B if bold else Font()
    for j, v in enumerate(vals):
        c = ws.cell(row=rr, column=2 + j, value=v if v is not None else "")
        c.number_format = fmt; c.alignment = RIGHT
        if v is not None and not isinstance(v, str): c.font = F_IN
        elif isinstance(v, str) and v.startswith("=Assumptions"): c.font = F_LINK
note(ws, 26, [
    "*2026E to 2028E are my own estimates. 2026E = reported Q1 and Q2 plus my Q3 and Q4 (Quarterly tab); 2026E EPS sums the quarters.",
    "2027E and 2028E operating profit: prior year + fall-through x extra revenue - extra depreciation. EPS: (operating profit + other income less interest) x (1 - tax) / diluted shares.",
    "History from SEC XBRL company facts (Forms 10-K). Silicon Labs, expected to close in the first half of 2027, is excluded.",
    f"Note check: 2026E revenue {REV26:.2f}, EPS {EPS26:.2f}; 2027E {R27:.2f}, {EPS27:.2f}; 2028E {R28:.2f}, {EPS28:.2f}.",
])

# ===================================================================== VALUATION
ws = sheet("Valuation", "VALUATION: 25x 2028E EPS, TWELVE MONTHS OUT", [58, 16, 70])
head(ws, 4, ["Line", "Value", "Working / note"])
val = [
    ("eps28", "2028E EPS, US$", "=Financials!G15", USD2, "The market will be pricing 2028 by October 2027."),
    ("mult", "Base multiple, x", "=" + A["pe"], MULT0, "Assumptions tab."),
    ("base", "Base-case value per share, US$", "=B5*B6", USD2, "EPS x multiple."),
    ("px", "Reference price, US$", "=" + A["price"], USD2, "Nasdaq close, 5 October 2026."),
    ("upb", "Base value against the price", "=B7/B8-1", PCT1, "Within +/- 15%: no call on value alone."),
    ("wtd", "Probability-weighted value, US$", "=Scenarios!K12", USD2, "Scenarios tab."),
    ("call", "CALL", '=IF(B9>0.15,"LONG",IF(B9<-0.15,"SHORT","NO CALL"))', None, "Rule of thumb: 15% either side of the price. Conviction is judgement. Tommy Lau's call, 6 Oct 2026."),
    (None, "", None, None, ""),
    ("mcap", "Market value on diluted shares, US$bn", f"={A['price']}*{A['diluted']}", BN1, ""),
    ("nd", "Net debt, US$bn", f"={A['debt']}-{A['cash']}", BN2, "30 June 2026, before Silicon Labs (about US$7.5bn)."),
    ("ev", "Enterprise value, US$bn", "=B13+B14", BN1, ""),
    ("dy", "Dividend yield at the new rate", f"=4*{A['dps']}/{A['price']}", '0.00%', "US$6.08 a year."),
    ("rise", "Change since 30 Sep 2025", f"={A['price']}/{A['px25']}-1", PCT1, ""),
    ("hi", "Price below the 52-week high", f"=1-{A['price']}/{A['hi52']}", PCT1, ""),
    (None, "", None, None, ""),
    ("pe26", "P/E on my 2026E", "=Financials!E17", MULT, ""),
    ("pe27", "P/E on my 2027E", "=Financials!F17", MULT, ""),
    ("pec27", "P/E on the 2027 consensus", f"={A['price']}/{A['cons27']}", MULT, "stockanalysis.com US$10.64."),
    ("pe28", "P/E on my 2028E", "=Financials!G17", MULT, ""),
    ("vc", "My 2027E EPS against consensus", f"=Financials!F15/{A['cons27']}-1", PCT1, ""),
    (None, "", None, None, ""),
    ("need", "2028 EPS the price needs at the base multiple, US$", f"={A['price']}/{A['pe']}", USD2, "What the market prices in."),
    ("needop", "2028 operating profit that needs, US$bn", f"=B26*{A['diluted']}/(1-{A['tax']})-4*{A['nonopq']}", BN2, ""),
    ("needrev", "2028 revenue that needs on my margin bridge, US$bn", f"=Financials!F7+(B27-Financials!F11+{A['ddep28']})/{A['fall']}", BN2, "Against my US$26.35bn."),
    ("vs22", "That revenue against the 2022 peak", "=B28/Financials!B7-1", PCT1, "2022 revenue US$20.03bn."),
    ("gm22", "2022 gross margin", "=Financials!B10", PCT1, "Against 61.4% in Q2 2026."),
    ("fcf26", "2026 free cash flow on TI's framework at my revenue, US$bn", "=9.5+(Financials!E7-22)*0.5", BN2, "TI: US$8 to 9bn at US$20bn revenue, US$9 to 10bn at US$22bn (Q2 2026 call); midpoints interpolated."),
    ("fcfps", "Per share, US$", f"=B31/{A['diluted']}", USD2, ""),
    ("fcfy", "Free cash flow yield at the price", f"=B32/{A['price']}", '0.00%', "Includes CHIPS Act cash."),
    ("revisit", "Price that puts the base case 15% above", "=B7/1.15", USD2, "The 'revisit below about US$244' line."),
    ("dc28", "Data centre share of revenue, 2028E", "=Financials!G9", PCT1, ""),
]
V = {}
for i, (key, label, f, fmt, nt) in enumerate(val):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=label)
    if f is not None:
        c = ws.cell(row=rr, column=2, value=f)
        if fmt: c.number_format = fmt
        if key in ("base", "call"):
            ws.cell(row=rr, column=1).font = F_B; c.font = F_B; c.fill = FILL_KEY
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
    if key: V[key] = f"Valuation!$B${rr}"

# ===================================================================== SCENARIOS
ws = sheet("Scenarios", "THREE CASES, TWELVE MONTHS OUT, AND WHICH DRIVER MATTERS", [30] + [12] * 11)
head(ws, 4, ["Case", "DC growth 27", "DC growth 28", "Rest growth 27", "Rest growth 28", "2027 revenue",
             "2028 revenue", "2028 op. profit", "2028 EPS", "Multiple", "Value, US$", "vs price"])
DC26 = "Financials!$E$5"; RE26 = "Financials!$E$6"; RV26 = "Financials!$E$7"; OP26C = "Financials!$E$11"


def scen_row(rr, name, vals, mult, link_base=False):
    ws.cell(row=rr, column=1, value=name).font = F_B
    for j, v in enumerate(vals):
        if isinstance(v, str) and v.startswith("="):
            c = ws.cell(row=rr, column=2 + j, value="=" + v.split("!")[1])
        elif isinstance(v, str):
            c = ws.cell(row=rr, column=2 + j, value="=" + A[v]); c.font = F_LINK
        else:
            c = ws.cell(row=rr, column=2 + j, value=v); c.font = F_INB; c.fill = FILL_ASSUME
        c.number_format = PCT1
    c = ws.cell(row=rr, column=10, value=("=" + A["pe"]) if isinstance(mult, str) else mult)
    if isinstance(mult, str): c.font = F_LINK
    else: c.font = F_INB; c.fill = FILL_ASSUME
    c.number_format = MULT0
    ws.cell(row=rr, column=6, value=f"={DC26}*(1+B{rr})+{RE26}*(1+D{rr})").number_format = BN2
    ws.cell(row=rr, column=7, value=f"={DC26}*(1+B{rr})*(1+C{rr})+{RE26}*(1+D{rr})*(1+E{rr})").number_format = BN2
    ws.cell(row=rr, column=8, value=f"={OP26C}+{A['fall']}*(G{rr}-{RV26})-{A['ddep27']}-{A['ddep28']}").number_format = BN2
    ws.cell(row=rr, column=9, value=f"=(H{rr}+4*{A['nonopq']})*(1-{A['tax']})/{A['diluted']}").number_format = USD2
    ws.cell(row=rr, column=11, value=f"=I{rr}*J{rr}").number_format = USD2
    ws.cell(row=rr, column=12, value=f"=K{rr}/{A['price']}-1").number_format = PCT0


base_keys = ["dcg27", "dcg28", "reg27", "reg28"]
for i, (name, a, b, c_, d, mult, p) in enumerate(SCEN):
    scen_row(5 + i, name, base_keys if name == "Base" else [a, b, c_, d], "pe" if name == "Base" else mult)
ws["A10"] = "Probability"; ws["A10"].font = F_B
for j, s in enumerate(SCEN):
    ws.cell(row=9, column=2 + j, value=s[0]).font = F_B
    c = ws.cell(row=10, column=2 + j, value=s[6]); c.font = F_INB; c.fill = FILL_ASSUME; c.number_format = PCT0
ws["E10"] = "=SUM(B10:D10)"; ws["E10"].number_format = PCT0; ws["F10"] = "must be 100%"; ws["F10"].font = F_NOTE
ws["A12"] = "Probability-weighted value, US$"; ws["A12"].font = F_B
ws["K12"] = "=B10*K5+C10*K6+D10*K7"; ws["K12"].number_format = USD2; ws["K12"].font = F_B; ws["K12"].fill = FILL_KEY
ws["L12"] = f"=K12/{A['price']}-1"; ws["L12"].number_format = PCT1
ws["A14"] = "Which driver matters: swap one driver from bear to bull, the other held at base"; ws["A14"].font = Font(bold=True, color=NAVY)
head(ws, 15, ["Case", "DC growth 27", "DC growth 28", "Rest growth 27", "Rest growth 28", "2027 revenue",
              "2028 revenue", "2028 op. profit", "2028 EPS", "Multiple", "Value, US$", "vs price"])
scen_row(16, "Data centre bear, rest base", ["=Scenarios!B5", "=Scenarios!C5", "reg27", "reg28"], "pe")
scen_row(17, "Data centre bull, rest base", ["=Scenarios!B7", "=Scenarios!C7", "reg27", "reg28"], "pe")
scen_row(18, "Rest bear, data centre base", ["dcg27", "dcg28", "=Scenarios!D5", "=Scenarios!E5"], "pe")
scen_row(19, "Rest bull, data centre base", ["dcg27", "dcg28", "=Scenarios!D7", "=Scenarios!E7"], "pe")
ws["A21"] = "2028 EPS swing from data centre, US$"; ws["I21"] = "=I17-I16"; ws["I21"].number_format = USD2
ws["A22"] = "2028 EPS swing from the rest of TI, US$"; ws["I22"] = "=I19-I18"; ws["I22"].number_format = USD2
ws["A23"] = "Ratio, rest to data centre"; ws["I23"] = "=I22/I21"; ws["I23"].number_format = '0.0"x"'
note(ws, 25, ["All cases start from my 2026E (data centre US$2.77bn, rest US$19.27bn) and use the same margin bridge, in both directions.",
              "Bear: the analog upturn turns in 2027 and data centre slows. Bull: the upturn runs and data centre keeps compounding.",
              "Silicon Labs is excluded from every case."])

# ===================================================================== SENSITIVITY
ws = sheet("Sensitivity", "VALUE PER SHARE: 2028E EPS AGAINST MULTIPLE (US$)", [24, 14, 14, 14])
head(ws, 4, ["2028 EPS \\ multiple", "", "", ""])
for j, (pe, link) in enumerate(zip(SENS_PE, [False, True, False])):
    c = ws.cell(row=4, column=2 + j, value=("=" + A["pe"]) if link else pe)
    c.number_format = MULT0; c.font = Font(bold=True, color="FFFFFF"); c.fill = FILL_HDR; c.alignment = RIGHT
for i, (e, link) in enumerate(zip(SENS_EPS, [False, True, False])):
    rr = 5 + i
    c = ws.cell(row=rr, column=1, value="=Financials!G15" if link else e)
    c.number_format = USD2; c.font = F_LINK if link else F_IN
    for j in range(3):
        col = get_column_letter(2 + j)
        ws.cell(row=rr, column=2 + j, value=f"=$A{rr}*{col}$4").number_format = USD0
ws["C6"].fill = FILL_KEY
ws["A9"] = "Cells above the reference price"; ws["A9"].font = F_B
ws["B9"] = f'=COUNTIF(B5:D7,">"&{A["price"]})'; ws["C9"] = "of 9"
note(ws, 11, ["Middle row links to the model's 2028E EPS and the middle column to the base multiple."])

# ===================================================================== PEERS
ws = sheet("Peers", "PEER MULTIPLES (FORWARD P/E, stockanalysis.com, 6 OCTOBER 2026)", [26, 16, 14, 80])
head(ws, 4, ["Company", "Ticker", "Forward P/E", "What they make"])
for i, (co, tk, pe, what) in enumerate(PEERS):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=co).font = F_B if co == NAME else Font()
    ws.cell(row=rr, column=2, value=tk)
    c = ws.cell(row=rr, column=3, value=("=" + A["fwdpe"]) if co == NAME else pe)
    c.number_format = MULT; c.font = F_LINK if co == NAME else F_IN
    ws.cell(row=rr, column=4, value=what).alignment = WRAP
ws["A13"] = "Median of the six peers"; ws["C13"] = "=MEDIAN(C6:C11)"; ws["C13"].number_format = MULT
ws["A14"] = "Median excluding Monolithic Power"; ws["C14"] = "=MEDIAN(C6:C10)"; ws["C14"].number_format = MULT
note(ws, 16, ["Infineon's figure is from an intraday price on 6 October 2026; the others use closes of 5 October 2026.",
              "Multiples only, to set the base multiple. No view on the peers' shares is expressed."])

# ===================================================================== SOURCES
ws = sheet("Sources", "SOURCES", [6, 130])
head(ws, 4, ["No.", "Source and what it supports"])
srcs = [
    "TI quarterly earnings releases (Exhibit 99 to Form 8-K), 23 April 2024 to 22 July 2026: revenue, gross and operating profit, EPS, depreciation, capital spending, cash flow, CHIPS Act proceeds, balance sheet, segment revenue, Q3 2026 guide.",
    "TI Form 10-K for 2025, filed 6 February 2026: revenue by end market, the 40% 300mm cost advantage, 2026 capital spending of US$2 to 3bn, CHIPS Act terms.",
    "TI Form 10-Q for the second quarter of 2026, filed 24 July 2026: Silicon Labs terms and the US$5bn delayed draw term loan.",
    "SEC XBRL company facts for TI (CIK 97476): annual revenue, gross and operating profit, EPS, cash flow, capital spending, depreciation and dividends, 2019 to 2025.",
    "TI earnings calls of 27 January 2026 (data centre US$1.5bn in 2025, end-market growth, 2026 depreciation of US$2.2 to 2.4bn), 22 April 2026 (data centre about 90% y/y, as reported) and 22 July 2026 (data centre doubled y/y and up about 20% q/q, loadings, fall-through of 70 to 85%, free cash flow framework, inventory days).",
    "TI 8-K of 4 February 2026: agreement to acquire Silicon Labs for US$231 a share, enterprise value about US$7.5bn. TI 8-K of 17 September 2026: quarterly dividend raised to US$1.52.",
    "TI announcement of 1 October 2026: Q3 2026 results call on 21 October 2026, 3:30pm Central time.",
    "Yahoo Finance chart API, retrieved 6 October 2026: month-end closes, the 2 and 5 October closes, 52-week range.",
    "stockanalysis.com, retrieved 6 October 2026: consensus revenue, EPS and target; forward P/E for TI and peers. MarketBeat for the results date cross-check.",
    "The Physical Layer, 'The most expensive chip in an AI rack runs on Texas Instruments' cheapest ones', 2 October 2026: https://thephysicallayer.fyi/journal/ti-feeds-the-chip/",
    "Estimates for 2026 to 2028 are the author's own. Personal research, not investment advice. I hold no position in Texas Instruments.",
]
for i, s in enumerate(srcs):
    ws.cell(row=5 + i, column=1, value=i + 1)
    ws.cell(row=5 + i, column=2, value=s).alignment = WRAP

# ===================================================================== COVER
ws = wb["Cover"]
head(ws, 4, ["Item", "Value", "Note"])
cover = [
    ("Call", "=" + V["call"], None, "Tommy Lau's call, published 6 Oct 2026."),
    ("Conviction", CALL["conviction"], None, "Judgement, not formula."),
    ("Reference price, US$", "=" + A["price"], USD2, "Nasdaq close, 5 October 2026."),
    ("Base value, US$", "=" + V["base"], USD2, "25x 2028E EPS."),
    ("Base value against the price", "=" + V["upb"], PCT1, ""),
    ("Probability-weighted value, US$", "=" + V["wtd"], USD2, "25 / 50 / 25 bear / base / bull."),
    ("2026E / 2027E / 2028E EPS, US$", '=TEXT(Financials!E15,"0.00")&" / "&TEXT(Financials!F15,"0.00")&" / "&TEXT(Financials!G15,"0.00")', None, "Own estimates."),
    ("Horizon", "12 months, to October 2027", None, ""),
    ("Wrong if", "See note", None, "Full year 2027 revenue of US$26bn or more with data centre above 15% of it; or gross margin of 65% or more in any quarter to June 2027; or the shares trade above US$340 by October 2027 without either."),
    ("Revisit if", "See note", None, "A price below about US$244; or data centre at 15% or more of revenue while industrial still grows; or a gross margin of 65% or more."),
]
for i, (lab, v, fmt, nt) in enumerate(cover):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=lab).font = F_B
    c = ws.cell(row=rr, column=2, value=v)
    if fmt: c.number_format = fmt
    c.alignment = RIGHT
    c.font = F_LINK if isinstance(v, str) and v.startswith("=") else F_IN
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
ws["A17"] = "Tabs"; ws["A17"].font = Font(bold=True, color=NAVY)
for i, (t, d) in enumerate([
    ("Assumptions", "Every input, with its source. Change the yellow cells to move the value."),
    ("Quarterly", "Q1 2024 to Q2 2026 reported, my Q3 and Q4 2026; margins, depreciation, capex, free cash flow."),
    ("DataCentre", "Data centre growth rates as TI states them, and my derivation of its size and share."),
    ("Financials", "2022A, 2024A to 2028E: data centre and the rest, the margin bridge, EPS, P/E, capex and free cash flow."),
    ("Valuation", "Base value, market value, multiples, what the price needs, free cash flow yield."),
    ("Scenarios", "Bear / base / bull, the probability-weighted value, and which driver moves earnings most."),
    ("Sensitivity", "2028E EPS against multiple."),
    ("Peers", "Forward multiples used to set the base multiple."),
    ("Sources", "Every source used."),
]):
    ws.cell(row=18 + i, column=1, value=t); ws.cell(row=18 + i, column=2, value=d)
ws["A28"] = "Colour key: blue = hardcoded input; yellow fill = my own assumption; green = link to another tab; black = formula."
ws["A28"].font = F_NOTE
ws["A29"] = f"Companion note: {DATE}_{COMPANY}_Initiation.pdf. Personal research, not investment advice."; ws["A29"].font = F_NOTE

for w in wb.worksheets:
    w.freeze_panes = "A5"
wb.active = 0
wb.save(XLSX)
print("model written", XLSX)
