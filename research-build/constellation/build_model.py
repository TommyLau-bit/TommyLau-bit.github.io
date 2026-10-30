"""Constellation Energy Corporation (Nasdaq: CEG) initiation model. Earnings on Constellation's own base and enhanced
framework to 2028E, P/E valuation, scenarios, sensitivity, Crane arithmetic. Live formulas throughout. House style follows
the Vertiv and Eaton models: navy header rows, blue font for hardcoded inputs, yellow fill on my own assumptions, green
font for cross-sheet links, black for formulas.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from data import *

NAVY = "1F3864"; MUTED = "595959"
F_TTL = Font(bold=True, color=NAVY, size=12)
F_SUB = Font(color=MUTED, size=9)
F_HDR = Font(bold=True, color="FFFFFF", size=10)
F_IN = Font(color="0000FF")
F_INB = Font(color="0000FF", bold=True)
F_LINK = Font(color="008000")
F_B = Font(bold=True)
F_NOTE = Font(color=MUTED, size=9, italic=True)
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_ASSUME = PatternFill("solid", fgColor="FFFF00")
FILL_KEY = PatternFill("solid", fgColor="D9E2F3")
WRAP = Alignment(wrap_text=True, vertical="top")
RIGHT = Alignment(horizontal="right")

USD2 = '"US$"#,##0.00'
USD0 = '"US$"#,##0'
M0 = '#,##0'
M1 = '#,##0.0'
M2 = '#,##0.00'
PCT1 = '0.0%'
PCT0 = '0%'
MULTF = '0.0"x"'
MULT0 = '0"x"'
SUB = f"Tommy Lau | Constellation Energy Corporation (Nasdaq: CEG) | {DATE} | Personal research, not investment advice."

wb = openpyxl.Workbook()


def sheet(name, title, widths):
    ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet(name)
    ws.title = name
    ws["A1"] = title; ws["A1"].font = F_TTL
    ws["A2"] = SUB; ws["A2"].font = F_SUB
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def head(ws, row, cells, start=1):
    for i, c in enumerate(cells, start):
        x = ws.cell(row=row, column=i, value=c)
        x.font = F_HDR; x.fill = FILL_HDR
        x.alignment = Alignment(horizontal="left" if i == start else "right", vertical="center", wrap_text=True)


def note(ws, row, texts):
    for i, t in enumerate(texts):
        ws.cell(row=row + i, column=1, value=t).font = F_NOTE


def put(ws, ref, v, fmt=None, font=None, fill=None):
    c = ws[ref]; c.value = v
    if fmt: c.number_format = fmt
    if font: c.font = font
    elif isinstance(v, str) and v.startswith("=Assumptions"): c.font = F_LINK
    if fill: c.fill = fill
    c.alignment = RIGHT
    return c


# =====================================================================  COVER (filled last)
ws = sheet("Cover", "CONSTELLATION ENERGY CORPORATION (NASDAQ: CEG)  INITIATION MODEL", [42, 24, 100])

# =====================================================================  ASSUMPTIONS
ws = sheet("Assumptions", "ASSUMPTIONS AND INPUTS", [64, 14, 10, 110])
ws["A3"] = ("Blue = hardcoded input. Yellow fill = my own assumption (change these to move the value). "
            "Green = link to another tab. Black = formula. US$ unless stated. Year end 31 December.")
ws["A3"].font = F_NOTE
head(ws, 5, ["Input", "Value", "Unit", "Source / note"])
A = {}
Q2R = "Q2 2026 earnings release, Form 8-K Ex. 99.1, 6 August 2026"
OUT = "2026 Business and Earnings Outlook, Form 8-K Ex. 99.1, 31 March 2026"
rows = [
    ("MARKET", None, None, None, None),
    ("price", "Reference share price", PRICE, "US$", "Nasdaq close, 29 October 2026. Yahoo Finance history, read 30 October 2026; matches stockanalysis.com."),
    ("hiclose", "Highest close, 15 October 2025", HI_CLOSE, "US$", "Yahoo Finance. Intraday high US$412.70 the same day."),
    ("pxdec", "Close, 31 December 2025", PX_DEC25, "US$", "Yahoo Finance."),
    ("shares_out", "Common shares outstanding, 31 July 2026", SHARES_OUT, "m", "Form 10-Q for the quarter to 30 June 2026, cover page: 354,307,379."),
    ("cash", "Cash and cash equivalents, 30 June 2026", CASH, "US$m", Q2R + "."),
    ("debt", "Debt, 30 June 2026", DEBT, "US$m", Q2R + ": short-term borrowings 5,226 + long-term debt due within one year 363 + long-term debt 19,111."),
    ("nci", "Noncontrolling interests, 30 June 2026", NCI, "US$m", Q2R + "."),
    ("sales", "Pending sale proceeds, before closing adjustments and tax", SALE_PROCEEDS, "US$m", "Form 10-Q: PJM assets to LS Power US$5.0bn; Brazos Valley US$860m; both expected to close by end 2026."),
    ("GUIDANCE, FRAMEWORK AND CONSENSUS", None, None, None, None),
    ("g_lo", "2026 adjusted operating EPS guide, low", GUIDE26["lo"], "US$", Q2R + ". Raised from US$11.00 to 12.00."),
    ("g_hi", "2026 adjusted operating EPS guide, high", GUIDE26["hi"], "US$", Q2R + "."),
    ("eps24", "2024 adjusted operating EPS", ADJ_EPS["2024A"], "US$", "Q4 2025 earnings release, 24 February 2026."),
    ("eps25", "2025 adjusted operating EPS", ADJ_EPS["2025A"], "US$", "Q4 2025 earnings release, 24 February 2026."),
    ("q1", "Q1 2026 adjusted operating EPS", 2.74, "US$", "Q1 2026 earnings release, 11 May 2026."),
    ("q2", "Q2 2026 adjusted operating EPS", 2.55, "US$", Q2R + "."),
    ("b26", "Base EPS 2026, midpoint of US$6.65 to 6.75", BASE["2026"], "US$", OUT + ". Includes assets to be divested in H2 2026."),
    ("b27", "Base EPS 2027, midpoint of US$7.60 to 7.70", BASE["2027"], "US$", OUT + "."),
    ("b29", "Base EPS 2029, midpoint of US$11.40 to 11.90", BASE["2029"], "US$", OUT + ". Base EPS growth of 20%+ from 2026 to 2029."),
    ("ppa_lo", "2029 EPS per 1 GW nuclear PPA at a US$20/MWh premium to the PTC floor", PPA_SENS[0], "US$", OUT + ": US$0.40 to 1.00 for US$20 to 50/MWh."),
    ("cons26", "2026 consensus adjusted EPS", CONS_EPS26, "US$", "stockanalysis.com (S&P Global), checked 28 October 2026, retrieved 30 October 2026. Cross-check only."),
    ("cons27", "2027 consensus adjusted EPS", CONS_EPS27, "US$", "stockanalysis.com, as above. Cross-check only."),
    ("constp", "Average analyst target", CONS_TP, "US$", "stockanalysis.com: 25 analysts, low US$238, high US$399. Cross-check only."),
    ("CRANE AND THE FLEET", None, None, None, None),
    ("crane_mw", "Crane output", CRANE_MW, "MW", "Constellation release of 20 September 2024; maximum facility output 859 MW per the FERC waiver affidavit."),
    ("crane_capex", "Crane restart capital spending, estimate", CRANE_CAPEX, "US$m", "Constellation Q3 2024 results, 4 November 2024."),
    ("cir", "Capacity interconnection rights moved from Eddystone 3 and 4", CIR_MW, "MW", "FERC order, 1 June 2026; 380 MW each per the waiver request of 31 March 2026."),
    ("nuc_twh", "Expected nuclear output, 2026", NUC_TWH_26, "TWh", OUT + " (million MWh, including Salem and STP)."),
    ("avail", "Nuclear output available to contract", AVAIL_TWH, "TWh", OUT + " (million MWh a year)."),
    ("mtax", "Marginal statutory tax rate", MARG_TAX, "%", Q2R + "."),
    ("repl", "Replacement cost of the Calpine natural gas assets", REPL_BN, "US$bn", OUT + ": ~US$65bn, over 2x the amount paid."),
    ("calp_ev", "Calpine enterprise value at announcement", CALPINE_EV, "US$bn", OUT + ": US$26.6bn enterprise value, 27.7 GW acquired."),
    ("goog_mw", "Google uprates", GOOG_UPRATE_MW, "MW", "Constellation and Google release, 6 October 2026; first uprate by 2028."),
    ("goog_capex", "Google uprate investment, at least", GOOG_CAPEX, "US$m", "As above: more than US$4.3bn."),
    ("cap_mw", "PJM 2028/2029 cleared capacity, all Constellation plants", CAP_MW_ALL, "MW", "Form 8-K, 14 July 2026 (15,700 MW nuclear)."),
    ("cap_px", "PJM 2028/2029 clearing price", CAP_PRICE, "US$/MW-day", "Form 8-K, 14 July 2026."),
    ("MY ESTIMATES", None, None, None, None),
    ("q3e", "Q3 2026 adjusted operating EPS", Q3E, "US$", "MY ESTIMATE. MarketBeat consensus US$3.84; Q3 2025 was US$3.04."),
    ("q4e", "Q4 2026 adjusted operating EPS", Q4E, "US$", "MY ESTIMATE. Q4 2025 was US$2.30; Calpine adds."),
    ("enh27", "Enhanced share of total EPS, 2027", ENH27, "%", "MY ASSUMPTION. Midpoint of Constellation's 35 to 40%."),
    ("enh28", "Enhanced share of total EPS, 2028", ENH28, "%", "MY ASSUMPTION. Between the 2027 range and the 30 to 35% for 2029."),
    ("crane_cf", "Crane capacity factor", CRANE_CF, "%", "MY ASSUMPTION. Fleet 92 to 95% in recent quarters; used only for the per-MWh sensitivity."),
    ("pe_base", "Base multiple, P/E on 2028E adjusted operating EPS", BASE_PE, "x", "MY ASSUMPTION. About one turn below the 21.3x forward P/E (stockanalysis.com); peers 9.6x to 14.4x."),
]
MINE = {"q3e", "q4e", "enh27", "enh28", "crane_cf", "pe_base"}
r = 6
for key, label, val, unit, src in rows:
    if label is None:
        ws.cell(row=r, column=1, value=key).font = Font(bold=True, color=NAVY)
        r += 1; continue
    ws.cell(row=r, column=1, value=label)
    c = ws.cell(row=r, column=2, value=val)
    c.font = F_INB if key in MINE else F_IN
    if key in MINE: c.fill = FILL_ASSUME
    c.number_format = PCT1 if unit == "%" else (MULT0 if unit == "x" else (M2 if unit == "m" else (USD2 if unit == "US$" else M0)))
    ws.cell(row=r, column=3, value=unit)
    ws.cell(row=r, column=4, value=src).alignment = WRAP
    A[key] = f"Assumptions!$B${r}"
    r += 1
L = lambda k: "=" + A[k]

# =====================================================================  EARNINGS
ws = sheet("Earnings", "ADJUSTED OPERATING EPS ON CONSTELLATION'S BASE AND ENHANCED FRAMEWORK (US$ PER SHARE)", [52, 13, 13, 13, 13, 13, 13])
head(ws, 4, ["US$ per share", "2024A", "2025A", "2026E*", "2027E*", "2028E*", "2029 (company)"])
labs = ["Base EPS (Constellation midpoint; 2028 my path)", "Enhanced share of total", "ADJUSTED OPERATING EPS", "Growth",
        "Enhanced EPS (total less base)", "Guide midpoint / consensus", "Mine against it", "Implied 2029 total at 32.5% enhanced"]
for i, lab in enumerate(labs):
    ws.cell(row=5 + i, column=1, value=lab).font = F_B if i in (0, 2) else Font()
put(ws, "B7", L("eps24"), USD2); put(ws, "C7", L("eps25"), USD2)
put(ws, "D5", L("b26"), USD2); put(ws, "E5", L("b27"), USD2)
put(ws, "F5", f"=SQRT({A['b27']}*{A['b29']})", USD2)
put(ws, "G5", L("b29"), USD2)
put(ws, "D7", f"={A['q1']}+{A['q2']}+{A['q3e']}+{A['q4e']}", USD2, F_B, FILL_KEY)
put(ws, "D6", "=1-D5/D7", PCT1)
put(ws, "E6", L("enh27"), PCT1); put(ws, "F6", L("enh28"), PCT1)
put(ws, "E7", "=E5/(1-E6)", USD2, F_B, FILL_KEY); put(ws, "F7", "=F5/(1-F6)", USD2, F_B, FILL_KEY)
for col, prev in [("C", "B"), ("D", "C"), ("E", "D"), ("F", "E")]:
    put(ws, f"{col}8", f"={col}7/{prev}7-1", PCT1)
for col in "DEF":
    put(ws, f"{col}9", f"={col}7-{col}5", USD2)
put(ws, "D10", f"=({A['g_lo']}+{A['g_hi']})/2", USD2); put(ws, "E10", L("cons27"), USD2)
put(ws, "D11", "=D7/D10-1", PCT1); put(ws, "E11", "=E7/E10-1", PCT1)
put(ws, "G12", "=G5/(1-0.325)", USD2)
ws["A14"] = "Consensus 2027 against my 2027E"; put(ws, "B14", "=E10/E7-1", PCT1)
ws["A15"] = "Base EPS growth a year, 2026 to 2029 (company midpoints)"; put(ws, "B15", f"=({A['b29']}/{A['b26']})^(1/3)-1", PCT1)
ws["A16"] = "2027 guide midpoint that makes a long if 2028E is scaled in proportion to 2027E"; put(ws, "B16", "=Valuation!B23/(F7/E7)", USD2)
ws["A17"] = "Test guide midpoint, US$"; put(ws, "B17", 13.0, USD2, F_INB, FILL_ASSUME)
ws["A18"] = "Value if 2028E scales in proportion, US$"; put(ws, "B18", f"=B17*F7/E7*{A['pe_base']}", USD2)
ws["A19"] = "Value on my framework path alone (base 2027 implied by the guide at 37.5% enhanced), US$"
put(ws, "B19", f"=SQRT(B17*(1-{A['enh27']})*{A['b29']})/(1-{A['enh28']})*{A['pe_base']}", USD2)
note(ws, 21, [
    "*2026E = Q1 and Q2 actual plus my Q3 and Q4. 2027E and 2028E = base EPS / (1 - enhanced share), on Constellation's framework of 31 March 2026.",
    "Base earnings (Constellation): contracts, nuclear at the PTC floor with 2% inflation, minimum fossil earnings, 10-year average commercial margins.",
    "Enhanced: prices above base assumptions, commercial margins above average, volatility. Company ranges: ~40% (2026), 35 to 40% (2027), 30 to 35% (2029).",
    "The 2028 base figure is my geometric interpolation; Constellation gave 2026, 2027 and 2029 only. Its 2029 projection includes announced long-term offtakes (Crane).",
])

# =====================================================================  VALUATION
ws = sheet("Valuation", "VALUATION: P/E ON 2028E ADJUSTED OPERATING EPS, TWELVE MONTHS OUT", [64, 16, 70])
head(ws, 4, ["Item", "Value", "Working / note"])
vrows = [
    (5, "2028E adjusted operating EPS, US$", "=Earnings!F7", USD2, "My estimate."),
    (6, "Base multiple", L("pe_base"), MULTF, "About one turn below the forward P/E today."),
    (7, "BASE VALUE PER SHARE, US$", "=B5*B6", USD2, ""),
    (8, "Reference price, US$", L("price"), USD2, "Nasdaq close, 29 October 2026."),
    (9, "Base value against price", "=B7/B8-1", PCT1, "Band: LONG at +15% or more, SHORT at -25% or worse, NO CALL between."),
    (10, "VIEW ON THE BAND", '=IF(B9>=0.15,"LONG",IF(B9<=-0.25,"SHORT","NO CALL"))', None, "Tommy Lau's call, 30 October 2026: NO CALL, Medium conviction."),
    (11, "Target, US$", CALL["target"] or "None (no call)", USD2, "None for a no call."),
    (12, "Target against price", "=IF(ISNUMBER(B11),B11/B8-1,\"n.a.\")", PCT1, ""),
    (14, "Market value, US$bn", f"={A['shares_out']}*B8/1000", '#,##0.0', "Shares outstanding at 31 July 2026."),
    (15, "Net debt, US$m, 30 June 2026", f"={A['debt']}-{A['cash']}", M0, "Before the pending sale proceeds."),
    (16, "Enterprise value, US$bn", f"=B14+(B15+{A['nci']})/1000", '#,##0.0', "Before decommissioning, leases and pensions."),
    (17, "P/E on 2026 guide midpoint", "=B8/Earnings!D10", MULTF, ""),
    (18, "P/E on 2026E (mine)", "=B8/Earnings!D7", MULTF, ""),
    (19, "P/E on 2027 consensus", "=B8/" + A["cons27"], MULTF, "stockanalysis.com."),
    (20, "P/E on 2027E (mine)", "=B8/Earnings!E7", MULTF, ""),
    (21, "P/E on 2028E (mine): what the price pays", "=B8/Earnings!F7", MULTF, ""),
    (22, "Multiple of 2028E needed for a value 15% above the price", "=B8*1.15/B5", MULTF, ""),
    (23, "2028E EPS needed at the base multiple for a long, US$", "=B8*1.15/B6", USD2, ""),
    (24, "Price below which the base value is 15% or more above it, US$", "=B7/1.15", USD2, "Revisit level for a long."),
    (25, "Price above which the base value is 25% or more below it, US$", "=B7/0.75", USD2, "Revisit level for a short."),
    (26, "Price against the highest close", f"=B8/{A['hiclose']}-1", PCT1, ""),
    (27, "Price change in 2026 to date", f"=B8/{A['pxdec']}-1", PCT1, ""),
    (28, "Average analyst target against price", f"={A['constp']}/B8-1", PCT1, "Cross-check only."),
    (29, "Dividend yield at the price", f"={DIV_Q}*4/B8", '0.00%', "Quarterly dividend US$0.4265 (Q4 2025 release)."),
]
for rr, lab, f, fmt, nt in vrows:
    ws.cell(row=rr, column=1, value=lab)
    c = ws.cell(row=rr, column=2, value=f)
    if fmt: c.number_format = fmt
    c.alignment = RIGHT
    if rr in (7, 10):
        ws.cell(row=rr, column=1).font = F_B; c.font = F_B; c.fill = FILL_KEY
    elif isinstance(f, str) and (f.startswith("=Assumptions") or f.startswith("=Earnings")): c.font = F_LINK
    elif rr == 11: c.font = F_INB; c.fill = FILL_ASSUME
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP

# =====================================================================  SCENARIOS
ws = sheet("Scenarios", "THREE CASES, TWELVE MONTHS OUT (VALUE = MULTIPLE x 2028E EPS)", [10, 16, 16, 12, 14, 14, 12])
head(ws, 4, ["Case", "2028 base EPS", "Enhanced share", "Multiple", "2028E EPS", "Value, US$", "vs price"])
for i, s in enumerate(SCEN):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=s[0]).font = F_B
    if s[0] == "Base":
        for col, ref in [(2, "=Earnings!F5"), (3, "=" + A["enh28"]), (4, "=" + A["pe_base"])]:
            c = ws.cell(row=rr, column=col, value=ref); c.font = F_LINK
    else:
        for col, v in [(2, s[1]), (3, s[2]), (4, s[3])]:
            c = ws.cell(row=rr, column=col, value=v); c.font = F_INB; c.fill = FILL_ASSUME
    ws.cell(row=rr, column=2).number_format = USD2; ws.cell(row=rr, column=3).number_format = PCT1
    ws.cell(row=rr, column=4).number_format = MULT0
    ws.cell(row=rr, column=5, value=f"=B{rr}/(1-C{rr})").number_format = USD2
    ws.cell(row=rr, column=6, value=f"=E{rr}*D{rr}").number_format = USD2
    ws.cell(row=rr, column=7, value=f"=F{rr}/{A['price']}-1").number_format = PCT1
ws["A9"] = "Probability"; ws["A9"].font = F_B
for j, s in enumerate(SCEN):
    ws.cell(row=8, column=2 + j, value=s[0]).font = F_B
    c = ws.cell(row=9, column=2 + j, value=s[4]); c.font = F_INB; c.fill = FILL_ASSUME; c.number_format = PCT0
ws["E9"] = "=SUM(B9:D9)"; ws["E9"].number_format = PCT0; ws["F9"] = "must be 100%"; ws["F9"].font = F_NOTE
ws["A11"] = "Probability-weighted value, US$"; ws["A11"].font = F_B
ws["E11"] = "=B9*F5+C9*F6+D9*F7"; ws["E11"].number_format = USD2; ws["E11"].font = F_B; ws["E11"].fill = FILL_KEY
ws["F11"] = f"=E11/{A['price']}-1"; ws["F11"].number_format = PCT1
ws["A12"] = "Check: base case equals the Valuation tab"; ws["E12"] = "=F6-Valuation!B7"; ws["E12"].number_format = '0.000'
ws["A13"] = "Bull add-on (MY ASSUMPTION: no start date disclosed; half a year in 2028): 2.7 GW at the low-end 2029 PPA sensitivity, US$"; ws["E13"] = f"=2.7*{A['ppa_lo']}/2"; ws["E13"].number_format = USD2
note(ws, 15, [
    "Bear: PJM caps capacity prices for longer and forward power prices fall; Crane slips; base 2028 EPS US$8.60 (about 9% below my path), enhanced 30%; 15x, near Vistra.",
    "Base: Constellation's framework at the midpoints; 20x.",
    "Bull: the 2,700 MW Google supply deal priced at the low end of Constellation's own sensitivity (about US$0.50 in 2028), enhanced holds at 40%; 23x.",
])

# =====================================================================  SENSITIVITY
ws = sheet("Sensitivity", "VALUE PER SHARE: 2028E EPS AGAINST THE P/E (US$)", [30, 14, 14, 14])
head(ws, 4, ["2028E EPS \\ P/E", "", "", ""])
for j, (m, link) in enumerate(zip(SENS_PE, [False, True, False])):
    c = ws.cell(row=4, column=2 + j, value=("=" + A["pe_base"]) if link else m)
    c.number_format = MULT0; c.font = Font(bold=True, color="FFFFFF"); c.fill = FILL_HDR; c.alignment = RIGHT
for i, (e, link) in enumerate(zip(SENS_EPS, [False, True, False])):
    rr = 5 + i
    c = ws.cell(row=rr, column=1, value="=Earnings!F7" if link else e)
    c.number_format = USD2; c.font = F_LINK if link else F_IN
    for j in range(3):
        col = get_column_letter(2 + j)
        ws.cell(row=rr, column=2 + j, value=f"=$A{rr}*{col}$4").number_format = USD0
ws["C6"].fill = FILL_KEY
ws["A9"] = "Cells above the reference price"; ws["A9"].font = F_B
ws["B9"] = f'=COUNTIF(B5:D7,">"&{A["price"]})'; ws["C9"] = "of 9"
ws["A10"] = "Cells 15% or more above the price"; ws["A10"].font = F_B
ws["B10"] = f'=COUNTIF(B5:D7,">="&{A["price"]}*1.15)'; ws["C10"] = "of 9"
ws["A11"] = "Cells 25% or more below the price"; ws["A11"].font = F_B
ws["B11"] = f'=COUNTIF(B5:D7,"<="&{A["price"]}*0.75)'; ws["C11"] = "of 9"
note(ws, 13, ["The middle row links to the model's 2028E EPS and the middle column to the base multiple."])

# =====================================================================  CRANE
ws = sheet("Crane", "CRANE AND THE GRID PLACE: THE CLAIM'S ARITHMETIC", [70, 16, 70])
head(ws, 4, ["Item", "Value", "Note"])
crows = [
    (5, "Crane output, TWh a year", f"={A['crane_mw']}*8760*{A['crane_cf']}/1000000", M2, "My capacity factor."),
    (6, "Share of 2026 expected nuclear output", f"=B5/{A['nuc_twh']}", PCT1, ""),
    (7, "After-tax EPS per US$10/MWh of Crane margin, US$", f"=B5*1000000*10*(1-{A['mtax']})/({A['shares_out']}*1000000)", '0.000', "Marginal statutory rate; shares at 31 July 2026."),
    (8, "Crane restart cost per kW, US$", f"={A['crane_capex']}*1000/{A['crane_mw']}", M0, "US$1.6bn estimate of November 2024."),
    (9, "Calpine gas assets: replacement cost over enterprise value paid", f"={A['repl']}/{A['calp_ev']}", MULTF, "Constellation's own estimates; my division."),
    (10, "Google uprates cost per kW, US$", f"={A['goog_capex']}*1000/{A['goog_mw']}", M0, "At least; 'more than US$4.3bn'."),
    (11, "Eddystone rights against Crane output", f"={A['cir']}/{A['crane_mw']}", PCT1, ""),
    (12, "PJM 2028/2029 capacity revenue, US$m a year", f"={A['cap_mw']}*{A['cap_px']}*365/1000000", M0, "Cleared MW x price x 365; my arithmetic."),
]
for rr, lab, f, fmt, nt in crows:
    ws.cell(row=rr, column=1, value=lab)
    c = ws.cell(row=rr, column=2, value=f); c.number_format = fmt; c.alignment = RIGHT
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
note(ws, 14, [
    "Crane timeline: restart announced 20 Sep 2024 (target 2028); PJM Reliability Resource Initiative selection 2 May 2025 brought the target to 2027;",
    "PJM Phase I report 31 Oct 2025: contingent 765 kV and 500 kV lines with in-service dates as late as Dec 2030; soonest full deliverability 31 Dec 2030;",
    "waiver request 31 Mar 2026 (ER26-2028-000): able to generate H2 2027, NRC approvals early 2027; FERC granted 1 Jun 2026; NRC fuel licence amendment approved (Q2 2026).",
    "The claim's test: Crane at full output on the grid by 31 Dec 2028.",
])

# =====================================================================  PEERS
ws = sheet("Peers", "PEER MULTIPLES (stockanalysis.com, UPDATED 29 OCTOBER 2026)", [22, 14, 13, 14, 70])
head(ws, 4, ["Company", "Ticker", "Forward P/E", "EV/EBITDA", "What they run"])
for i, (co, tk, pe, eve, what, grp) in enumerate(PEERS):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=co).font = F_B if grp == "self" else Font()
    ws.cell(row=rr, column=2, value=tk)
    for j, v in enumerate([pe, eve]):
        c = ws.cell(row=rr, column=3 + j, value=v); c.number_format = MULTF; c.font = F_IN
    ws.cell(row=rr, column=5, value=what).alignment = WRAP
last = 4 + len(PEERS)
ws[f"A{last + 2}"] = "Peer median (excluding Constellation)"; ws[f"A{last + 2}"].font = F_B
ws[f"C{last + 2}"] = f"=MEDIAN(C6:C{last})"; ws[f"C{last + 2}"].number_format = MULTF
ws[f"A{last + 3}"] = "Constellation premium over the peer median, turns"; ws[f"A{last + 3}"].font = F_B
ws[f"C{last + 3}"] = f"=C5-C{last + 2}"; ws[f"C{last + 3}"].number_format = MULTF
note(ws, last + 5, ["Multiples only, to set the range. No view on the peers' shares is expressed. Talen had a trailing net loss, so its EV/EBITDA is not meaningful."])

# =====================================================================  SOURCES
ws = sheet("Sources", "SOURCES", [6, 150])
head(ws, 4, ["No.", "Source and what it supports"])
srcs = [
    "Constellation second quarter 2026 earnings release, Form 8-K Ex. 99.1, 6 August 2026: Q2 adjusted EPS, raised 2026 guide, balance sheet, FERC waiver, NRC fuel licence, Brazos Valley sale, nuclear output.",
    "Constellation Form 10-Q, quarter to 30 June 2026, filed 6 August 2026: Note 3 (Crane CIRs), Note 2 (Calpine purchase and divestitures), Note 13 (DOE loan), repurchases, shares outstanding.",
    "Constellation first quarter 2026 release, 11 May 2026, and fourth quarter 2025 release, 24 February 2026: quarterly and annual adjusted EPS.",
    "Constellation 2026 Business and Earnings Outlook, Form 8-K Ex. 99.1, 31 March 2026: base and enhanced earnings, 2029 framework, sensitivities, nuclear volumes, Calpine replacement cost.",
    "Constellation Form 8-K, 14 July 2026: PJM 2028/2029 capacity auction. Form 8-K, 2 June 2026: secondary offering and repurchase. Form 8-K/A, 20 April 2026: Calpine closing, 7 January 2026.",
    "Constellation request for limited waiver, FERC docket ER26-2028-000, 31 March 2026, and Habtemichael affidavit: contingent facilities, December 2030, Eddystone 380 MW units, H2 2027.",
    "FERC order of 1 June 2026, as quoted by Utility Dive (5 June 2026) and the American Nuclear Society (5 June 2026).",
    "Constellation and Google announcement, 6 October 2026: 890 MW of uprates, 11 units, 20-year PPA, more than US$4.3bn, separate 2,700 MW 15-year agreement.",
    "Constellation Q3 2024 results, 4 November 2024: about US$1.6bn of capital spending for the Crane restart.",
    "Yahoo Finance CEG history pages, read 30 October 2026; stockanalysis.com (S&P Global data) for consensus, targets and peer multiples; MarketBeat for Q3 2026 consensus and date (estimated).",
    "The Physical Layer, 'Constellation can bring a closed reactor back by 2027. The grid connection was the hard part.', 30 October 2026: https://thephysicallayer.fyi/journal/constellation-restarts-the-reactor/",
    "Estimates for the second half of 2026, 2027 and 2028 are the author's own. Personal research, not investment advice.",
]
for i, s in enumerate(srcs):
    ws.cell(row=5 + i, column=1, value=i + 1)
    ws.cell(row=5 + i, column=2, value=s).alignment = WRAP

# =====================================================================  COVER
ws = wb["Cover"]
head(ws, 4, ["Item", "Value", "Note"])
cover = [
    ("Call (Tommy Lau, 30 October 2026)", "=Valuation!B10", None, "LONG at +15% or more, SHORT at -25% or worse, NO CALL between."),
    ("Conviction", CALL["conviction"], None, "Judgement, not formula."),
    ("Base value per share, US$", "=Valuation!B7", USD2, "Base multiple x 2028E adjusted operating EPS."),
    ("Reference price, US$", "=" + A["price"], USD2, "Nasdaq close, 29 October 2026."),
    ("Base value against price", "=Valuation!B9", PCT1, ""),
    ("Probability-weighted value, US$", "=Scenarios!E11", USD2, "25 / 50 / 25 bear / base / bull."),
    ("2028E adjusted operating EPS, US$", "=Earnings!F7", USD2, "My estimate."),
    ("Horizon", "12 months, to October 2027", None, ""),
    ("Revisit if", "See note", None, REVISIT_IF),
    ("Wrong if", "See note", None, WRONG_IF),
]
for i, (lab, v, fmt, nt) in enumerate(cover):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=lab).font = F_B
    c = ws.cell(row=rr, column=2, value=v)
    if fmt: c.number_format = fmt
    c.alignment = RIGHT
    c.font = F_LINK if isinstance(v, str) and v.startswith("=") else F_IN
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
ws.row_dimensions[13].height = 44; ws.row_dimensions[14].height = 58
ws["A17"] = "Tabs"; ws["A17"].font = Font(bold=True, color=NAVY)
for i, (t, d) in enumerate([
    ("Assumptions", "Every input, with its source. Change the yellow cells to move the value."),
    ("Earnings", "Adjusted operating EPS on Constellation's base and enhanced framework, 2024A to 2028E."),
    ("Valuation", "Base value, the band, multiples at the price, revisit levels."),
    ("Scenarios", "Bear / base / bull and the probability-weighted value."),
    ("Sensitivity", "2028E EPS against the multiple."),
    ("Crane", "The restart and grid-place arithmetic behind the claim."),
    ("Peers", "Forward P/E and EV/EBITDA used to set the multiple."),
    ("Sources", "Every source used."),
]):
    ws.cell(row=18 + i, column=1, value=t)
    ws.cell(row=18 + i, column=2, value=d)
ws["A27"] = "Colour key: blue = hardcoded input; yellow fill = my own assumption; green = link to another tab; black = formula."
ws["A27"].font = F_NOTE
ws["A28"] = f"Companion note: {DATE}_{COMPANY}_Initiation.pdf. Personal research, not investment advice."
ws["A28"].font = F_NOTE

for w in wb.worksheets:
    w.freeze_panes = "A5"
wb.active = 0
wb.save(XLSX)
print("model written", XLSX)
