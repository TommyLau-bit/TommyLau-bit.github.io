"""Oklo initiation model, 6 Oct 2026. Live formulas throughout.

House style follows the Vertiv and Bloom Energy models: navy header rows, blue font for hardcoded inputs,
yellow fill on my own assumptions, green font for cross-sheet links, black for formulas. US$ million unless stated.
Method: a risked value of the Aurora fleet (plant by plant NPV, probability of the first plant), plus cash, less
corporate spending, per fully diluted share, twelve months out (October 2027).
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from data import *

NAVY = "1F3864"; MUTED = "595959"
F_TTL = Font(bold=True, color=NAVY, size=12)
F_SUB = Font(color=MUTED, size=9)
F_HDR = Font(bold=True, color="FFFFFF", size=10)
F_IN = Font(color="0000FF")
F_INB = Font(color="0000FF", bold=True)
F_LINK = Font(color="008000")
F_B = Font(bold=True)
F_NOTE = Font(color=MUTED, size=9, italic=True)
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_ASSUME = PatternFill("solid", fgColor="FFFF00")
FILL_KEY = PatternFill("solid", fgColor="D9E2F3")
WRAP = Alignment(wrap_text=True, vertical="top")
RIGHT = Alignment(horizontal="right")
USD2 = '"US$"#,##0.00'; USD0 = '"US$"#,##0'; M1 = '#,##0.0'; M0 = '#,##0'; M3 = '#,##0.000'
PCT1 = '0.0%'; PCT0 = '0%'; MULT = '0.0"x"'
SUB = f"Tommy Lau | Oklo (NYSE: OKLO) | {DATE_LONG} | US$ million unless stated | Personal research, not investment advice."

wb = openpyxl.Workbook()


def sheet(name, title, widths):
    ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet(name)
    ws.title = name
    ws["A1"] = title; ws["A1"].font = F_TTL
    ws["A2"] = SUB; ws["A2"].font = F_SUB
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def head(ws, row, cells, start=1):
    for i, c in enumerate(cells, start):
        x = ws.cell(row=row, column=i, value=c)
        x.font = F_HDR; x.fill = FILL_HDR
        x.alignment = Alignment(horizontal="left" if i == start else "right", vertical="center", wrap_text=True)


def put(ws, ref, v, fmt=None, bold=False, key=False):
    c = ws[ref]; c.value = v
    if fmt: c.number_format = fmt
    if isinstance(v, str) and v.startswith("="):
        c.font = Font(color="008000", bold=bold) if "!" in v else Font(bold=bold)
    else:
        c.font = F_INB if bold else F_IN
    if key: c.fill = FILL_KEY
    c.alignment = RIGHT
    return c


sheet("Cover", "OKLO (NYSE: OKLO)  INITIATION MODEL", [38, 24, 80])

# ===================================================================== ASSUMPTIONS
ws = sheet("Assumptions", "INPUTS AND ASSUMPTIONS", [62, 14, 9, 80])
ws["A3"] = ("Blue = hardcoded input. Yellow fill = my own assumption (change these to move the value). Green = link to "
            "another tab. Black = formula. US$ million unless stated; per-share figures in US$.")
ws["A3"].font = F_NOTE
head(ws, 5, ["Input", "Value", "Unit", "Source / note"])
A = {}
MINE = {"ocf26", "capex26", "ocf27", "capex27", "unitmw", "cf", "opex", "esc", "life", "rplant", "rdev", "inlrem", "inlcod"}
rows = [
    ("MARKET (retrieved 6 October 2026)", None, None, None, None),
    ("price", "Share price (NYSE: OKLO)", PRICE, "US$", "NYSE close, 5 October 2026. Yahoo Finance."),
    ("hi52", "52-week high (intraday, 15 October 2025)", HI52, "US$", "Yahoo Finance."),
    ("hiclose", "Highest close (14 October 2025)", HI_CLOSE, "US$", "Yahoo Finance."),
    ("lo52", "52-week low (intraday, 14 September 2026)", LO52, "US$", "Yahoo Finance."),
    ("constp", "Average analyst target", CONS_TP, "US$", "stockanalysis.com, 25 analysts; median US$74.50, range US$14 to 130."),
    ("conseps26", "2026 consensus EPS", CONS_EPS_26, "US$", "stockanalysis.com."),
    ("SHARES (10-Q Q2 2026; 8-K and 424B5 of 11 Sep 2026)", None, None, None, None),
    ("shjun", "Shares outstanding, 30 Jun 2026, m", SH_JUN, "m", "10-Q balance sheet: 185,090,155."),
    ("atmsh", "Shares sold under the May 2026 ATM to 10 Sep 2026, m", ATM_MAY_SH, "m", "8-K of 11 Sep 2026: 17,971,448 shares for about US$1bn gross."),
    ("atmgross", "Gross proceeds of the May 2026 ATM", ATM_MAY_GROSS, "US$m", "8-K: approximately US$1,000m."),
    ("atmq2sh", "Of which sold by 30 Jun 2026, m", ATM_Q2_SH, "m", "10-Q Note 8: 10,712,054."),
    ("atmq2gross", "Of which gross proceeds by 30 Jun 2026", ATM_Q2_GROSS, "US$m", "10-Q Note 8: 680.4."),
    ("atmfee", "ATM commission", ATM_FEE, "%", "Up to 1.5% (10-Q, 424B5)."),
    ("options", "Stock options outstanding, 30 Jun 2026, m", OPTIONS, "m", "424B5: 5,738,353 at a weighted exercise price of US$2.07."),
    ("rsus", "Restricted stock units outstanding, 30 Jun 2026, m", RSUS, "m", "424B5: 3,677,081."),
    ("CASH (10-Q Q2 2026, 10-K 2025)", None, None, None, None),
    ("cashjun", "Cash, equivalents and marketable securities, 30 Jun 2026", CASH_JUN, "US$m", "10-Q: 1,644.7 + 820.5 + 541.1 = 3,006.3. Excludes 16.9 restricted."),
    ("cashdec", "Same, 31 Dec 2025", CASH_DEC25, "US$m", "10-K: 788.4 + 439.5 + 184.6."),
    ("raisedh1", "Net proceeds from share sales, H1 2026", RAISED_H1_26, "US$m", "10-Q cash flow statement: 1,851.9."),
    ("h1ocf", "Operating cash use, H1 2026", -H1["ocf"], "US$m", "10-Q: 65.5."),
    ("h1capex", "Capital spending, H1 2026", H1["capex"], "US$m", "10-Q: 126.9."),
    ("gocflo", "2026 operating cash use guide, low", GUIDE_26["ocf"][0], "US$m", "Q2 business update, 7 Aug 2026 (prior 80 to 100)."),
    ("gocfhi", "2026 operating cash use guide, high", GUIDE_26["ocf"][1], "US$m", "Same."),
    ("gcaplo", "2026 capital spending guide, low", GUIDE_26["capex"][0], "US$m", "Same (prior 350 to 450)."),
    ("gcaphi", "2026 capital spending guide, high", GUIDE_26["capex"][1], "US$m", "Same."),
    ("MY ASSUMPTIONS: CASH", None, None, None, None),
    ("ocf26", "Operating cash use, 2026", OCF26, "US$m", "MINE. Inside the guide."),
    ("capex26", "Capital spending, 2026", CAPEX26, "US$m", "MINE. Midpoint of the guide."),
    ("ocf27", "Operating cash use, 2027", OCF27, "US$m", "MINE."),
    ("capex27", "Capital spending, 2027", CAPEX27, "US$m", "MINE. Aurora-INL construction, fuel fabrication, Ohio site work."),
    ("MY ASSUMPTIONS: THE PLANT", None, None, None, None),
    ("unitmw", "Aurora unit size, MWe", UNIT_MW, "MW", "MINE. Oklo's 75 MWe design (10-Q)."),
    ("cf", "Capacity factor", CF, "%", "MINE."),
    ("opex", "Operating cost, fuel, insurance and decommissioning, US$/MWh", OPEX_MWH, "US$", "MINE. Oklo expects 70 to 80 long-term staff for the INL plant and fuel facility (22 Sep 2025)."),
    ("esc", "Escalation of price and cost, a year", ESC, "%", "MINE."),
    ("life", "Plant life, years", LIFE, "yrs", "MINE. Matches decades-long PPAs."),
    ("rplant", "Discount rate, operating contracted plant", R_PLANT, "%", "MINE."),
    ("rdev", "Discount rate, from first power back to October 2027", R_DEV, "%", "MINE. Development and funding risk; the first-plant probability is separate."),
    ("inlrem", "Aurora-INL capital still to spend after Sep 2027", INL_REM, "US$m", "MINE. Oklo has not disclosed the total; the CFO said in August it is still narrowing it with Kiewit."),
    ("inlcod", "Aurora-INL first power (year, mid-year = .5)", INL_COD, "year", "MINE. Oklo targets 2028 (Q2 2026 update)."),
]
r = 6
for key, label, val, unit, src in rows:
    if label is None:
        ws.cell(row=r, column=1, value=key).font = Font(bold=True, color=NAVY); r += 1; continue
    ws.cell(row=r, column=1, value=label)
    c = ws.cell(row=r, column=2, value=val)
    mine = key in MINE
    c.font = F_INB if mine else F_IN
    if mine: c.fill = FILL_ASSUME
    c.number_format = PCT1 if unit == "%" else ("0.00" if unit in ("US$",) else ("0.0" if unit == "year" else M3 if unit == "m" else M1))
    ws.cell(row=r, column=3, value=unit)
    ws.cell(row=r, column=4, value=src).alignment = WRAP
    A[key] = f"Assumptions!$B${r}"
    r += 1

# ===================================================================== SHARES AND CASH
ws = sheet("Cash", "SHARES AND CASH: WHAT IS IN THE BANK AND WHO PAID FOR IT", [64, 16, 70])
head(ws, 4, ["Line", "Value", "Working / note"])
L = [
    ("Shares sold 1 Jul to 10 Sep 2026, m", f"={A['atmsh']}-{A['atmq2sh']}", M3, "May 2026 ATM total less the part sold by 30 June."),
    ("Gross proceeds 1 Jul to 10 Sep 2026", f"={A['atmgross']}-{A['atmq2gross']}", M1, "Approximate: the 8-K gives about US$1bn."),
    ("Average sale price, 1 Jul to 10 Sep 2026, US$", "=B6/B5", USD2, "Q1 2026 US$96.95; Q2 2026 US$63.51 (10-Q)."),
    ("Net proceeds after commission", f"=B6*(1-{A['atmfee']})", M1, ""),
    ("Shares outstanding by 10 Sep 2026, at least, m", f"={A['shjun']}+B5", M3, "Excludes awards vested since June and any sales under the September ATM."),
    ("Fully diluted shares, m", f"=B9+{A['options']}+{A['rsus']}", M3, "Options at US$2.07 counted in full; the strike proceeds are trivial."),
    (None, None, None, ""),
    ("Rise in cash and securities, H1 2026", f"={A['cashjun']}-{A['cashdec']}", M1, "Where the September piece wrote about US$1.9bn."),
    ("Raised from share sales, H1 2026 (net)", f"={A['raisedh1']}", M1, ""),
    ("Spent and invested, H1 2026", "=B13-B12", M1, "Operating cash, capital spending, acquisitions, other investments."),
    (None, None, None, ""),
    ("2026 spending, mine (operating cash use plus capital spending)", f"={A['ocf26']}+{A['capex26']}", M1, "Guide 520 to 650."),
    ("H2 2026 spending, mine", f"=B16-{A['h1ocf']}-{A['h1capex']}", M1, ""),
    ("Q3 2026 spending, mine (half of H2)", "=B17/2", M1, ""),
    ("Cash and securities, 30 Sep 2026, mine", f"={A['cashjun']}+B8-B18", M1, "Before any sales under the September ATM."),
    ("Cash and securities, end 2026, mine", f"={A['cashjun']}+B8-B17", M1, ""),
    ("2027 spending, mine", f"={A['ocf27']}+{A['capex27']}", M1, ""),
    ("Cash and securities, end Sep 2027, mine", "=B20-0.75*B21", M1, "Used in the valuation."),
    ("Cash per fully diluted share, Sep 2027, US$", "=B22/B10", USD2, "Before any further spending."),
    ("Market value at the price, US$m", f"={A['price']}*B9", M0, ""),
    ("Enterprise value, US$m", "=B24-B19", M0, "Market value less my 30 Sep 2026 cash."),
    ("Market value against cash and securities (x)", "=B24/B19", '0.00"x"', ""),
    ("Share count rise since 30 Jun 2024", f"=B9/{Q_SHARES[0]}-1", PCT1, "122.1m at 30 Jun 2024 (10-Q)."),
]
for i, (label, f, fmt, nt) in enumerate(L):
    rr = 5 + i
    if label is None: continue
    ws.cell(row=rr, column=1, value=label)
    put(ws, f"B{rr}", f, fmt)
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
FDR = "Cash!$B$10"; CASH27 = "Cash!$B$22"

# ===================================================================== HISTORY
ws = sheet("History", "REPORTED HISTORY, Q2 2024 TO Q2 2026, AND ANNUAL", [44] + [11] * 9)
head(ws, 4, ["US$ million"] + Q)
hist = [("Cash and marketable securities, end", Q_CASHSEC, M1), ("Shares outstanding, end, m", Q_SHARES, M1),
        ("Capital spending", Q_CAPEX, M1), ("Operating cash flow", Q_OCF, M1), ("Net loss", Q_NETLOSS, M1)]
for i, (lab, vals, fmt) in enumerate(hist):
    ws.cell(row=5 + i, column=1, value=lab)
    for j, v in enumerate(vals):
        put(ws, f"{get_column_letter(2 + j)}{5 + i}", v, fmt)
ws["A10"] = "Operating cash use plus capital spending"
for j in range(len(Q)):
    c = get_column_letter(2 + j); put(ws, f"{c}10", f"={c}7-{c}8", M1)
head(ws, 12, ["Annual", "2024A", "2025A", "H1 2026A", "2026 guide", "2026E*", "2027E*"])
ann = [
    ("Revenue", [0, 0, H1["rev"], "", CONS_REV_26, ""]),
    ("Operating loss", [OPLOSS_H[0], OPLOSS_H[1], H1["oploss"], "", "", ""]),
    ("Net loss", [NETLOSS_H[0], NETLOSS_H[1], H1["netloss"], "", "", ""]),
    ("Operating cash use", [-OCF_H[0], -OCF_H[1], f"={A['h1ocf']}", "120 to 150", f"={A['ocf26']}", f"={A['ocf27']}"]),
    ("Capital spending", [CAPEX_H[0], CAPEX_H[1], f"={A['h1capex']}", "400 to 500", f"={A['capex26']}", f"={A['capex27']}"]),
    ("Cash and securities, end of period", [CASHSEC_H[0], f"={A['cashdec']}", f"={A['cashjun']}", "", "=Cash!B20", "=Cash!B20-Cash!B21"]),
    ("Shares outstanding, end, m", [SH_H[0], SH_H[1], f"={A['shjun']}", "", "=Cash!B9", ""]),
    ("Net proceeds from share sales", [0, RAISED_25, f"={A['raisedh1']}", "", "", ""]),
]
for i, (lab, vals) in enumerate(ann):
    ws.cell(row=13 + i, column=1, value=lab)
    for j, v in enumerate(vals):
        put(ws, f"{get_column_letter(2 + j)}{13 + i}", v, M1)
ws["A22"] = ("*2026E revenue is the stockanalysis.com consensus (US$2.74m); spending estimates are mine. Q2 2024 to Q2 2026 from "
             "Oklo 10-Qs and the 10-K via SEC XBRL; quarterly flows derived from year-to-date figures.")
ws["A22"].font = F_NOTE

# ===================================================================== FLEET
ws = sheet("Fleet", "THE FLEET: WHEN UNITS COME ON AND WHAT EACH IS WORTH TODAY", [30, 14, 18, 18])
head(ws, 4, ["First power (mid-year)", "Meta MW", "Further units (1 = yes)", "Discount to Oct 2027"])
yrs = list(range(2030, 2041))
for i, y in enumerate(yrs):
    rr = 5 + i
    put(ws, f"A{rr}", y, "0")
    put(ws, f"B{rr}", META.get(y, 0), M0)
    put(ws, f"C{rr}", 1 if y in PIPE_YEARS else 0, "0")
    put(ws, f"D{rr}", f"=1/(1+{A['rdev']})^(A{rr}+0.5-2027.75)", "0.000")
last = 4 + len(yrs)
ws[f"A{last + 2}"] = "Meta MW, discounted (S_meta)"; put(ws, f"B{last + 2}", f"=SUMPRODUCT(B5:B{last},D5:D{last})", M1)
ws[f"A{last + 3}"] = "Further years, discounted (S_rate)"; put(ws, f"B{last + 3}", f"=SUMPRODUCT(C5:C{last},D5:D{last})", "0.000")
ws[f"A{last + 4}"] = "Aurora-INL discount"; put(ws, f"B{last + 4}", f"=1/(1+{A['rdev']})^({A['inlcod']}-2027.75)", "0.000")
ws[f"A{last + 5}"] = "Present value factor, 40 years with escalation"; put(ws, f"B{last + 5}", f"=(1-((1+{A['esc']})/(1+{A['rplant']}))^{A['life']})/({A['rplant']}-{A['esc']})", "0.000")
ws[f"A{last + 6}"] = "Further units a year: count of years"; put(ws, f"B{last + 6}", f"=SUM(C5:C{last})", "0")
ws[f"A{last + 8}"] = ("Meta: phase one as early as 2030 and 1.2 GW by 2034 (Oklo and Meta, 9 Jan 2026); the yearly split is mine. "
                      "Further units from 2033 to 2040 stand for Switch and other customers; the rate is set per scenario.")
ws[f"A{last + 8}"].font = F_NOTE
SMETA = f"Fleet!$B${last + 2}"; SRATE = f"Fleet!$B${last + 3}"; DINL = f"Fleet!$B${last + 4}"; PVF = f"Fleet!$B${last + 5}"
NYRS = f"Fleet!$B${last + 6}"

head(ws, last + 10, ["Corporate and non-plant spending from Oct 2027", "US$m a year", "Discount", "Present value"])
for i, cst in enumerate(CORP):
    rr = last + 11 + i
    put(ws, f"A{rr}", i + 1, "0"); c = put(ws, f"B{rr}", cst, M0); c.fill = FILL_ASSUME; c.font = F_INB
    put(ws, f"C{rr}", f"=1/(1+{A['rdev']})^(A{rr}-0.5)", "0.000"); put(ws, f"D{rr}", f"=B{rr}*C{rr}", M1)
cl = last + 10 + len(CORP)
ws[f"A{cl + 1}"] = "Present value, base and bull"; put(ws, f"D{cl + 1}", f"=SUM(D{last + 11}:D{cl})", M1, bold=True)
ws[f"A{cl + 2}"] = "Present value, bear (wind-down: first three years at 150)"
put(ws, f"D{cl + 2}", f"=150*(C{last + 11}+C{last + 12}+C{last + 13})", M1, bold=True)
PVC = f"Fleet!$D${cl + 1}"; PVCB = f"Fleet!$D${cl + 2}"


def margin_f(pr):
    return f"(8760*{A['cf']}*({pr}-{A['opex']})/1000000)"


def npvmw_f(pr, cx):
    return f"({margin_f(pr)}*{PVF}-{cx}/1000)"


def inl_f(pr):
    return f"({margin_f(pr)}*{A['unitmw']}*{PVF}-{A['inlrem']})"


def value_f(pr, cx, p, rate):
    fleet = f"{p}*({inl_f(pr)}*{DINL}+MAX(0,{npvmw_f(pr, cx)})*({SMETA}+{rate}*{SRATE}))"
    return f"=IF({p}=0,({CASH27}-{PVCB})/{FDR},({fleet}+{CASH27}-{PVC})/{FDR})"


# ===================================================================== SCENARIOS
ws = sheet("Scenarios", "BEAR / BASE / BULL, TWELVE MONTHS OUT", [22, 12, 13, 12, 13, 10, 14, 14, 14, 14, 12, 11, 60])
head(ws, 4, ["Case", "Price US$/MWh", "Cost US$/kW", "P(first plant)", "MW a year after Meta", "Weight",
             "NPV per MW, US$m", "Risked fleet, US$m", "Value per share", "Change vs price", "GW by 2040", "", "What happens"])
desc = ["Aurora-INL slips past 2029 and the cost does not close against the power price; no fleet; spending winds down",
        "Aurora-INL runs in 2028 with 60% probability; Meta's 1.2 GW by 2034, then 300 MW a year",
        "Aurora-INL on time; costs fall fast; Switch and others contract: 900 MW a year from 2033"]
for i, (nm, pr, cx, p, rate, w) in enumerate(SCEN):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=nm).font = F_B
    for col, v, fmt in [("B", pr, M0), ("C", cx, M0), ("D", p, PCT0), ("E", rate, M0), ("F", w, PCT0)]:
        c = put(ws, f"{col}{rr}", v, fmt); c.fill = FILL_ASSUME
    put(ws, f"G{rr}", "=" + npvmw_f(f"B{rr}", f"C{rr}"), "0.00")
    put(ws, f"H{rr}", f"=IF(D{rr}=0,0,D{rr}*({inl_f(f'B{rr}')}*{DINL}+MAX(0,G{rr})*({SMETA}+E{rr}*{SRATE})))", M0)
    put(ws, f"I{rr}", value_f(f"B{rr}", f"C{rr}", f"D{rr}", f"E{rr}"), USD2, bold=True, key=True)
    put(ws, f"J{rr}", f"=I{rr}/{A['price']}-1", PCT0)
    put(ws, f"K{rr}", f"=IF(D{rr}=0,0,({A['unitmw']}+SUM(Fleet!$B$5:$B${last})+E{rr}*{NYRS})/1000)", "0.0")
    ws.cell(row=rr, column=13, value=desc[i]).alignment = WRAP
ws["A9"] = "Probability-weighted"; ws["A9"].font = F_B
put(ws, "I9", "=SUMPRODUCT(F5:F7,I5:I7)", USD2, bold=True, key=True); put(ws, "J9", f"=I9/{A['price']}-1", PCT0)
ws["A11"] = ("Value = (probability x risked NPV of the fleet discounted to Oct 2027 + my cash at Sep 2027 - present value of "
             "corporate spending) / fully diluted shares. A unit whose NPV is negative is not built (MAX(0, ...)); Aurora-INL is "
             "built regardless. No tax and no tax credits are modelled. The bear case keeps only cash, less a wind-down.")
ws["A11"].font = F_NOTE

# ===================================================================== SENSITIVITY
ws = sheet("Sensitivity", "VALUE PER SHARE: POWER PRICE AGAINST COST PER KW (BASE PROBABILITY AND BUILD RATE)", [26, 14, 14, 14])
head(ws, 4, ["Cost US$/kW \\ price US$/MWh"] + [f"{p}" for p in SENS_PRICE])
for j, pr in enumerate(SENS_PRICE):
    c = put(ws, f"{get_column_letter(2 + j)}3", pr, M0); c.fill = FILL_ASSUME
for i, cx in enumerate(SENS_CAPEX):
    rr = 5 + i
    c = put(ws, f"A{rr}", cx, M0); c.fill = FILL_ASSUME
    for j in range(3):
        col = get_column_letter(2 + j)
        put(ws, f"{col}{rr}", value_f(f"{col}$3", f"$A{rr}", "Scenarios!$D$6", "Scenarios!$E$6"), USD2,
            key=(i == 1 and j == 1))
ws["A9"] = "Cells above the price"; put(ws, "B9", f"=COUNTIF(B5:D7,\">\"&{A['price']})", "0")
ws["A10"] = "Highest cell"; put(ws, "B10", "=MAX(B5:D7)", USD2)
ws["A11"] = "Lowest cell"; put(ws, "B11", "=MIN(B5:D7)", USD2)
ws["A13"] = ("Where a cell's units do not pay, only Aurora-INL is built but the base corporate spending is kept, so the "
             "lowest cells sit below the bear case's wind-down value.")
ws["A13"].font = F_NOTE

# ===================================================================== VALUATION
ws = sheet("Valuation", "VALUE, WHAT THE PRICE NEEDS, AND THE PLANT", [66, 16, 60])
head(ws, 4, ["Line", "Value", "Note"])
pr, cx, p, rt = "Scenarios!$B$6", "Scenarios!$C$6", "Scenarios!$D$6", "Scenarios!$E$6"
V = [
    ("Base value per share, Oct 2027, US$", "=Scenarios!I6", USD2, ""),
    ("Base against the price", f"=B5/{A['price']}-1", PCT1, ""),
    ("Probability-weighted value, US$", "=Scenarios!I9", USD2, ""),
    ("Weighted against the price", f"=B7/{A['price']}-1", PCT1, ""),
    ("Rule on value (long if base +15% or more, short if -25% or less)", f"=IF(B6>=0.15,\"LONG\",IF(B6<=-0.25,\"SHORT\",\"NO CALL\"))", None, "The rule reads short. The target is the probability-weighted value, rounded, because the base rests on undisclosed cost and price."),
    ("Call", CALL["direction"], None, "Tommy Lau's call, published 7 Oct 2026."),
    ("Margin per MW in the first year at the base price, US$m", "=" + margin_f(pr), "0.000", ""),
    ("NPV per MW at first power, base, US$m", "=" + npvmw_f(pr, cx), "0.00", ""),
    ("Break-even cost per kW at the base price, US$", f"={margin_f(pr)}*{PVF}*1000", M0, "Above this, a unit destroys value."),
    ("Flat power price that just repays the base cost, US$/MWh", f"=({cx}/1000/((1-(1+{A['rplant']})^-{A['life']})/{A['rplant']}))/(8760*{A['cf']}/1000000)+{A['opex']}", M0, "No escalation; for comparison only."),
    ("Risked fleet value the price needs, US$m", f"={A['price']}*{FDR}-{CASH27}+{PVC}", M0, ""),
    ("Base risked fleet value, US$m", "=Scenarios!H6", M0, ""),
    ("Further MW a year the price needs, base economics", f"=((B15/{p}-{inl_f(pr)}*{DINL})/{npvmw_f(pr, cx)}-{SMETA})/{SRATE}", M0, "Base: 300."),
    ("GW by 2040 the price needs", f"=({A['unitmw']}+SUM(Fleet!$B$5:$B${last})+B17*{NYRS})/1000", "0.0", "Announced: Switch 12 GW + Meta 1.2 GW."),
    ("Or: cost per kW the price needs, at the base fleet, US$", f"=({margin_f(pr)}*{PVF}-(B15/{p}-{inl_f(pr)}*{DINL})/({SMETA}+{rt}*{SRATE}))*1000", M0, ""),
    ("Market value, US$m", "=Cash!B24", M0, ""),
    ("Enterprise value, US$m", "=Cash!B25", M0, ""),
    ("Below the highest close", f"=1-{A['price']}/{A['hiclose']}", PCT1, ""),
    ("Cover level: close the short below, US$ (base value)", "=B5", USD2, ""),
    ("Short wrong if the shares close above, US$", 60, USD2, "On the path to the bull case."),
    ("Consensus target against the price", f"={A['constp']}/{A['price']}-1", PCT1, ""),
    ("Darlington BWRX-300, four units, C$ per kW", f"={DARLINGTON[0]}/{DARLINGTON[1]}*1000", M0, "OPG, 8 May 2025: C$20.9bn for 1.2 GW."),
    ("Darlington, first unit with shared costs, C$ per kW", f"={DARLINGTON[2]}/{DARLINGTON[3]}*1000", M0, "C$7.7bn for 300 MW."),
]
for i, (label, f, fmt, nt) in enumerate(V):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=label)
    put(ws, f"B{rr}", f, fmt, bold=rr in (5, 7))
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP

# ===================================================================== PEERS
ws = sheet("Peers", "PEERS: ENTERPRISE VALUE PER GIGAWATT OF ANNOUNCED CUSTOMER CAPACITY", [16, 13, 11, 12, 13, 13, 11, 13, 70])
head(ws, 4, ["Company", "Ticker", "Close 5 Oct", "Shares m", "Net cash US$m", "EV US$m", "Announced GW", "EV per GW, US$bn", "What they make"])
for i, (nm, tk, px, sh, nc, gw, what) in enumerate(PEERS):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=nm); ws.cell(row=rr, column=2, value=tk)
    if nm == "Oklo":
        put(ws, f"C{rr}", f"={A['price']}", USD2); put(ws, f"D{rr}", "=Cash!B9", M1); put(ws, f"E{rr}", "=Cash!B19", M0)
    else:
        put(ws, f"C{rr}", px, USD2); put(ws, f"D{rr}", sh, M1); put(ws, f"E{rr}", nc, M0)
    put(ws, f"F{rr}", f"=C{rr}*D{rr}-E{rr}", M0)
    if gw:
        put(ws, f"G{rr}", gw, "0.0"); put(ws, f"H{rr}", f"=F{rr}/G{rr}/1000", "0.00")
    else:
        ws.cell(row=rr, column=7, value="n/m"); ws.cell(row=rr, column=8, value="n/m")
    ws.cell(row=rr, column=9, value=what).alignment = WRAP
ws["A10"] = ("Closes from Yahoo Finance, 5 Oct 2026. Peer shares and net cash from stockanalysis.com, retrieved 6 Oct 2026. Oklo: "
             "my share count and my 30 Sep 2026 cash. Announced GW is customer capacity under agreements, none of it under a "
             "binding power purchase agreement. Multiples only; no view on the peers' shares.")
ws["A10"].font = F_NOTE

# ===================================================================== SOURCES
ws = sheet("Sources", "SOURCES", [130])
for i, s in enumerate([
    "Oklo Form 10-Q for the quarter to 30 June 2026, filed 7 August 2026: balance sheet, statements of operations and cash flows, ATM programmes, customer agreements, Meta prepayment agreement, Centrus letter of intent, DOE and NRC status, Groves.",
    "Oklo Form 10-Q/A for the quarter to 31 March 2026, filed 17 June 2026 (certification only; figures unchanged).",
    "Oklo Form 10-K for 2025, filed 17 March 2026: non-binding agreements (risk factors), right of first refusal payment, 2025 figures.",
    "Oklo Form 8-K and 424B5, 11 September 2026: end of the May 2026 ATM (17,971,448 shares, about US$1bn gross), new US$1bn ATM, options and RSUs outstanding.",
    "SEC XBRL company facts (CIK 1849056) for quarterly history.",
    "Oklo second quarter 2026 business update and call, 7 August 2026: 2026 guidance, Aurora-INL in 2028, CFO remarks.",
    "Oklo and Meta announcement, 9 January 2026. Oklo groundbreaking release, 22 September 2025.",
    "Ontario Power Generation and Ontario government, Darlington BWRX-300 approval, 8 May 2025.",
    "Yahoo Finance chart API for prices; stockanalysis.com for peers, consensus, analyst targets and short interest; retrieved 6 October 2026.",
    "Estimates and the valuation are the author's own. Personal research, not investment advice.",
]):
    ws.cell(row=4 + i, column=1, value=s).alignment = WRAP

# ===================================================================== COVER
ws = wb["Cover"]
cov = [
    ("Company", "Oklo Inc. (NYSE: OKLO)", None, "Designs Aurora fast reactors it intends to own, and sells the power."),
    ("Date", DATE_LONG, None, ""),
    ("Price", f"={A['price']}", USD2, "NYSE close, 5 October 2026."),
    ("Base value, Oct 2027", "=Scenarios!I6", USD2, "Risked fleet plus cash, less corporate spending, per diluted share."),
    ("Probability-weighted value", "=Scenarios!I9", USD2, "25 / 50 / 25."),
    ("Call", CALL["direction"], None, "Tommy Lau's call, published 7 Oct 2026."),
    ("Target, US$", CALL["target"], USD2, "Probability-weighted value, rounded."),
    ("Conviction", CALL["conviction"], None, ""),
]
for i, (k, v, fmt, nt) in enumerate(cov):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=k).font = F_B
    c = ws.cell(row=rr, column=2, value=v)
    if fmt: c.number_format = fmt
    c.alignment = RIGHT
    c.font = F_LINK if isinstance(v, str) and v.startswith("=") else F_IN
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
ws["A15"] = "Tabs"; ws["A15"].font = Font(bold=True, color=NAVY)
for i, (t, d) in enumerate([
    ("Assumptions", "Every input, with its source. Change the yellow cells to move the value."),
    ("Cash", "Shares sold, fully diluted count, cash raised against cash kept, my cash path to Sep 2027."),
    ("History", "Quarterly and annual reported figures, the 2026 guide and my spending estimates."),
    ("Fleet", "When units come online, discount factors, corporate spending."),
    ("Scenarios", "Bear / base / bull and the probability-weighted value."),
    ("Sensitivity", "Power price against cost per kW."),
    ("Valuation", "Plant economics, what the price needs, revisit levels."),
    ("Peers", "Enterprise value per announced gigawatt."),
    ("Sources", "Every source used."),
]):
    ws.cell(row=16 + i, column=1, value=t); ws.cell(row=16 + i, column=2, value=d)
ws["A27"] = "Colour key: blue = hardcoded input; yellow fill = my own assumption; green = link to another tab; black = formula."
ws["A27"].font = F_NOTE
ws["A28"] = f"Companion note: {DATE}_{COMPANY}_Initiation.pdf. Personal research, not investment advice."; ws["A28"].font = F_NOTE

order = ["Cover", "Assumptions", "Cash", "History", "Fleet", "Scenarios", "Sensitivity", "Valuation", "Peers", "Sources"]
wb._sheets = [wb[n] for n in order]
for w in wb.worksheets:
    w.freeze_panes = "A5"
wb.active = 0
wb.save(XLSX)
print("model written", XLSX)
