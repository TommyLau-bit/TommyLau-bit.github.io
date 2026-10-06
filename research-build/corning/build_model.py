"""Corning initiation model, 6 Oct 2026. Live formulas throughout.

House style follows the Vertiv, TSMC and Texas Instruments models: navy header rows, blue font for hardcoded inputs,
yellow fill on my own assumptions, green font for cross-sheet links, black for formulas. Core (non-GAAP) basis.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from data import *

NAVY = "1F3864"; MUTED = "595959"
F_TTL = Font(bold=True, color=NAVY, size=12)
F_SUB = Font(color=MUTED, size=9)
F_HDR = Font(bold=True, color="FFFFFF", size=10)
F_IN = Font(color="0000FF")
F_INB = Font(color="0000FF", bold=True)
F_LINK = Font(color="008000")
F_B = Font(bold=True)
F_NOTE = Font(color=MUTED, size=9, italic=True)
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_ASSUME = PatternFill("solid", fgColor="FFFF00")
FILL_KEY = PatternFill("solid", fgColor="D9E2F3")
WRAP = Alignment(wrap_text=True, vertical="top")
RIGHT = Alignment(horizontal="right")

USD2 = '"US$"#,##0.00'; USD0 = '"US$"#,##0'; BN2 = '#,##0.00'; BN3 = '#,##0.000'; BN1 = '#,##0.0'
PCT1 = '0.0%'; PCT0 = '0%'; MULT = '0.0"x"'; MULT0 = '0"x"'
SUB = f"Tommy Lau | Corning (NYSE: GLW) | {DATE_LONG} | Core (non-GAAP) basis | Personal research, not investment advice."

wb = openpyxl.Workbook()


def sheet(name, title, widths):
    ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet(name)
    ws.title = name
    ws["A1"] = title; ws["A1"].font = F_TTL
    ws["A2"] = SUB; ws["A2"].font = F_SUB
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def head(ws, row, cells, start=1):
    for i, c in enumerate(cells, start):
        x = ws.cell(row=row, column=i, value=c)
        x.font = F_HDR; x.fill = FILL_HDR
        x.alignment = Alignment(horizontal="left" if i == start else "right", vertical="center", wrap_text=True)


def note(ws, row, texts):
    for i, t in enumerate(texts):
        ws.cell(row=row + i, column=1, value=t).font = F_NOTE


sheet("Cover", "CORNING (NYSE: GLW)  INITIATION MODEL", [38, 24, 80])

# ===================================================================== ASSUMPTIONS
ws = sheet("Assumptions", "ASSUMPTIONS AND INPUTS", [52, 14, 10, 92])
ws["A3"] = ("Blue = hardcoded input. Yellow fill = my own assumption (change these to move the value). Green = link to "
            "another tab. Black = formula. US$ billion unless stated; per-share figures in US$. Core (non-GAAP) basis.")
ws["A3"].font = F_NOTE
head(ws, 5, ["Input", "Value", "Unit", "Source / note"])
A = {}
MINE = {"q3opt", "q4opt", "q3rest", "q4rest", "q3om", "q4om", "restm26", "corpq", "corp27", "corp28", "atm27", "atm28", "pe"}
rows = [
    ("MARKET (retrieved 6 October 2026)", None, None, None, None),
    ("price", "Share price (NYSE: GLW)", PRICE, "US$", "NYSE close, 5 October 2026. Yahoo Finance."),
    ("hi52", "52-week high (intraday, 29 June 2026)", HI52, "US$", "Yahoo Finance."),
    ("px25", "Month-end close, 30 Sep 2025", PRICE_SEP25, "US$", "Yahoo Finance."),
    ("px24", "Month-end close, 31 Dec 2024", PRICE_DEC24, "US$", "Yahoo Finance."),
    ("atmpre", "Close on 11 Sep 2026", ATM_PRE, "US$", "Day the at-the-market programme was filed. stockanalysis.com daily history."),
    ("atmpost", "Close on 14 Sep 2026", ATM_POST, "US$", "Next trading day."),
    ("shout", "Shares outstanding, 24 Jul 2026", SHARES_OUT, "bn", "Prospectus supplement of 11 Sep 2026: 861,388,331."),
    ("diluted", "Diluted shares, Q2 2026", DILUTED, "bn", "Q2 2026 release: 875 million, includes the 3m pre-funded NVIDIA warrant."),
    ("cash", "Cash, 30 Jun 2026", CASH, "US$bn", "Q2 2026 release."),
    ("debt", "Total debt, 30 Jun 2026", DEBT, "US$bn", "Q2 2026 release: long-term 7.756 + current 0.668."),
    ("debt1y", "Debt due within a year, 30 Jun 2026", DEBT_DUE_1Y, "US$bn", "Q2 2026 release and 10-Q."),
    ("dps", "Quarterly dividend", DPS_Q, "US$", "10-Q: US$0.28 a quarter."),
    ("cons26", "2026 consensus core EPS", CONS_EPS_26, "US$", "Nasdaq.com (Zacks), 8 estimates, retrieved 6 October 2026."),
    ("cons27", "2027 consensus core EPS", CONS_EPS_27, "US$", "Nasdaq.com (Zacks), 8 estimates."),
    ("cons28", "2028 consensus core EPS", CONS_EPS_28, "US$", "Nasdaq.com (Zacks), 4 estimates."),
    ("fwdpe", "Forward P/E shown by stockanalysis.com", FWD_PE_SA, "x", "stockanalysis.com, 6 October 2026."),
    ("trailpe", "Trailing P/E (GAAP) shown by stockanalysis.com", TRAIL_PE_SA, "x", "stockanalysis.com, 6 October 2026."),
    ("constp", "Average analyst target", CONS_TP, "US$", "stockanalysis.com, 17 analysts, 6 October 2026."),
    ("REPORTED (Corning releases, Form 10-K, Form 10-Q, 8-Ks)", None, None, None, None),
    ("h1sales", "Core sales, H1 2026", H1_26_SALES, "US$bn", "Q2 2026 release: 4.345 + 4.738."),
    ("h1opt", "Optical Communications sales, H1 2026", H1_26_OPT, "US$bn", "10-Q segment note."),
    ("h1optni", "Optical Communications net income, H1 2026", H1_26_OPT_NI, "US$bn", "10-Q segment note."),
    ("h1segni", "Segment net income incl. Life Sciences and Emerging, H1 2026", H1_26_SEGNI, "US$bn", "10-Q segment note."),
    ("h1ni", "Core net income, H1 2026", H1_26_CORE_NI, "US$bn", "Q2 2026 release."),
    ("h1eps", "Core EPS, H1 2026", H1_26_EPS, "US$", "Q2 2026 release."),
    ("g3lo", "Q3 2026 core sales guide, low", Q3_GUIDE_SALES[0], "US$bn", "Release of 28 July 2026."),
    ("g3hi", "Q3 2026 core sales guide, high", Q3_GUIDE_SALES[1], "US$bn", "Same."),
    ("g3elo", "Q3 2026 core EPS guide, low", Q3_GUIDE_EPS[0], "US$", "Same."),
    ("g3ehi", "Q3 2026 core EPS guide, high", Q3_GUIDE_EPS[1], "US$", "Same."),
    ("ent", "Enterprise Networks sales, Q2 2026", ENT_Q2_26, "US$bn", "Q2 2026 call: US$1.27bn, up 65%."),
    ("entg", "Enterprise Networks growth, Q2 2026", ENT_Q2_26_G, "%", "Q2 2026 release and call."),
    ("atm", "At-the-market programme, maximum", ATM, "US$bn", "8-K and prospectus supplement, 11 Sep 2026. Goldman Sachs; 1% commission."),
    ("capex26", "2026 capital spending outlook", CAPEX_GUIDE_26, "US$bn", "10-Q for Q2 2026: approximately US$2.0bn."),
    ("h1capex", "Capital spending, H1 2026", H1_26_CAPEX, "US$bn", "Q2 2026 release."),
    ("h1fcf", "Adjusted free cash flow, H1 2026", H1_26_FCF, "US$bn", "Q2 2026 release (Corning definition)."),
    ("h1dep", "Customer deposits and government incentives, H1 2026", H1_26_DEP_INFLOW, "US$bn", "Cash flow statement line; includes the US$1.0bn deposit received in Q2."),
    ("h1div", "Dividends paid, H1 2026", H1_26_DIV, "US$bn", "Q2 2026 release."),
    ("deposit", "Customer deposit booked in Q2 2026", DEPOSIT_Q2, "US$bn", "10-Q Note 2: long-term supply agreement to 31 Dec 2029."),
    ("nvfv", "Traditional warrant to a customer, grant-date fair value", NV_WARRANT_FV, "US$bn", "10-Q Note 12: 15m shares at US$180, consideration payable to a customer."),
    ("nvpf", "Cash for the pre-funded warrant", NV_PF_CASH, "US$bn", "8-K of 6 May 2026: NVIDIA, 3m shares."),
    ("MY ASSUMPTIONS", None, None, None, None),
    ("q3opt", "Optical sales, Q3 2026", Q3E_OPT, "US$bn", "MINE. About 7% on Q2."),
    ("q4opt", "Optical sales, Q4 2026", Q4E_OPT, "US$bn", "MINE."),
    ("q3rest", "Rest of Corning sales, Q3 2026", Q3E_REST, "US$bn", "MINE. Q3 total 4.95 = guide midpoint."),
    ("q4rest", "Rest of Corning sales, Q4 2026", Q4E_REST, "US$bn", "MINE. Q4 total about US$5.1bn, the US$20bn run-rate."),
    ("q3om", "Optical segment net margin, Q3 2026", Q3E_OPT_M, "%", "MINE. Q2 2026: 21.1%."),
    ("q4om", "Optical segment net margin, Q4 2026", Q4E_OPT_M, "%", "MINE."),
    ("restm26", "Rest of Corning segment net margin, H2 2026", REST_M_26, "%", "MINE. H1 2026: 15.2%; Solar profitability expected to improve from Q3."),
    ("corpq", "Core net income less segment net income, per quarter", CORP_Q, "US$bn", "MINE. Q2 2026: -0.166; H1: -0.318."),
    ("corp27", "Corporate line, 2027", CORP27, "US$bn", "MINE."),
    ("corp28", "Corporate line, 2028", CORP28, "US$bn", "MINE."),
    ("atm27", "Share of the programme sold, 2027 average", ATM_SOLD27, "%", "MINE. Sold at today's price."),
    ("atm28", "Share of the programme sold, 2028 average", ATM_SOLD28, "%", "MINE."),
    ("pe", "Target multiple on 2028E core EPS (base)", BASE_PE, "x", "MINE. Between the fibre makers' median (Peers tab) and Amphenol."),
]
r = 6
for key, label, val, unit, src in rows:
    if label is None:
        ws.cell(row=r, column=1, value=key).font = Font(bold=True, color=NAVY); r += 1; continue
    ws.cell(row=r, column=1, value=label)
    c = ws.cell(row=r, column=2, value=val)
    mine = key in MINE
    c.font = F_INB if mine else F_IN
    if mine: c.fill = FILL_ASSUME
    c.number_format = PCT1 if unit == "%" else (MULT0 if unit == "x" else (BN3 if unit == "bn" else "#,##0.000"))
    ws.cell(row=r, column=3, value=unit)
    ws.cell(row=r, column=4, value=src).alignment = WRAP
    A[key] = f"Assumptions!$B${r}"
    r += 1

# ===================================================================== QUARTERLY
ws = sheet("Quarterly", "QUARTERLY HISTORY, Q1 2024 TO Q2 2026, AND MY Q3 / Q4 2026 (CORE BASIS)", [46] + [10] * 12)
QH = Q + ["Q3 26E", "Q4 26E"]
head(ws, 4, ["US$ billion unless stated"] + QH)
C = [get_column_letter(2 + j) for j in range(12)]


def qrow(row, label, vals, fmt, bold=False):
    ws.cell(row=row, column=1, value=label).font = F_B if bold else Font()
    for j, v in enumerate(vals):
        c = ws.cell(row=row, column=2 + j, value=v if v is not None else "")
        c.number_format = fmt; c.alignment = RIGHT
        if v is not None and not (isinstance(v, str) and v.startswith("=")):
            c.font = F_IN


bn = lambda xs: [x / 1000 for x in xs]
qrow(5, "Optical Communications sales", bn(Q_OPT) + ["=" + A["q3opt"], "=" + A["q4opt"]], BN3, True)
qrow(6, "Rest of Corning sales (derived)", [f"={c}7-{c}5" for c in C[:10]] + ["=" + A["q3rest"], "=" + A["q4rest"]], BN3)
qrow(7, "Core sales", bn(Q_SALES) + [f"={C[10]}5+{C[10]}6", f"={C[11]}5+{C[11]}6"], BN3, True)
qrow(8, "Optical growth y/y (derived)", [None] * 4 + [f"={C[j]}5/{C[j-4]}5-1" for j in range(4, 12)], PCT1)
qrow(9, "Optical share of core sales (derived)", [f"={c}5/{c}7" for c in C], PCT1)
qrow(10, "Optical segment net income", bn(Q_OPT_NI) + [f"={C[10]}5*{A['q3om']}", f"={C[11]}5*{A['q4om']}"], BN3)
qrow(11, "Optical segment net margin (derived)", [f"={c}10/{c}5" for c in C], PCT1, True)
qrow(12, "Core EPS, US$ (reported; Q3 / Q4 26 mine)", Q_EPS + [
    f"=({C[10]}10+{C[10]}6*{A['restm26']}+{A['corpq']})/{A['diluted']}",
    f"=({C[11]}10+{C[11]}6*{A['restm26']}+{A['corpq']})/{A['diluted']}"], USD2, True)
ws["A14"] = "Checks"; ws["A14"].font = Font(bold=True, color=NAVY)
ws["A15"] = "Trailing four quarters core EPS to Q2 26"; ws["B15"] = "=SUM(H12:K12)"; ws["B15"].number_format = USD2
ws["A16"] = "Q3 26 core EPS guide midpoint"; ws["B16"] = f"=({A['g3elo']}+{A['g3ehi']})/2"; ws["B16"].number_format = USD2
ws["A17"] = "Q3 26 core sales guide midpoint"; ws["B17"] = f"=({A['g3lo']}+{A['g3hi']})/2"; ws["B17"].number_format = BN2
ws["A18"] = "Q4 26E annualised sales run-rate (Springboard: US$20bn by end 2026)"; ws["B18"] = "=M7*4"; ws["B18"].number_format = BN1
ws["A19"] = "Enterprise share of optical, Q2 26"; ws["B19"] = f"={A['ent']}/K5"; ws["B19"].number_format = PCT1
ws["A20"] = "Carrier sales, Q2 26 (derived)"; ws["B20"] = f"=K5-{A['ent']}"; ws["B20"].number_format = BN3
ws["A21"] = "Carrier sales, Q2 25 (derived)"; ws["B21"] = f"=G5-{A['ent']}/(1+{A['entg']})"; ws["B21"].number_format = BN3
note(ws, 23, ["Source: Corning quarterly earnings releases (Exhibit 99 to Form 8-K), 30 April 2024 to 28 July 2026. Core measures are non-GAAP.",
              "Q3 26E and Q4 26E are my estimates; EPS = (optical net income + rest x margin + corporate line) / diluted shares."])

# ===================================================================== FINANCIALS
ws = sheet("Financials", "FINANCIAL SUMMARY AND FORECASTS (US$ billion unless stated, core basis)", [50, 12, 12, 12, 12, 12, 12])
head(ws, 4, ["", "2023A", "2024A", "2025A", "2026E*", "2027E*", "2028E*"])
b_ = SCEN[1]
fin = [
    ("Optical Communications sales", [OPT_H[0], OPT_H[1], OPT_H[2], "=SUM(Quarterly!J5:M5)", "=E5*(1+Scenarios!B6)", "=F5*(1+Scenarios!C6)"], BN2, False),
    ("Rest of Corning sales", ["=B7-B5", "=C7-C5", "=D7-D5", "=SUM(Quarterly!J6:M6)", "=E6*(1+Scenarios!F6)", "=F6*(1+Scenarios!G6)"], BN2, False),
    ("Core sales", [SALES_H[0], SALES_H[1], SALES_H[2], "=E5+E6", "=F5+F6", "=G5+G6"], BN2, True),
    ("Core sales growth", [None, "=C7/B7-1", "=D7/C7-1", "=E7/D7-1", "=F7/E7-1", "=G7/F7-1"], PCT1, False),
    ("Optical share of sales", ["=B5/B7", "=C5/C7", "=D5/D7", "=E5/E7", "=F5/F7", "=G5/G7"], PCT1, False),
    ("Optical segment net income", [OPT_NI_H[0], OPT_NI_H[1], OPT_NI_H[2], "=SUM(Quarterly!J10:M10)", "=F5*Scenarios!D6", "=G5*Scenarios!E6"], BN2, False),
    ("Optical segment net margin", ["=B10/B5", "=C10/C5", "=D10/D5", "=E10/E5", "=F10/F5", "=G10/G5"], PCT1, True),
    ("Rest of Corning segment net income", [None, None, None, f"=({A['h1segni']}-{A['h1optni']})+(Quarterly!L6+Quarterly!M6)*{A['restm26']}", "=F6*Scenarios!H6", "=G6*Scenarios!I6"], BN2, False),
    ("Corporate line (core net income less segments)", [None, None, None, f"=({A['h1ni']}-{A['h1segni']})+2*{A['corpq']}", "=" + A["corp27"], "=" + A["corp28"]], BN2, False),
    ("Core net income", [CORE_NI_H[0], CORE_NI_H[1], CORE_NI_H[2], "=E10+E12+E13", "=F10+F12+F13", "=G10+G12+G13"], BN2, True),
    ("Core net margin", ["=B14/B7", "=C14/C7", "=D14/D7", "=E14/E7", "=F14/F7", "=G14/G7"], PCT1, False),
    ("Diluted shares, bn", [None, None, None, "=" + A["diluted"], f"={A['diluted']}+{A['atm']}/{A['price']}*{A['atm27']}", f"={A['diluted']}+{A['atm']}/{A['price']}*{A['atm28']}"], BN3, False),
    ("Core EPS, US$", [EPS_H[0], EPS_H[1], EPS_H[2], f"={A['h1eps']}+Quarterly!L12+Quarterly!M12", "=F14/F16", "=G14/G16"], USD2, True),
    ("Core EPS growth", [None, "=C17/B17-1", "=D17/C17-1", "=E17/D17-1", "=F17/E17-1", "=G17/F17-1"], PCT1, False),
    ("P/E at reference price", [f"={A['price']}/{c}17" for c in "BCDEFG"], MULT, False),
    ("Consensus core EPS, US$", [None, None, None, "=" + A["cons26"], "=" + A["cons27"], "=" + A["cons28"]], USD2, False),
    ("Mine against consensus", [None, None, None, "=E17/E20-1", "=F17/F20-1", "=G17/G20-1"], PCT1, False),
    ("Core operating margin (reported)", [OPM_H[0], OPM_H[1], OPM_H[2], None, None, None], PCT1, False),
    ("Capital expenditures", [CAPEX_H[0], CAPEX_H[1], CAPEX_H[2], "=" + A["capex26"], None, None], BN2, False),
    ("Adjusted free cash flow (Corning definition)", [FCF_H[0], FCF_H[1], FCF_H[2], None, None, None], BN2, False),
    ("Of which customer deposits and government incentives", [DEP_FLOW_H[0], DEP_FLOW_H[1], DEP_FLOW_H[2], None, None, None], BN2, False),
    ("Dividends paid", [DIV_H[0], DIV_H[1], DIV_H[2], None, None, None], BN2, False),
]
for i, (label, vals, fmt, bold) in enumerate(fin):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=label).font = F_B if bold else Font()
    for j, v in enumerate(vals):
        c = ws.cell(row=rr, column=2 + j, value=v if v is not None else "")
        c.number_format = fmt; c.alignment = RIGHT
        if v is not None and not isinstance(v, str): c.font = F_IN
        elif isinstance(v, str) and v.startswith("=Assumptions"): c.font = F_LINK
note(ws, 28, [
    "*2026E to 2028E are my own estimates. 2026E = reported H1 plus my Q3 and Q4 (Quarterly tab). 2027E and 2028E use the base case on the Scenarios tab (row 6).",
    "Core net income = optical segment net income + rest of Corning segment net income + corporate line. History from Corning's fourth quarter releases and Form 10-K.",
    f"Note check: 2026E sales {SALES26:.2f}, EPS {EPS26:.2f}; 2027E {S27:.2f}, {EPS27:.2f}; 2028E {S28:.2f}, {EPS28:.2f}.",
])

# ===================================================================== SCENARIOS
ws = sheet("Scenarios", "THREE CASES, TWELVE MONTHS OUT, AND WHICH DRIVER MATTERS", [30] + [11] * 15)
HDR = ["Case", "Opt g 27", "Opt g 28", "Opt m 27", "Opt m 28", "Rest g 27", "Rest g 28", "Rest m 27", "Rest m 28",
       "2028 optical", "2028 sales", "2028 core NI", "2028 EPS", "Multiple", "Value, US$", "vs price"]
head(ws, 4, HDR)
O26 = "Financials!$E$5"; R26 = "Financials!$E$6"


def scen_row(rr, name, vals, mult):
    ws.cell(row=rr, column=1, value=name).font = F_B
    for j, v in enumerate(vals):
        if isinstance(v, str):
            c = ws.cell(row=rr, column=2 + j, value=v)
        else:
            c = ws.cell(row=rr, column=2 + j, value=v); c.font = F_INB; c.fill = FILL_ASSUME
        c.number_format = PCT1
    ws.cell(row=rr, column=10, value=f"={O26}*(1+B{rr})*(1+C{rr})").number_format = BN2
    ws.cell(row=rr, column=11, value=f"=J{rr}+{R26}*(1+F{rr})*(1+G{rr})").number_format = BN2
    ws.cell(row=rr, column=12, value=f"=J{rr}*E{rr}+(K{rr}-J{rr})*I{rr}+{A['corp28']}").number_format = BN2
    ws.cell(row=rr, column=13, value=f"=L{rr}/Financials!$G$16").number_format = USD2
    c = ws.cell(row=rr, column=14, value=("=" + A["pe"]) if mult == "pe" else mult)
    if mult == "pe": c.font = F_LINK
    else: c.font = F_INB; c.fill = FILL_ASSUME
    c.number_format = MULT0
    ws.cell(row=rr, column=15, value=f"=M{rr}*N{rr}").number_format = USD2
    ws.cell(row=rr, column=16, value=f"=O{rr}/{A['price']}-1").number_format = PCT0


for i, s in enumerate(SCEN):
    scen_row(5 + i, s[0], list(s[1:9]), "pe" if s[0] == "Base" else s[9])
ws["A10"] = "Probability"; ws["A10"].font = F_B
for j, s in enumerate(SCEN):
    ws.cell(row=9, column=2 + j, value=s[0]).font = F_B
    c = ws.cell(row=10, column=2 + j, value=s[10]); c.font = F_INB; c.fill = FILL_ASSUME; c.number_format = PCT0
ws["E10"] = "=SUM(B10:D10)"; ws["E10"].number_format = PCT0; ws["F10"] = "must be 100%"; ws["F10"].font = F_NOTE
ws["A12"] = "Probability-weighted value, US$"; ws["A12"].font = F_B
ws["O12"] = "=B10*O5+C10*O6+D10*O7"; ws["O12"].number_format = USD2; ws["O12"].font = F_B; ws["O12"].fill = FILL_KEY
ws["P12"] = f"=O12/{A['price']}-1"; ws["P12"].number_format = PCT1
ws["A14"] = "Which driver matters: swing one part from bear to bull, the other held at base"; ws["A14"].font = Font(bold=True, color=NAVY)
head(ws, 15, HDR)
scen_row(16, "Optical bear, rest base", ["=B5", "=C5", "=D5", "=E5", "=F6", "=G6", "=H6", "=I6"], "pe")
scen_row(17, "Optical bull, rest base", ["=B7", "=C7", "=D7", "=E7", "=F6", "=G6", "=H6", "=I6"], "pe")
scen_row(18, "Rest bear, optical base", ["=B6", "=C6", "=D6", "=E6", "=F5", "=G5", "=H5", "=I5"], "pe")
scen_row(19, "Rest bull, optical base", ["=B6", "=C6", "=D6", "=E6", "=F7", "=G7", "=H7", "=I7"], "pe")
ws["A21"] = "2028 EPS swing from optical, US$"; ws["M21"] = "=M17-M16"; ws["M21"].number_format = USD2
ws["A22"] = "2028 EPS swing from the rest of Corning, US$"; ws["M22"] = "=M19-M18"; ws["M22"].number_format = USD2
ws["A23"] = "Ratio, optical to rest"; ws["M23"] = "=M21/M22"; ws["M23"].number_format = '0.0"x"'
note(ws, 25, ["All cases start from my 2026E and use 2028 diluted shares from the Financials tab. Margins are segment net margins.",
              "Bear: new fibre capacity (Corning's and China's) catches up, prices and margins slip. Bull: the new plants fill early at rising margins.",
              "The base row (6) drives the Financials tab for 2027E and 2028E."])

# ===================================================================== VALUATION
ws = sheet("Valuation", "VALUATION: 27x 2028E CORE EPS, TWELVE MONTHS OUT", [60, 16, 70])
head(ws, 4, ["Line", "Value", "Working / note"])
val = [
    ("eps28", "2028E core EPS, US$", "=Financials!G17", USD2, "The market will be pricing 2028 by October 2027."),
    ("mult", "Base multiple, x", "=" + A["pe"], MULT0, "Assumptions tab."),
    ("base", "Base-case value per share, US$", "=B5*B6", USD2, "EPS x multiple."),
    ("px", "Reference price, US$", "=" + A["price"], USD2, "NYSE close, 5 October 2026."),
    ("upb", "Base value against the price", "=B7/B8-1", PCT1, "Within +/- 15%: no call on value alone."),
    ("wtd", "Probability-weighted value, US$", "=Scenarios!O12", USD2, "Scenarios tab."),
    ("call", "CALL", '=IF(B9>0.15,"LONG",IF(B9<-0.15,"SHORT","NO CALL"))', None, "Rule of thumb: 15% either side of the price. Conviction is judgement. Tommy Lau's call, 6 Oct 2026."),
    (None, "", None, None, ""),
    ("mcap", "Market value on shares outstanding, US$bn", f"={A['price']}*{A['shout']}", BN1, ""),
    ("nd", "Net debt, US$bn", f"={A['debt']}-{A['cash']}", BN2, "30 June 2026."),
    ("ev", "Enterprise value, US$bn", "=B13+B14", BN1, ""),
    ("dy", "Dividend yield", f"=4*{A['dps']}/{A['price']}", '0.00%', "US$1.12 a year."),
    ("rise", "Change since 30 Sep 2025", f"={A['price']}/{A['px25']}-1", PCT1, ""),
    ("hi", "Price below the 52-week high", f"=1-{A['price']}/{A['hi52']}", PCT1, ""),
    (None, "", None, None, ""),
    ("pe26", "P/E on my 2026E", "=Financials!E19", MULT, ""),
    ("pe27", "P/E on my 2027E", "=Financials!F19", MULT, ""),
    ("pe28", "P/E on my 2028E", "=Financials!G19", MULT, ""),
    ("pec27", "P/E on the 2027 consensus", f"={A['price']}/{A['cons27']}", MULT, "Nasdaq.com (Zacks) US$4.29."),
    ("pettm", "P/E on trailing four quarters core EPS", f"={A['price']}/Quarterly!B15", MULT, "GAAP trailing: 73.5x (stockanalysis.com)."),
    ("pehist", "End-2024 price on 2025 core EPS, in hindsight", f"={A['px24']}/Financials!D17", MULT, "What the market paid for Corning before the re-rating."),
    (None, "", None, None, ""),
    ("need", "2028 EPS the price needs at the base multiple, US$", f"={A['price']}/{A['pe']}", USD2, "What the market prices in."),
    ("needni", "2028 optical net income that needs, rest at base, US$bn", f"=B27*Financials!G16-Financials!G12-Financials!G13", BN2, ""),
    ("needopt", "2028 optical sales that needs at the base margin, US$bn", "=B28/Scenarios!E6", BN2, "Against my US$14.34bn."),
    ("needx", "That against 2025 optical sales", "=B29/Financials!D5", '0.00"x"', "2025: US$6.27bn."),
    ("needg", "Implied optical growth a year, 2026E to 2028", "=(B29/Financials!E5)^0.5-1", PCT1, ""),
    ("revisit", "Price that puts the base case 15% above", "=B7/1.15", USD2, "The 'revisit below about US$127' line."),
]
V = {}
for i, (key, label, f, fmt, nt) in enumerate(val):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=label)
    if f is not None:
        c = ws.cell(row=rr, column=2, value=f)
        if fmt: c.number_format = fmt
        if key in ("base", "call"):
            ws.cell(row=rr, column=1).font = F_B; c.font = F_B; c.fill = FILL_KEY
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
    if key: V[key] = f"Valuation!$B${rr}"

# ===================================================================== FUNDING (the share sale question)
ws = sheet("Funding", "WHY A SOLD-OUT BUSINESS SET UP A US$2BN SHARE SALE: THE CASH ARITHMETIC", [64, 16, 70])
head(ws, 4, ["Line", "Value", "Working / note"])
fund = [
    ("Capital spending, 2025", f"=Financials!D23", BN2, "Corning Q4 2025 release."),
    ("Capital spending, 2026 outlook", "=" + A["capex26"], BN2, "10-Q: approximately US$2.0bn."),
    ("Step-up, 2026 on 2025", "=B6/B5-1", PCT0, ""),
    ("Capital spending, H1 2026", "=" + A["h1capex"], BN2, ""),
    ("Implied capital spending, H2 2026", "=B6-B8", BN2, "Second half spends more than the first."),
    ("Adjusted free cash flow, H1 2026", "=" + A["h1fcf"], BN2, ""),
    ("Less customer deposits and government incentives, H1 2026", "=" + A["h1dep"], BN2, "Includes the US$1.0bn deposit received in Q2."),
    ("Adjusted free cash flow without them, H1 2026", "=B10-B11", BN2, ""),
    ("Dividends paid, H1 2026", "=" + A["h1div"], BN2, ""),
    ("Debt due within a year, 30 Jun 2026", "=" + A["debt1y"], BN2, ""),
    ("Cash, 30 Jun 2026", "=" + A["cash"], BN2, "Plus an undrawn US$1.5bn revolving credit facility."),
    (None, None, None, ""),
    ("At-the-market programme, maximum, US$bn", "=" + A["atm"], BN2, "11 Sep 2026."),
    ("Shares if fully sold at the reference price, bn", f"={A['atm']}/{A['price']}", BN3, ""),
    ("Dilution against shares outstanding", f"=B18/{A['shout']}", '0.00%', ""),
    ("Share price change, 11 to 14 Sep 2026", f"={A['atmpost']}/{A['atmpre']}-1", PCT1, ""),
    (None, None, None, ""),
    ("Customer deposit booked in Q2 2026, US$bn", "=" + A["deposit"], BN2, "Long-term supply agreement to 31 Dec 2029."),
    ("Warrant to the same customer, fair value, US$bn", "=" + A["nvfv"], BN3, "Comes off revenue as Corning delivers."),
    ("Warrant value per dollar of deposit", "=B23/B22", PCT0, "About 30 cents."),
]
for i, (label, f, fmt, nt) in enumerate(fund):
    rr = 5 + i
    if label is None: continue
    ws.cell(row=rr, column=1, value=label)
    c = ws.cell(row=rr, column=2, value=f); c.number_format = fmt
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
note(ws, 27, ["Corning has not said why it set up the programme; the prospectus gives general corporate purposes, including capital spending, debt repayment and buybacks.",
              "Also relevant: Samsung Display holds 58m Corning shares under a lock-up expiring in 2027, and can offer 22m more to Corning in tranches through 2027 (10-K, 10-Q)."])

# ===================================================================== SENSITIVITY
ws = sheet("Sensitivity", "VALUE PER SHARE: 2028E CORE EPS AGAINST MULTIPLE (US$)", [24, 14, 14, 14])
head(ws, 4, ["2028 EPS \\ multiple", "", "", ""])
for j, (pe, link) in enumerate(zip(SENS_PE, [False, True, False])):
    c = ws.cell(row=4, column=2 + j, value=("=" + A["pe"]) if link else pe)
    c.number_format = MULT0; c.font = Font(bold=True, color="FFFFFF"); c.fill = FILL_HDR; c.alignment = RIGHT
for i, (e, link) in enumerate(zip(SENS_EPS, [False, True, False])):
    rr = 5 + i
    c = ws.cell(row=rr, column=1, value="=Financials!G17" if link else e)
    c.number_format = USD2; c.font = F_LINK if link else F_IN
    for j in range(3):
        col = get_column_letter(2 + j)
        ws.cell(row=rr, column=2 + j, value=f"=$A{rr}*{col}$4").number_format = USD0
ws["C6"].fill = FILL_KEY
ws["A9"] = "Cells above the reference price"; ws["A9"].font = F_B
ws["B9"] = f'=COUNTIF(B5:D7,">"&{A["price"]})'; ws["C9"] = "of 9"
note(ws, 11, ["Middle row links to the model's 2028E EPS and the middle column to the base multiple."])

# ===================================================================== PEERS
ws = sheet("Peers", "PEER MULTIPLES (FORWARD P/E, stockanalysis.com, 6 OCTOBER 2026)", [26, 16, 14, 80])
head(ws, 4, ["Company", "Ticker", "Forward P/E", "What they make"])
for i, (co, tk, pe, what) in enumerate(PEERS):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=co).font = F_B if co == NAME else Font()
    ws.cell(row=rr, column=2, value=tk)
    c = ws.cell(row=rr, column=3, value=("=" + A["fwdpe"]) if co == NAME else pe)
    c.number_format = MULT; c.font = F_LINK if co == NAME else F_IN
    ws.cell(row=rr, column=4, value=what).alignment = WRAP
ws["A14"] = "Median, fibre and cable makers (Prysmian, Fujikura, Sumitomo Electric)"; ws["C14"] = "=MEDIAN(C6:C8)"; ws["C14"].number_format = MULT
ws["A15"] = "Median, all seven peers"; ws["C15"] = "=MEDIAN(C6:C12)"; ws["C15"].number_format = MULT
note(ws, 17, ["Prysmian, Fujikura and Sumitomo Electric multiples are on their local listings. Multiples only, to set the base multiple.",
              "No view on the peers' shares is expressed."])

# ===================================================================== SOURCES
ws = sheet("Sources", "SOURCES", [6, 130])
head(ws, 4, ["No.", "Source and what it supports"])
srcs = [
    "Corning quarterly earnings releases (Exhibit 99 to Form 8-K), 30 April 2024 to 28 July 2026: core sales, core EPS, segment sales and net income, margins, cash flow, capital spending, adjusted free cash flow, Q3 2026 guide.",
    "Corning Form 10-K for 2025, filed 12 February 2026: segment sales, enterprise and carrier network sales, optical at 38% of segment sales, two end customers at 28% of optical sales, the Samsung Display share repurchase agreement.",
    "Corning Form 10-Q for Q2 2026, filed 29 July 2026: US$1.0bn customer deposit to 31 Dec 2029, warrant accounting (US$296m), contract liabilities, liquidity, credit facility, leverage covenant, 2026 capital spending of about US$2.0bn.",
    "Corning 8-K and press release of 6 May 2026: NVIDIA partnership, three new plants, 10x US connectivity capacity; pre-funded warrant (3m shares, US$500m) and traditional warrant (15m shares at US$180).",
    "Corning 8-K and prospectus supplement of 11 September 2026: US$2bn at-the-market programme with Goldman Sachs; use of proceeds; 861,388,331 shares outstanding at 24 July 2026.",
    "Corning Q2 2026 earnings call of 28 July 2026 (Investing.com transcript summary): Enterprise Networks US$1.27bn, capacity commentary, long-term contracts, second-half capital spending.",
    "Meta (27 January 2026), Amazon (8 June 2026) and Verizon (8 September 2026) agreement announcements; Corning Q1 2026 release (two further hyperscale agreements).",
    "Corning investor relations events feed: Q3 2026 results call on Tuesday 27 October 2026, 8:30am ET.",
    "Yahoo Finance chart API, retrieved 6 October 2026: month-end closes, the 5 October close, 52-week range. stockanalysis.com daily history for 11 and 14 September 2026.",
    "Nasdaq.com earnings forecast (Zacks), retrieved 6 October 2026: consensus core EPS 2026 to 2028. stockanalysis.com: forward and trailing P/E, average target, peer multiples.",
    "The Physical Layer, 'Corning makes the glass AI chips talk through', 1 October 2026: https://thephysicallayer.fyi/journal/corning-lays-the-glass/",
    "Estimates for 2026 to 2028 are the author's own. Personal research, not investment advice. I hold no position in Corning.",
]
for i, s in enumerate(srcs):
    ws.cell(row=5 + i, column=1, value=i + 1)
    ws.cell(row=5 + i, column=2, value=s).alignment = WRAP

# ===================================================================== COVER
ws = wb["Cover"]
head(ws, 4, ["Item", "Value", "Note"])
cover = [
    ("Call", "=" + V["call"], None, "Tommy Lau's call, published 6 Oct 2026."),
    ("Conviction", CALL["conviction"], None, "Judgement, not formula."),
    ("Reference price, US$", "=" + A["price"], USD2, "NYSE close, 5 October 2026."),
    ("Base value, US$", "=" + V["base"], USD2, "27x 2028E core EPS."),
    ("Base value against the price", "=" + V["upb"], PCT1, ""),
    ("Probability-weighted value, US$", "=" + V["wtd"], USD2, "25 / 50 / 25 bear / base / bull."),
    ("2026E / 2027E / 2028E core EPS, US$", '=TEXT(Financials!E17,"0.00")&" / "&TEXT(Financials!F17,"0.00")&" / "&TEXT(Financials!G17,"0.00")', None, "Own estimates."),
    ("Horizon", "12 months, to October 2027", None, ""),
    ("Wrong if", "See note", None, "Optical Communications sales of US$11.5bn or more in 2027 with a segment net margin of 24% or more; or the shares close above US$200 or below US$110 by October 2027."),
    ("Revisit if", "See note", None, "A price below about US$127; or Q4 2026 optical sales of US$2.4bn or more with a segment net margin of 23% or more, and the share sale programme largely unused."),
]
for i, (lab, v, fmt, nt) in enumerate(cover):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=lab).font = F_B
    c = ws.cell(row=rr, column=2, value=v)
    if fmt: c.number_format = fmt
    c.alignment = RIGHT
    c.font = F_LINK if isinstance(v, str) and v.startswith("=") else F_IN
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
ws["A17"] = "Tabs"; ws["A17"].font = Font(bold=True, color=NAVY)
for i, (t, d) in enumerate([
    ("Assumptions", "Every input, with its source. Change the yellow cells to move the value."),
    ("Quarterly", "Q1 2024 to Q2 2026 reported, my Q3 and Q4 2026; optical, the rest, margins, core EPS."),
    ("Financials", "2023A to 2028E: optical and the rest, segment margins, core net income, EPS, P/E, cash."),
    ("Scenarios", "Bear / base / bull, the probability-weighted value, and which part of Corning moves earnings most."),
    ("Valuation", "Base value, market value, multiples, what the price needs."),
    ("Funding", "The share sale question: capital spending, deposits, free cash flow, dilution."),
    ("Sensitivity", "2028E core EPS against multiple."),
    ("Peers", "Forward multiples used to set the base multiple."),
    ("Sources", "Every source used."),
]):
    ws.cell(row=18 + i, column=1, value=t); ws.cell(row=18 + i, column=2, value=d)
ws["A28"] = "Colour key: blue = hardcoded input; yellow fill = my own assumption; green = link to another tab; black = formula."
ws["A28"].font = F_NOTE
ws["A29"] = f"Companion note: {DATE}_{COMPANY}_Initiation.pdf. Personal research, not investment advice."; ws["A29"].font = F_NOTE

for w in wb.worksheets:
    w.freeze_panes = "A5"
wb.active = 0
wb.save(XLSX)
print("model written", XLSX)
