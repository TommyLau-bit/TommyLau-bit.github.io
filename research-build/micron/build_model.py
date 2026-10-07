"""Micron initiation model, 7 Oct 2026. Live formulas throughout.

House style follows the Vertiv and Broadcom models: navy header rows, blue font for hardcoded inputs, yellow fill on
my own assumptions, green font for cross-sheet links, black for formulas.
Non-GAAP basis unless marked GAAP. Fiscal years end on the Thursday closest to 31 August (FY2026 = year to 3 September 2026;
FY2027 = year to about 2 September 2027). HBM and industry supply statements are in calendar years.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from data import *
import data as _d

NAVY = "1F3864"; MUTED = "595959"
F_TTL = Font(bold=True, color=NAVY, size=12)
F_SUB = Font(color=MUTED, size=9)
F_HDR = Font(bold=True, color="FFFFFF", size=10)
F_IN = Font(color="0000FF")
F_INB = Font(color="0000FF", bold=True)
F_LINK = Font(color="008000")
F_B = Font(bold=True)
F_NOTE = Font(color=MUTED, size=9, italic=True)
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_ASSUME = PatternFill("solid", fgColor="FFFF00")
FILL_KEY = PatternFill("solid", fgColor="D9E2F3")
WRAP = Alignment(wrap_text=True, vertical="top")
RIGHT = Alignment(horizontal="right")

USD2 = '"US$"#,##0.00'; USD0 = '"US$"#,##0'; BN2 = '#,##0.00'; BN3 = '#,##0.000'; BN1 = '#,##0.0'
PCT1 = '0.0%'; PCT0 = '0%'; PCT2 = '0.00%'; MULT = '0.0"x"'; MULT0 = '0"x"'
SUB = f"Tommy Lau | Micron Technology, Inc. (Nasdaq: MU) | {DATE_LONG} | Non-GAAP basis; FY26 ended 3 Sep 2026, FY27 ends about 2 Sep 2027 | Personal research, not investment advice."

wb = openpyxl.Workbook()


def sheet(name, title, widths):
    ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet(name)
    ws.title = name
    ws["A1"] = title; ws["A1"].font = F_TTL
    ws["A2"] = SUB; ws["A2"].font = F_SUB
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def head(ws, row, cells, start=1):
    for i, c in enumerate(cells, start):
        x = ws.cell(row=row, column=i, value=c)
        x.font = F_HDR; x.fill = FILL_HDR
        x.alignment = Alignment(horizontal="left" if i == start else "right", vertical="center", wrap_text=True)


def note(ws, row, texts):
    for i, t in enumerate(texts):
        ws.cell(row=row + i, column=1, value=t).font = F_NOTE


def is_formula(v):
    return isinstance(v, str) and v.startswith("=")


def put(ws, r, c, v, fmt=None, bold=False):
    x = ws.cell(row=r, column=c, value=v if v is not None else "")
    if fmt: x.number_format = fmt
    x.alignment = RIGHT
    if v is not None and not is_formula(v) and not isinstance(v, str):
        x.font = Font(color="0000FF", bold=bold)
    elif isinstance(v, str) and v.startswith("=Assumptions"):
        x.font = Font(color="008000", bold=bold)
    elif bold:
        x.font = F_B
    return x


sheet("Cover", "MICRON TECHNOLOGY, INC. (NASDAQ: MU)  INITIATION MODEL", [38, 26, 90])

# ===================================================================== ASSUMPTIONS
ws = sheet("Assumptions", "ASSUMPTIONS AND INPUTS", [64, 14, 10, 92])
ws["A3"] = ("Blue = hardcoded input. Yellow fill = my own assumption (change these to move the value). Green = link to "
            "another tab. Black = formula. US$ billion unless stated; per-share figures in US$; shares in billions.")
ws["A3"].font = F_NOTE
head(ws, 5, ["Input", "Value", "Unit", "Source / note"])
A = {}
MINE = {"q2rev", "q3rev", "q4rev", "q2gm", "q3gm", "q4gm", "q2opex", "q3opex", "q4opex", "q2oth", "q3oth", "q4oth", "sh",
        "g28", "gm28", "opex28", "oth28", "da27", "capex27", "wc27", "hbm26", "hbm27", "hbm28", "nand27", "nand28",
        "bits", "nyears", "pf", "nm", "mult", "rate", "w28", "w29", "w30"}
rows = [
    ("MARKET (retrieved 7 October 2026)", None, None, None, None),
    ("price", "Share price (Nasdaq: MU)", PRICE, "US$", "Nasdaq close, 6 October 2026. Nasdaq.com historical data (Yahoo Finance refused requests on 7 Oct)."),
    ("hiclose", "Highest close of the past year (25 June 2026)", HI_CLOSE, "US$", "Nasdaq.com. Intraday high US$1,255.00, same day."),
    ("loclose", "Lowest close of the past year (10 October 2025)", LO_CLOSE, "US$", "Nasdaq.com. Intraday low US$179.61, same day."),
    ("pxdec25", "Close, 31 Dec 2025", PRICE_DEC25, "US$", "Nasdaq.com."),
    ("fy22hi", "Highest close of fiscal 2022 (14 January 2022)", FY22_HI[0], "US$", "Nasdaq.com."),
    ("fy23lo", "Lowest close after it (26 September 2022)", FY23_LO[0], "US$", "Nasdaq.com."),
    ("sharesout", "Shares outstanding", SHARES_OUT, "bn", "stockanalysis.com, 7 October 2026."),
    ("cashinv", "Cash and investments, 3 Sep 2026", CASH_INV, "US$bn", "Q4 FY26 release: cash 38.364, short-term 5.070, long-term 30.019."),
    ("debt", "Total debt, 3 Sep 2026", DEBT, "US$bn", "Q4 FY26 release: current 0.491, long-term 4.688."),
    ("deposits", "Customer cash deposits under SCAs on the balance sheet", DEPOSITS, "US$bn", "Q4 FY26 prepared remarks: US$12.7bn, returned to customers over time."),
    ("cons27", "FY27 consensus EPS", CONS_EPS_27, "US$", "stockanalysis.com (S&P Global, updated 6 Oct 2026), 39 analysts; range 161.00 to 214.70."),
    ("cons28", "FY28 consensus EPS", CONS_EPS_28, "US$", "stockanalysis.com (S&P Global)."),
    ("consrev27", "FY27 consensus revenue", CONS_REV_27, "US$bn", "stockanalysis.com (S&P Global)."),
    ("consrev28", "FY28 consensus revenue", CONS_REV_28, "US$bn", "stockanalysis.com (S&P Global)."),
    ("constp", "Average analyst target", CONS_TP, "US$", "stockanalysis.com, 49 analysts; range 361 to 2,200."),
    ("REPORTED AND GUIDED (Micron release, prepared remarks, calls)", None, None, None, None),
    ("fy25rev", "FY25 revenue", FY25["rev"], "US$bn", "Q4 FY26 release comparatives."),
    ("fy25ni", "FY25 non-GAAP net income", FY25["ni"], "US$bn", "Same."),
    ("fy25eps", "FY25 non-GAAP EPS", FY25["eps"], "US$", "Same."),
    ("fy25gm", "FY25 non-GAAP gross margin", FY25["gm"], "%", "Q4 FY26 release comparatives: 40.9%."),
    ("fy26rev", "FY26 revenue", FY26["rev"], "US$bn", "Q4 FY26 release."),
    ("fy26ni", "FY26 non-GAAP net income", FY26["ni"], "US$bn", "Same."),
    ("fy26eps", "FY26 non-GAAP EPS", FY26["eps"], "US$", "Same."),
    ("fy26gm", "FY26 non-GAAP gross margin", FY26["gm"], "%", "Prepared remarks, 30 Sep 2026."),
    ("fy26dram", "FY26 DRAM revenue (sum of quarters)", FY26["dram"], "US$bn", "Quarterly tab; 'surpassed $100 billion' (CEO)."),
    ("fy26nand", "FY26 NAND revenue (sum of quarters)", FY26["nand"], "US$bn", "Quarterly tab."),
    ("g1rev", "Q1 FY27 revenue guide", G1["rev"], "US$bn", "Release, 30 Sep 2026: US$61.5bn plus or minus US$1.5bn."),
    ("g1gm", "Q1 FY27 non-GAAP gross margin guide", G1["gm"], "%", "About 86.25%; Q1 is 'the floor for gross margins in fiscal 2027'."),
    ("g1opex", "Q1 FY27 non-GAAP operating expenses guide", G1["opex"], "US$bn", "About US$2.06bn."),
    ("g1eps", "Q1 FY27 non-GAAP EPS guide", G1["eps"], "US$", "US$38.15 plus or minus US$1.00, on about 1.15bn shares."),
    ("tax", "Tax rate, Q1 and FY27", TAX, "%", "Guided around 15.5%; held for FY28 (mine)."),
    ("opexup", "FY27 operating expense increase", OPEX_FY27_UP, "US$bn", "Guided about US$2.5bn on FY26's about US$6.8bn."),
    ("capexh1", "First-half FY27 capex", CAPEX_H1, "US$bn", "Guided about US$25bn; second half higher."),
    ("dps", "Dividend per share, a year", 0.60, "US$", "US$0.15 a quarter."),
    ("MY ASSUMPTIONS: FISCAL 2027 QUARTERS 2 TO 4 AND FISCAL 2028", None, None, None, None),
    ("sh", "Diluted shares, FY27 and FY28", SH, "bn", "MINE: the Q1 guide held; buybacks treated as cash kept, value-neutral at fair value."),
    ("q2rev", "Q2 FY27 revenue", Q27_REV[1], "US$bn", "MINE: 'sequential revenue growth each quarter' with slower price rises."),
    ("q3rev", "Q3 FY27 revenue", Q27_REV[2], "US$bn", "MINE."),
    ("q4rev", "Q4 FY27 revenue", Q27_REV[3], "US$bn", "MINE."),
    ("q2gm", "Q2 FY27 gross margin", Q27_GM[1], "%", "MINE: higher gross margins beyond Q1 (guided)."),
    ("q3gm", "Q3 FY27 gross margin", Q27_GM[2], "%", "MINE."),
    ("q4gm", "Q4 FY27 gross margin", Q27_GM[3], "%", "MINE."),
    ("q2opex", "Q2 FY27 operating expenses", Q27_OPEX[1], "US$bn", "MINE: FY27 total equals FY26's 6.8 plus the guided 2.5."),
    ("q3opex", "Q3 FY27 operating expenses", Q27_OPEX[2], "US$bn", "MINE."),
    ("q4opex", "Q4 FY27 operating expenses", Q27_OPEX[3], "US$bn", "MINE."),
    ("q2oth", "Q2 FY27 interest and other income", Q27_OTHER[1], "US$bn", "MINE. Q1 implied by the guide (Quarterly tab)."),
    ("q3oth", "Q3 FY27 interest and other income", Q27_OTHER[2], "US$bn", "MINE."),
    ("q4oth", "Q4 FY27 interest and other income", Q27_OTHER[3], "US$bn", "MINE."),
    ("g28", "FY28 revenue growth (base)", G28, "%", "MINE. Consensus about 14%."),
    ("gm28", "FY28 gross margin (base)", GM28, "%", "MINE."),
    ("opex28", "FY28 operating expenses", OPEX28, "US$bn", "MINE."),
    ("oth28", "FY28 interest and other income", OTHER28, "US$bn", "MINE."),
    ("da27", "FY27 depreciation and amortisation", DA27, "US$bn", "MINE. FY26 US$9.5bn (stockanalysis.com)."),
    ("capex27", "FY27 capital spending, net", CAPEX27, "US$bn", "MINE. First half about US$25bn guided, second half higher, which implies more than US$50bn."),
    ("wc27", "FY27 working capital build", WC27, "US$bn", "MINE."),
    ("hbm26", "FY26 HBM revenue", HBM26, "US$bn", "MINE. Not disclosed. About Micron's DRAM share of an industry HBM market near US$49bn in CY2026."),
    ("hbm27", "FY27 HBM revenue", SPLIT27["hbm"], "US$bn", "MINE. About 23% of a CY2027 market near US$69bn (Micron: about US$35bn in 2025, about US$100bn in 2028)."),
    ("hbm28", "FY28 HBM revenue", SPLIT28["hbm"], "US$bn", "MINE. About 22% of Micron's US$100bn CY2028 market forecast."),
    ("nand27", "FY27 NAND revenue", SPLIT27["nand"], "US$bn", "MINE. 26% of Q4 FY26 revenue."),
    ("nand28", "FY28 NAND revenue", SPLIT28["nand"], "US$bn", "MINE."),
    ("MY ASSUMPTIONS: NORMAL EARNINGS AND VALUE (base)", None, None, None, None),
    ("bits", "Bit growth a year to a mid-cycle year", BITS, "x", "MINE. Micron expects industry DRAM bits up in the low 20s % in CY2027 and CY2028."),
    ("nyears", "Years from FY26 to the mid-cycle year", N_YEARS, "yrs", "MINE. Around FY2030."),
    ("pf", "Mid-cycle revenue per bit against FY26's average", PF, "%", "MINE."),
    ("nm", "Mid-cycle net margin", NM, "%", "MINE. FY16 to FY25 GAAP: 18.8% in all; FY18 peak 46.5% (Normal tab)."),
    ("mult", "P/E on normal earnings", _d.MULT, "x", "MINE."),
    ("rate", "Discount rate on above-normal earnings", RATE, "%", "MINE."),
    ("w28", "Share of FY28 excess earned in FY28", W_UP[0], "%", "MINE: base keeps the shortage through FY28."),
    ("w29", "Share of FY28 excess earned again in FY29", W_UP[1], "%", "MINE: half."),
    ("w30", "Share of FY28 excess earned again in FY30", W_UP[2], "%", "MINE: none."),
]
r = 6
for key, label, val, unit, src in rows:
    if label is None:
        ws.cell(row=r, column=1, value=key).font = Font(bold=True, color=NAVY); r += 1; continue
    ws.cell(row=r, column=1, value=label)
    c = ws.cell(row=r, column=2, value=val)
    mine = key in MINE
    c.font = F_INB if mine else F_IN
    if mine: c.fill = FILL_ASSUME
    c.number_format = PCT2 if unit == "%" else (MULT0 if unit == "x" and key == "mult" else ("0.00" if unit == "x" else ("#,##0.000" if unit == "bn" else ("#,##0.00" if unit == "US$" else BN3))))
    ws.cell(row=r, column=3, value=unit)
    ws.cell(row=r, column=4, value=src).alignment = WRAP
    A[key] = f"Assumptions!$B${r}"
    r += 1

# ===================================================================== QUARTERLY
ws = sheet("Quarterly", "QUARTERLY HISTORY, Q4 FY25 TO Q4 FY26, AND FISCAL 2027 (Q1 GUIDE, Q2 TO Q4 MINE)", [50] + [11] * 9)
QH = QL + ["Q1 FY27 guide", "Q2 FY27E", "Q3 FY27E", "Q4 FY27E"]
head(ws, 4, ["US$ billion unless stated"] + QH)
C = [get_column_letter(2 + j) for j in range(9)]   # B..J ; F = Q4 FY26, G..J = FY27


def qrow(row, label, vals, fmt, bold=False):
    ws.cell(row=row, column=1, value=label).font = F_B if bold else Font()
    for j, v in enumerate(vals):
        put(ws, row, 2 + j, v, fmt)


qrow(5, "Revenue", Q_REV + ["=" + A["g1rev"], "=" + A["q2rev"], "=" + A["q3rev"], "=" + A["q4rev"]], BN3, True)
qrow(6, "  DRAM, including HBM", Q_DRAM + [None] * 4, BN1)
qrow(7, "  NAND", Q_NAND + [None] * 4, BN1)
qrow(8, "Revenue growth on the quarter", [None] + [f"={C[j]}5/{C[j - 1]}5-1" for j in range(1, 9)], PCT1)
qrow(9, "Gross margin, non-GAAP", Q_GM + ["=" + A["g1gm"], "=" + A["q2gm"], "=" + A["q3gm"], "=" + A["q4gm"]], PCT1)
qrow(10, "Gross profit", [f"={c}5*{c}9" for c in C], BN2)
qrow(11, "Operating expenses, non-GAAP", Q_OPEX + ["=" + A["g1opex"], "=" + A["q2opex"], "=" + A["q3opex"], "=" + A["q4opex"]], BN2)
qrow(12, "Operating income, non-GAAP", Q_OP[:1] + Q_OP[1:] + [f"={c}10-{c}11" for c in C[5:]], BN2, True)
qrow(13, "Operating margin", [None] + [f"={c}12/{c}5" for c in C[1:]], PCT1)
qrow(14, "Interest and other (Q1: implied by the EPS guide)",
     [None] * 5 + [f"=G17*G16/(1-{A['tax']})-G12", "=" + A["q2oth"], "=" + A["q3oth"], "=" + A["q4oth"]], BN3)
qrow(15, "Net income, non-GAAP", [None] * 5 + [f"=(G12+G14)*(1-{A['tax']})"] + [f"=({c}12+{c}14)*(1-{A['tax']})" for c in C[6:]], BN2, True)
qrow(16, "Diluted shares, bn", [None] * 4 + [SH_Q4] + ["=" + A["sh"]] * 4, "#,##0.000")
qrow(17, "Diluted EPS, US$ (reported; Q1 guide; Q2 to Q4 mine)", Q_EPS + ["=" + A["g1eps"]] + [f"={c}15/{c}16" for c in C[6:]], USD2, True)
ws["A19"] = "Business unit gross margin, fiscal 2026"; ws["A19"].font = Font(bold=True, color=NAVY)
for i, (k, v) in enumerate(BU_GM.items()):
    qrow(20 + i, k, [None] + v, PCT0)
ws["A25"] = "Business unit revenue, fiscal 2026"; ws["A25"].font = Font(bold=True, color=NAVY)
for i, (k, v) in enumerate(BU_REV.items()):
    qrow(26 + i, k, [None] + v, BN1)
ws["A31"] = "Checks"; ws["A31"].font = Font(bold=True, color=NAVY)
chk = [
    (32, "FY26 revenue, sum of quarters", "=SUM(C5:F5)", BN3),
    (33, "Revenue growth, Q4 FY26 on Q4 FY25", "=F5/B5-1", PCT1),
    (34, "Gap between core data centre and cloud unit margins, Q4 FY26, points", "=(F21-F20)*100", BN1),
    (35, "FY27 revenue, sum of quarters", "=SUM(G5:J5)", BN2),
    (36, "FY27 EPS, sum of quarters", "=SUM(G17:J17)", USD2),
]
for rr, lab, f, fmt in chk:
    ws.cell(row=rr, column=1, value=lab); c = ws.cell(row=rr, column=2, value=f); c.number_format = fmt
note(ws, 38, ["Source: Micron prepared remarks (24 June and 30 September 2026), Q4 FY26 release, earnings calls of 23 September 2025, 17 December 2025 and 18 March 2026.",
              "Q2 FY26 revenue is FY26 less the other three quarters. Q4 FY25 DRAM and NAND are derived from the year-on-year growth given on 30 September 2026.",
              "Q1 FY27 is the guide; interest and other is what makes the guided EPS. Q2 to Q4 FY27 are my estimates."])

# ===================================================================== FINANCIALS
ws = sheet("Financials", "FINANCIAL SUMMARY AND FORECASTS (US$ billion unless stated, non-GAAP basis)", [52, 12, 12, 12, 12])
head(ws, 4, ["", "FY25A", "FY26A", "FY27E*", "FY28E*"])
fin = [
    (5, "Revenue", ["=" + A["fy25rev"], "=" + A["fy26rev"], "=SUM(Quarterly!G5:J5)", f"=D5*(1+{A['g28']})"], BN2, True),
    (6, "Revenue growth", [None, "=C5/B5-1", "=D5/C5-1", "=E5/D5-1"], PCT1, False),
    (7, "  HBM (my estimate; not disclosed)", [None, "=" + A["hbm26"], "=" + A["hbm27"], "=" + A["hbm28"]], BN1, False),
    (8, "  Conventional DRAM", [None, f"={A['fy26dram']}-C7", "=D5-D7-D9", "=E5-E7-E9"], BN1, False),
    (9, "  NAND", [None, "=" + A["fy26nand"], "=" + A["nand27"], "=" + A["nand28"]], BN1, False),
    (10, "Gross margin", ["=" + A["fy25gm"], "=" + A["fy26gm"], "=SUM(Quarterly!G10:J10)/D5", "=" + A["gm28"]], PCT1, False),
    (11, "Operating expenses", [None, 6.8, "=SUM(Quarterly!G11:J11)", "=" + A["opex28"]], BN2, False),
    (12, "Operating income", [None, FY26["op"], "=D5*D10-D11", "=E5*E10-E11"], BN2, True),
    (13, "Interest and other", [None, None, "=SUM(Quarterly!G14:J14)", "=" + A["oth28"]], BN2, False),
    (14, "Net income", ["=" + A["fy25ni"], "=" + A["fy26ni"], "=SUM(Quarterly!G15:J15)", f"=(E12+E13)*(1-{A['tax']})"], BN2, True),
    (15, "Net margin", ["=B14/B5", "=C14/C5", "=D14/D5", "=E14/E5"], PCT1, False),
    (16, "Diluted EPS, US$", ["=" + A["fy25eps"], "=" + A["fy26eps"], "=SUM(Quarterly!G17:J17)", f"=E14/{A['sh']}"], USD2, True),
    (17, "Consensus EPS, US$", [None, None, "=" + A["cons27"], "=" + A["cons28"]], USD2, False),
    (18, "Mine against consensus", [None, None, "=D16/D17-1", "=E16/E17-1"], PCT1, False),
    (19, "Revenue against consensus", [None, None, f"=D5/{A['consrev27']}-1", f"=E5/{A['consrev28']}-1"], PCT1, False),
    (20, "P/E at reference price", [f"={A['price']}/B16", f"={A['price']}/C16", f"={A['price']}/D16", f"={A['price']}/E16"], MULT, False),
    (21, "P/E on consensus", [None, None, f"={A['price']}/D17", f"={A['price']}/E17"], MULT, False),
    (23, "Free cash flow, FY27 (net income + D&A - capex - working capital)", [None, FY26["fcf"], f"=D14+{A['da27']}-{A['capex27']}-{A['wc27']}", None], BN2, False),
    (24, "Net cash after deposits, start of period", [None, None, f"={A['cashinv']}-{A['debt']}-{A['deposits']}", None], BN2, False),
    (25, "Net cash after deposits, end of FY27 (about October 2027); per share on 1.15bn diluted shares, while market value uses 1.13bn outstanding", [None, None, f"=D24+D23-{A['dps']}*{A['sh']}", None], BN2, True),
    (26, "  per share, US$", [None, None, f"=D25/{A['sh']}", None], USD2, True),
]
for rr, label, vals, fmt, bold in fin:
    ws.cell(row=rr, column=1, value=label).font = F_B if bold else Font()
    for j, v in enumerate(vals):
        put(ws, rr, 2 + j, v, fmt)
note(ws, 28, [
    "*FY27E sums the Q1 guide and my Q2 to Q4 (Quarterly tab). FY28E is my base case. HBM and the conventional DRAM split are my estimates; Micron reports DRAM including HBM.",
    "FY26 operating expenses are the four quarters rounded (1.3, 1.4, 1.5, 2.6); FY26 operating income is the sum of the quarters (6.4, 16.5, 33.7, 44.6). FY26 free cash flow is Micron's adjusted figure.",
    "Consensus from stockanalysis.com (S&P Global data, updated 6 October 2026).",
    f"Note check: FY27E revenue {REV27:.1f}, EPS {EPS27:.2f}; FY28E revenue {REV28:.1f}, EPS {EPS28:.2f}; net cash per share at October 2027 {NC_PS:.1f}.",
])

# ===================================================================== NORMAL
ws = sheet("Normal", "THE CYCLE AND NORMAL (MID-CYCLE) EARNINGS", [40, 12, 12, 12])
head(ws, 4, ["Fiscal year", "Revenue, US$bn", "GAAP net income", "Net margin"])
for i, (y, rv, ni) in enumerate(zip(YRS10, REV10, NI10)):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=y)
    put(ws, rr, 2, rv, BN3); put(ws, rr, 3, ni, BN3); put(ws, rr, 4, f"=C{rr}/B{rr}", PCT1)
ws["A17"] = "FY16 to FY25, all ten years"; ws["A17"].font = F_B
ws["B17"] = "=SUM(B5:B14)"; ws["C17"] = "=SUM(C5:C14)"; ws["D17"] = "=C17/B17"
for c, f in (("B17", BN2), ("C17", BN2), ("D17", PCT1)): ws[c].number_format = f
ws["A18"] = "FY18, the last cycle's peak"; ws["D18"] = "=D7"; ws["D18"].number_format = PCT1
head(ws, 20, ["Normal year (base)", "Value", "", ""])
nrows = [
    (21, "FY26 revenue", "=" + A["fy26rev"], BN2),
    (22, "Bit growth factor", f"={A['bits']}^{A['nyears']}", "0.000"),
    (23, "Revenue per bit against FY26's average", "=" + A["pf"], PCT0),
    (24, "Normal revenue, US$bn", "=B21*B22*B23", BN1),
    (25, "Normal net margin", "=" + A["nm"], PCT0),
    (26, "Normal net income, US$bn", "=B24*B25", BN2),
    (27, "Normal EPS, US$", f"=B26/{A['sh']}", USD2),
]
for rr, lab, f, fmt in nrows:
    ws.cell(row=rr, column=1, value=lab); c = ws.cell(row=rr, column=2, value=f); c.number_format = fmt
ws["A27"].font = F_B; ws["B27"].fill = FILL_KEY
note(ws, 29, ["Revenue and GAAP net income: Micron 10-K XBRL data on SEC EDGAR for FY16 to FY21; stockanalysis.com (S&P Global) for FY22 to FY26.",
              "The normal year is mine: bits grow 20% a year for four years from FY26, and revenue per bit settles at 55% of FY26's average."])

# ===================================================================== SCENARIOS
ws = sheet("Scenarios", "THREE CASES, TWELVE MONTHS OUT (VALUE AT OCTOBER 2027)", [12] + [11] * 15)
head(ws, 4, ["Case", "Rev. per bit", "Normal margin", "Normal P/E", "FY28 growth", "FY28 GM", "w FY28", "w FY29", "w FY30",
             "Normal EPS", "FY28 EPS", "Net cash / sh", "Excess / sh", "Normal value / sh", "Value, US$", "vs price"])
for i, s in enumerate(SCEN):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=s[0]).font = F_B
    if s[0] == "Base":
        for col, key, fmt in [(2, "pf", PCT0), (3, "nm", PCT0), (4, "mult", MULT0), (5, "g28", PCT0), (6, "gm28", PCT1),
                              (7, "w28", PCT0), (8, "w29", PCT0), (9, "w30", PCT0)]:
            c = ws.cell(row=rr, column=col, value="=" + A[key]); c.font = F_LINK; c.number_format = fmt
    else:
        vals = [s[1], s[2], s[3], s[4], s[5], s[6][0], s[6][1], s[6][2]]
        fmts = [PCT0, PCT0, MULT0, PCT0, PCT1, PCT0, PCT0, PCT0]
        for col, (v, fmt) in enumerate(zip(vals, fmts), 2):
            c = ws.cell(row=rr, column=col, value=v); c.font = F_INB; c.fill = FILL_ASSUME; c.number_format = fmt
    ws.cell(row=rr, column=10, value=f"=Normal!$B$21*Normal!$B$22*B{rr}*C{rr}/{A['sh']}").number_format = USD2
    ws.cell(row=rr, column=11, value=f"=(Financials!$D$5*(1+E{rr})*F{rr}-{A['opex28']}+{A['oth28']})*(1-{A['tax']})/{A['sh']}").number_format = USD2
    ws.cell(row=rr, column=12, value="=Financials!$D$26").number_format = USD2
    ws.cell(row=rr, column=13, value=f"=(K{rr}-J{rr})*(G{rr}/(1+{A['rate']})+H{rr}/(1+{A['rate']})^2+I{rr}/(1+{A['rate']})^3)").number_format = USD2
    ws.cell(row=rr, column=14, value=f"=J{rr}*D{rr}").number_format = USD2
    ws.cell(row=rr, column=15, value=f"=L{rr}+M{rr}+N{rr}").number_format = USD2
    ws.cell(row=rr, column=16, value=f"=O{rr}/{A['price']}-1").number_format = PCT0
ws["A10"] = "Probability"; ws["A10"].font = F_B
for j, s in enumerate(SCEN):
    ws.cell(row=9, column=2 + j, value=s[0]).font = F_B
    c = ws.cell(row=10, column=2 + j, value=s[7]); c.font = F_INB; c.fill = FILL_ASSUME; c.number_format = PCT0
ws["E10"] = "=SUM(B10:D10)"; ws["E10"].number_format = PCT0; ws["F10"] = "must be 100%"; ws["F10"].font = F_NOTE
ws["A12"] = "Probability-weighted value, US$"; ws["A12"].font = F_B
ws["O12"] = "=B10*O5+C10*O6+D10*O7"; ws["O12"].number_format = USD2; ws["O12"].font = F_B; ws["O12"].fill = FILL_KEY
ws["P12"] = f"=O12/{A['price']}-1"; ws["P12"].number_format = PCT1
note(ws, 14, ["Value at October 2027 = net cash after deposits per share + above-normal earnings to FY30, discounted + normal EPS x normal P/E.",
              "Excess = (FY28 EPS - normal EPS) x (w FY28 / 1.1 + w FY29 / 1.1^2 + w FY30 / 1.1^3). FY27 and the cash at October 2027 are the same in all three cases.",
              "Bear: prices turn in 2027 as new cleanrooms ramp; FY28 revenue falls 25%; normal year at the ten-year margin. Bull: the shortage outlasts 2028, which Micron's 'no line of sight' to balance allows, and floors hold margins near 40%.",
              "FY27 is held fixed in all cases on purpose: Q1 is guided and more than 75% of 2027 output is committed, much of it under contract."])

# ===================================================================== VALUATION
ws = sheet("Valuation", "VALUATION: NET CASH + ABOVE-NORMAL EARNINGS + NORMAL EARNINGS, AT OCTOBER 2027", [72, 16, 70])
head(ws, 4, ["Line", "Value", "Working / note"])
val = [
    ("nc", "Net cash after deposits per share, October 2027, US$", "=Financials!D26", USD2, "Financials tab."),
    ("exc", "Above-normal earnings, discounted, per share, US$", "=Scenarios!M6", USD2, "Base case."),
    ("term", "Normal EPS x normal P/E, US$", "=Scenarios!N6", USD2, "Base case."),
    ("base", "Base-case value per share, US$", "=B5+B6+B7", USD2, ""),
    ("px", "Reference price, US$", "=" + A["price"], USD2, "Nasdaq close, 6 October 2026."),
    ("upb", "Base value against the price", "=B8/B9-1", PCT1, ""),
    ("wtd", "Probability-weighted value, US$", "=Scenarios!O12", USD2, "Scenarios tab."),
    ("rule", "Rule on value alone (long at 15% above, short at 25% below)", '=IF(B10>=0.15,"LONG",IF(B10<=-0.25,"SHORT","NO CALL"))', None, "Mechanical only; wider on the short side because a short's losses are open-ended."),
    ("call", "DRAFT VIEW" if CALL["draft"] else "CALL", CALL["direction"], None, "Draft for Tommy Lau's decision." if CALL["draft"] else "Tommy Lau's call."),
    ("tgt", "Target, US$", CALL["target"] if CALL["target"] else "None (no call)", USD2, ""),
    ("mcap", "Market value, US$bn", f"={A['price']}*{A['sharesout']}", BN1, "On 1.13bn shares outstanding."),
    ("ev", "Enterprise value, US$bn (net of cash after deposits)", f"=B15-({A['cashinv']}-{A['debt']}-{A['deposits']})", BN1, ""),
    ("pe26", "P/E on FY26 EPS", "=Financials!C20", MULT, ""),
    ("pe27", "P/E on my FY27E", "=Financials!D20", MULT, ""),
    ("pe28", "P/E on my FY28E", "=Financials!E20", MULT, ""),
    ("pec27", "P/E on FY27 consensus", "=Financials!D21", MULT, "stockanalysis.com."),
    ("pec28", "P/E on FY28 consensus", "=Financials!E21", MULT, ""),
    ("pen", "P/E on my normal EPS", f"={A['price']}/Normal!B27", MULT, ""),
    ("k", "Discount factor on the FY28 excess (base weights)", f"={A['w28']}/(1+{A['rate']})+{A['w29']}/(1+{A['rate']})^2+{A['w30']}/(1+{A['rate']})^3", "0.0000", ""),
    ("need", "Normal EPS the price needs, US$", f"=({A['price']}-B5-B23*Scenarios!K6)/({A['mult']}-B23)", USD2, "At my base FY28, cash and multiple. What the market prices in."),
    ("needm", "That as a net margin on my normal revenue", f"=B24*{A['sh']}/Normal!B24", PCT1, "FY18 peak 46.5%; FY16 to FY25 18.8%."),
    ("needrev", "Or as normal revenue at my 30% margin, US$bn", f"=B24*{A['sh']}/{A['nm']}", BN1, ""),
    ("need26", "Needed normal EPS against FY26 EPS", f"=B24/{A['fy26eps']}", PCT0, ""),
    ("need28", "Needed normal EPS against my FY28 EPS", "=B24/Financials!E16", PCT0, ""),
    ("rlong", "Price at or below which the base is 15% above (points long)", "=B8/1.15", USD2, ""),
    ("rshort", "Price at or above which the base is 25% below (points short)", "=B8/0.75", USD2, ""),
    ("nlong", "Normal EPS at or above which the base is 15% above the price", f"=({A['price']}*1.15-B5-B23*Scenarios!K6)/({A['mult']}-B23)", USD2, ""),
    ("nshort", "Normal EPS at or below which the base is 25% below the price", f"=({A['price']}*0.75-B5-B23*Scenarios!K6)/({A['mult']}-B23)", USD2, ""),
    ("fy22pe", "Last cycle: P/E at fiscal 2022's highest close on FY22 GAAP EPS of US$7.75", f"={A['fy22hi']}/7.75", MULT, ""),
    ("fy23fall", "Fall from that close to 26 September 2022", f"={A['fy23lo']}/{A['fy22hi']}-1", PCT0, ""),
    ("offhi", "Against the highest close (25 Jun 2026)", f"={A['price']}/{A['hiclose']}-1", PCT1, ""),
    ("rise", "Change since 31 Dec 2025", f"={A['price']}/{A['pxdec25']}-1", PCT0, ""),
]
V = {}
for i, (key, label, f, fmt, nt) in enumerate(val):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=label)
    c = ws.cell(row=rr, column=2, value=f)
    if fmt: c.number_format = fmt
    if not is_formula(f): c.font = F_IN
    if key in ("base", "call"):
        ws.cell(row=rr, column=1).font = F_B; c.font = F_B; c.fill = FILL_KEY
    if key == "call": c.font = F_INB
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
    V[key] = f"Valuation!$B${rr}"
assert V["base"] == "Valuation!$B$8" and V["k"] == "Valuation!$B$23" and V["need"] == "Valuation!$B$24" and V["mcap"] == "Valuation!$B$15"

# ===================================================================== SENSITIVITY
ws = sheet("Sensitivity", "VALUE PER SHARE: NORMAL NET MARGIN AGAINST THE MULTIPLE ON NORMAL EARNINGS (US$)", [30, 14, 14, 14])
head(ws, 4, ["Normal net margin \\ normal P/E", "", "", ""])
for j, (m, link) in enumerate(zip(SENS_X, [False, True, False])):
    c = ws.cell(row=4, column=2 + j, value=("=" + A["mult"]) if link else m)
    c.number_format = MULT0; c.font = Font(bold=True, color="FFFFFF"); c.fill = FILL_HDR; c.alignment = RIGHT
for i, (m, link) in enumerate(zip(SENS_NM, [False, True, False])):
    rr = 5 + i
    c = ws.cell(row=rr, column=1, value=("=" + A["nm"]) if link else m)
    c.number_format = PCT0; c.font = F_LINK if link else F_IN
    for j in range(3):
        col = get_column_letter(2 + j)
        ws.cell(row=rr, column=2 + j,
                value=f"=Valuation!$B$5+Valuation!$B$23*(Scenarios!$K$6-Normal!$B$24*$A{rr}/{A['sh']})+Normal!$B$24*$A{rr}/{A['sh']}*{col}$4").number_format = USD0
ws["C6"].fill = FILL_KEY
ws["A9"] = "Cells above the reference price"; ws["A9"].font = F_B
ws["B9"] = f'=COUNTIF(B5:D7,">"&{A["price"]})'; ws["C9"] = "of 9"
ws["A10"] = "Cells 25% or more below the price"; ws["A10"].font = F_B
ws["B10"] = f'=COUNTIF(B5:D7,"<="&{A["price"]}*0.75)'; ws["C10"] = "of 9"
note(ws, 12, ["Middle row links to my base normal margin and the middle column to the base multiple. Normal revenue, FY28 EPS and net cash as in the base."])

# ===================================================================== PEERS
ws = sheet("Peers", "PEER MULTIPLES ON CONSENSUS FOR THE YEAR ENDING IN 2027 (stockanalysis.com, 7 OCTOBER 2026)", [24, 16, 14, 8, 14, 16, 12, 40])
head(ws, 4, ["Company", "Listing", "Price", "Currency", "Consensus EPS", "Year", "P/E", "What they make"])
ws.cell(row=5, column=1, value="Micron").font = F_B; ws.cell(row=5, column=2, value="Nasdaq: MU")
c = ws.cell(row=5, column=3, value="=" + A["price"]); c.font = F_LINK; c.number_format = BN2
ws.cell(row=5, column=4, value="US$")
c = ws.cell(row=5, column=5, value="=" + A["cons27"]); c.font = F_LINK; c.number_format = BN2
ws.cell(row=5, column=6, value="FY to Aug 2027"); ws["G5"] = "=C5/E5"; ws["G5"].number_format = MULT
ws.cell(row=5, column=8, value="DRAM, HBM, NAND")
for i, (co, tk, px, cur, eps, yr, what, nt) in enumerate(PEERS):
    rr = 6 + i
    ws.cell(row=rr, column=1, value=co); ws.cell(row=rr, column=2, value=tk)
    put(ws, rr, 3, px, BN2); ws.cell(row=rr, column=4, value=cur); put(ws, rr, 5, eps, BN2)
    ws.cell(row=rr, column=6, value=yr); ws.cell(row=rr, column=7, value=f"=C{rr}/E{rr}").number_format = MULT
    ws.cell(row=rr, column=8, value=what)
co, tk, mc, ni, yr, what = KIOXIA
ws["A9"] = co; ws["B9"] = tk; put(ws, 9, 3, mc, BN2); ws["D9"] = "JPY tn"; put(ws, 9, 5, ni, BN2)
ws["F9"] = yr; ws["G9"] = "=C9/E9"; ws["G9"].number_format = MULT; ws["H9"] = what
ws["A11"] = "Median, four peers"; ws["A11"].font = F_B; ws["G11"] = "=MEDIAN(G6:G9)"; ws["G11"].number_format = MULT
note(ws, 13, ["Prices: Micron and SanDisk Nasdaq closes of 6 October 2026; SK Hynix, Samsung and Kioxia closes of 7 October 2026. Consensus from stockanalysis.com (S&P Global).",
              "Kioxia: market value over consensus net income (no per-share consensus shown), both in JPY trillion. Fiscal year-ends differ by up to eight months, so the multiples are not on one basis; multiples only.",
              "No view on the peers' shares is expressed."])

# ===================================================================== SOURCES
ws = sheet("Sources", "SOURCES", [6, 130])
head(ws, 4, ["No.", "Source and what it supports"])
srcs = [
    "Micron fourth quarter and fiscal 2026 results release, Exhibit 99.1 to Form 8-K, 30 September 2026: revenue, GAAP and non-GAAP margins, net income, EPS, business unit revenue, cash flow, cash, investments, debt, diluted shares, Q1 FY27 guide, dividend.",
    "Micron fiscal Q4 2026 prepared remarks, 30 September 2026: DRAM US$39.8bn and NAND US$14.1bn, bits and prices; business unit margins; HBM agreements for CY2027; 26 SCAs, over 35% of revenue through 2030, US$32bn of commitments, RPO about US$150bn, US$12.7bn of deposits; industry outlook; fab timing; FY27 opex, tax and capex.",
    "Micron fiscal Q3 2026 prepared remarks, 24 June 2026: Q3 figures; SCA terms, floors and ceilings at CQ2 2026 market prices; supply constraints including energy infrastructure; HBM4 ramp.",
    "Micron earnings calls: Q4 FY25 (23 September 2025; HBM nearly US$2bn, gross margin 45.7%), Q1 FY26 (17 December 2025; three to one trade ratio; CY2026 HBM agreed; HBM market about US$35bn in 2025 to about US$100bn in 2028), Q2 FY26 (18 March 2026), Q4 FY26 (30 September 2026; 2026 HBM prices set the year before; more than 75% of 2027 output committed; questions on peak earnings and de-speccing).",
    "SEC EDGAR XBRL company concept data for Micron (CIK 723125): revenue and net income, FY2016 to FY2021.",
    "Nasdaq.com historical prices, retrieved 7 October 2026. (Yahoo Finance chart API refused requests on 7 October 2026.)",
    "stockanalysis.com, retrieved 7 October 2026: consensus EPS and revenue (S&P Global, updated 6 October 2026), average target, shares outstanding, FY22 to FY26 annual figures, and peer prices and consensus.",
    "The Physical Layer, 'Micron sells AI memory before the year starts, and each bit takes three times the wafer', 7 October 2026: https://thephysicallayer.fyi/journal/micron-three-times-the-wafer/",
    "Estimates for FY27 Q2 to FY28, the HBM split, normal earnings and the valuation are the author's own. Personal research, not investment advice. I hold no position in Micron.",
]
for i, s in enumerate(srcs):
    ws.cell(row=5 + i, column=1, value=i + 1)
    ws.cell(row=5 + i, column=2, value=s).alignment = WRAP

# ===================================================================== COVER
ws = wb["Cover"]
head(ws, 4, ["Item", "Value", "Note"])
cover = [
    ("Draft view" if CALL["draft"] else "Call", "=" + V["call"], None, "Draft for Tommy Lau's decision; not yet a call." if CALL["draft"] else "Tommy Lau's call."),
    ("Target, US$", "=" + V["tgt"], USD2, "No target for a no call."),
    ("Conviction", CALL["conviction"], None, "Judgement, not formula."),
    ("Reference price, US$", "=" + A["price"], USD2, "Nasdaq close, 6 October 2026."),
    ("Base value, US$", "=" + V["base"], USD2, "Net cash + above-normal earnings to FY30 + 12x normal EPS, at October 2027."),
    ("Base value against the price", "=" + V["upb"], PCT1, ""),
    ("Probability-weighted value, US$", "=" + V["wtd"], USD2, "25 / 50 / 25 bear / base / bull."),
    ("Normal EPS the price needs, US$", "=" + V["need"], USD2, "Against my base of US$39.63."),
    ("FY27E / FY28E EPS, US$", '=TEXT(Financials!D16,"0.00")&" / "&TEXT(Financials!E16,"0.00")', None, "Own estimates, non-GAAP, diluted."),
    ("Horizon", "12 months, to October 2027", None, ""),
]
for i, (lab, v, fmt, nt) in enumerate(cover):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=lab).font = F_B
    c = ws.cell(row=rr, column=2, value=v)
    if fmt: c.number_format = fmt
    c.alignment = RIGHT
    c.font = F_LINK if is_formula(v) else F_IN
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
ws["A17"] = "Tabs"; ws["A17"].font = Font(bold=True, color=NAVY)
for i, (t, d) in enumerate([
    ("Assumptions", "Every input, with its source. Change the yellow cells to move the value."),
    ("Quarterly", "Q4 FY25 to Q4 FY26 reported, the Q1 FY27 guide and my Q2 to Q4 FY27; business unit margins."),
    ("Financials", "FY25A to FY28E: revenue by line, margins, EPS, consensus, P/E, cash at October 2027."),
    ("Normal", "Ten years of the cycle and my normal (mid-cycle) year."),
    ("Scenarios", "Bear / base / bull value at October 2027 and the probability-weighted value."),
    ("Valuation", "Base value, multiples, what the price needs, revisit levels."),
    ("Sensitivity", "Normal net margin against the multiple on normal earnings."),
    ("Peers", "Memory peers on consensus for the year ending in 2027."),
    ("Sources", "Every source used."),
]):
    ws.cell(row=18 + i, column=1, value=t); ws.cell(row=18 + i, column=2, value=d)
ws["A29"] = "Colour key: blue = hardcoded input; yellow fill = my own assumption; green = link to another tab; black = formula."
ws["A29"].font = F_NOTE
ws["A30"] = f"Companion note: {DATE}_{COMPANY}_Initiation.pdf. Personal research, not investment advice."; ws["A30"].font = F_NOTE

order = ["Cover", "Assumptions", "Quarterly", "Financials", "Normal", "Scenarios", "Valuation", "Sensitivity", "Peers", "Sources"]
wb._sheets = [wb[n] for n in order]
for w in wb.worksheets:
    w.freeze_panes = "A5"
wb.active = 0
wb.save(XLSX)
print("model written", XLSX)
