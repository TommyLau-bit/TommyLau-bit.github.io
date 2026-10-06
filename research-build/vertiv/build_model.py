"""Vertiv (NYSE: VRT) initiation model, 6 Oct 2026. Live formulas throughout.

House style follows 2026-07-30_EGPEnergy_Model.xlsx: navy header rows, navy bold titles, blue font for
hardcoded inputs, yellow fill on my own assumptions, green font for cross-sheet links, black for formulas.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from data import *

NAVY = "1F3864"; MUTED = "595959"
F_TTL = Font(bold=True, color=NAVY, size=12)
F_SUB = Font(color=MUTED, size=9)
F_HDR = Font(bold=True, color="FFFFFF", size=10)
F_IN = Font(color="0000FF")
F_INB = Font(color="0000FF", bold=True)
F_LINK = Font(color="008000")
F_B = Font(bold=True)
F_NOTE = Font(color=MUTED, size=9, italic=True)
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_ASSUME = PatternFill("solid", fgColor="FFFF00")
FILL_KEY = PatternFill("solid", fgColor="D9E2F3")
THIN = Side(style="thin", color="BFBFBF")
WRAP = Alignment(wrap_text=True, vertical="top")

USD2 = '"US$"#,##0.00'
USD0 = '"US$"#,##0'
BN2 = '#,##0.00'
PCT1 = '0.0%'
PCT0 = '0%'
MULT = '0.0"x"'
MULT0 = '0"x"'

wb = openpyxl.Workbook()


def sheet(name, title, widths, sub="Tommy Lau | Vertiv (NYSE: VRT) | 6 October 2026 | Personal research, not investment advice."):
    ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet(name)
    ws.title = name
    ws["A1"] = title; ws["A1"].font = F_TTL
    ws["A2"] = sub; ws["A2"].font = F_SUB
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def head(ws, row, cells, start=1):
    for i, c in enumerate(cells, start):
        x = ws.cell(row=row, column=i, value=c)
        x.font = F_HDR; x.fill = FILL_HDR
        x.alignment = Alignment(horizontal="left" if i == start else "right", vertical="center", wrap_text=True)


def put(ws, ref, value, fmt=None, font=None, fill=None, align=None):
    c = ws[ref]; c.value = value
    if fmt: c.number_format = fmt
    if font: c.font = font
    if fill: c.fill = fill
    if align: c.alignment = align
    return c


# =====================================================================  COVER
ws = sheet("Cover", "VERTIV HOLDINGS CO (NYSE: VRT)  INITIATION MODEL", [34, 22, 70])
# filled after the other tabs exist (needs their cell addresses)

# =====================================================================  ASSUMPTIONS
ws = sheet("Assumptions", "ASSUMPTIONS AND INPUTS", [40, 14, 12, 92])
ws["A3"] = ("Blue = hardcoded input. Yellow fill = my own assumption (change these to move the target). "
            "Green = link to another tab. Black = formula. US$ billion unless stated.")
ws["A3"].font = F_NOTE
head(ws, 5, ["Input", "Value", "Unit", "Source / note"])
A = {}
rows = [
    ("MARKET", None, None, None, None),
    ("price", "Reference share price", PRICE, "US$", "NYSE close, 5 October 2026. Yahoo Finance, retrieved 6 October 2026."),
    ("peak", "Intraday peak share price, 14 May 2026", PEAK, "US$", "Yahoo Finance, retrieved 6 October 2026."),
    ("shares", "Diluted shares", SHARES_DIL, "m", "Vertiv Form 10-Q, quarter to 30 June 2026, filed 29 July 2026."),
    ("cash", "Cash and cash equivalents, 30 Jun 2026", CASH, "US$m", "Vertiv Form 10-Q, quarter to 30 June 2026."),
    ("stinv", "Short-term investments, 30 Jun 2026", ST_INV, "US$m", "Vertiv Form 10-Q, quarter to 30 June 2026."),
    ("debt", "Total debt, 30 Jun 2026", DEBT, "US$m", "Vertiv Form 10-Q, quarter to 30 June 2026."),
    ("cons27", "2027 consensus adjusted EPS", CONSENSUS_27, "US$", "MarketBeat, retrieved 6 October 2026. Cross-check only."),
    ("fwdpe", "Forward P/E shown by stockanalysis.com", FWD_PE_SA, "x", "stockanalysis.com, retrieved 6 October 2026. Cross-check only."),
    ("REPORTED", None, None, None, None),
    ("s24", "2024 net sales", SALES["2024A"], "US$bn", "Sum of the four 2024 quarterly results releases (Form 8-K, Exhibit 99.1)."),
    ("s25", "2025 net sales", SALES["2025A"], "US$bn", "Sum of the four 2025 quarterly results releases. Rounded quarters on the Quarterly tab sum to 10.24."),
    ("m24", "2024 adjusted operating margin", MARGIN["2024A"], "%", "Vertiv results releases."),
    ("m25", "2025 adjusted operating margin", MARGIN["2025A"], "%", "Vertiv results releases."),
    ("e24", "2024 adjusted EPS", EPS["2024A"], "US$", "Sum of the four 2024 quarterly results releases."),
    ("e25", "2025 adjusted EPS", EPS["2025A"], "US$", "Sum of the four 2025 quarterly results releases."),
    ("2026 GUIDE (release of 29 July 2026)", None, None, None, None),
    ("gs_lo", "2026 net sales guide, low", GUIDE["sales_lo"], "US$bn", "Raised from US$13.25 to 13.75bn in February 2026."),
    ("gs_hi", "2026 net sales guide, high", GUIDE["sales_hi"], "US$bn", "Vertiv second quarter 2026 results release."),
    ("gm_lo", "2026 adjusted operating margin guide, low", GUIDE["m_lo"], "%", "Vertiv second quarter 2026 results release."),
    ("gm_hi", "2026 adjusted operating margin guide, high", GUIDE["m_hi"], "%", "Vertiv second quarter 2026 results release."),
    ("ge_lo", "2026 adjusted EPS guide, low", GUIDE["eps_lo"], "US$", "Bottom of the guide is part of the falsifier (US$6.65)."),
    ("ge_hi", "2026 adjusted EPS guide, high", GUIDE["eps_hi"], "US$", "Vertiv second quarter 2026 results release."),
    ("gop", "2026 adjusted operating profit guide, midpoint", GUIDE["op_mid"], "US$bn", "Vertiv second quarter 2026 results release, midpoint."),
    ("gmay", "2026 adjusted EPS guide at the May peak, midpoint", 6.35, "US$", "Guide of the time, US$6.30 to 6.40. Used only for the peak-multiple cross-check."),
    ("MY ASSUMPTIONS", None, None, None, None),
    ("g27", "2027 net sales growth", G27, "%", "MY ASSUMPTION. Inside Vertiv's long-run target of 20 to 22% organic growth a year."),
    ("g28", "2028 net sales growth", G28, "%", "MY ASSUMPTION. Same basis."),
    ("m27", "2027 adjusted operating margin", M27, "%", "MY ASSUMPTION. About one point a year towards the 2030 target of about 27% or more (Investor Conference, 19 May 2026)."),
    ("m28", "2028 adjusted operating margin", M28, "%", "MY ASSUMPTION. Same basis."),
    ("pe", "Target multiple on 2028E EPS (base)", BASE_PE, "x", "MY ASSUMPTION. In line with Eaton, nVent and Trane at about 28.5 to 29x forward (Peers tab)."),
    ("rnd", "Target price rounding", 10, "US$", "Base value rounded to the nearest US$10 for the published target."),
]
r = 6
for key, label, val, unit, src in rows:
    if label is None:
        ws.cell(row=r, column=1, value=key).font = Font(bold=True, color=NAVY)
        r += 1; continue
    ws.cell(row=r, column=1, value=label)
    c = ws.cell(row=r, column=2, value=val)
    mine = key in ("g27", "g28", "m27", "m28", "pe")
    c.font = F_INB if mine else F_IN
    if mine: c.fill = FILL_ASSUME
    c.number_format = PCT1 if unit == "%" else (MULT0 if unit == "x" and key == "pe" else (MULT if unit == "x" else ("#,##0.0" if unit in ("US$m", "m") else BN2)))
    ws.cell(row=r, column=3, value=unit)
    ws.cell(row=r, column=4, value=src).alignment = WRAP
    A[key] = f"Assumptions!$B${r}"
    r += 1

# =====================================================================  FINANCIALS
ws = sheet("Financials", "FINANCIAL SUMMARY AND FORECASTS (US$ billion unless stated)", [44, 13, 13, 13, 13, 13])
head(ws, 4, ["", "2024A", "2025A", "2026 guide", "2027E*", "2028E*"])
L = lambda k: "=" + A[k]
fin = {}
lines = [
    ("sales", "Net sales", [L("s24"), L("s25"), f"=AVERAGE({A['gs_lo']},{A['gs_hi']})", "=D5*(1+" + A["g27"] + ")", "=E5*(1+" + A["g28"] + ")"], BN2),
    ("growth", "Net sales growth", [None, "=C5/B5-1", "=D5/C5-1", "=E5/D5-1", "=F5/E5-1"], PCT1),
    ("margin", "Adjusted operating margin", [L("m24"), L("m25"), f"=AVERAGE({A['gm_lo']},{A['gm_hi']})", L("m27"), L("m28")], PCT1),
    ("op", "Adjusted operating profit", ["=B5*B7", "=C5*C7", L("gop"), "=E5*E7", "=F5*F7"], BN2),
    ("ratio", "EPS per US$bn of adjusted operating profit", [None, None, f"=AVERAGE({A['ge_lo']},{A['ge_hi']})/D8", "=D9", "=E9"], "0.0000"),
    ("eps", "Adjusted EPS, US$", [L("e24"), L("e25"), f"=AVERAGE({A['ge_lo']},{A['ge_hi']})", "=E8*E9", "=F8*F9"], USD2),
    ("epsg", "Adjusted EPS growth", [None, "=C10/B10-1", "=D10/C10-1", "=E10/D10-1", "=F10/E10-1"], PCT1),
    ("pe", "P/E at reference price", [None, None, "=" + A["price"] + "/D10", "=" + A["price"] + "/E10", "=" + A["price"] + "/F10"], MULT),
]
r = 5
for key, label, vals, fmt in lines:
    ws.cell(row=r, column=1, value=label).font = F_B if key in ("sales", "op", "eps") else Font()
    for j, v in enumerate(vals):
        c = ws.cell(row=r, column=2 + j, value=v if v is not None else "")
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right")
        if isinstance(v, str) and v.startswith("=Assumptions"):
            c.font = F_LINK
    fin[key] = r
    r += 1
notes = [
    "*2027E and 2028E are my own estimates, not Vertiv's. 2026 is the midpoint of Vertiv's own guide of 29 July 2026.",
    "2024A and 2025A adjusted operating profit are derived as sales x margin; they are not separately sourced figures.",
    "EPS bridge: earnings are tied to adjusted operating profit at the ratio implied by the 2026 guide midpoints (US$6.70 EPS per US$3.325bn",
    "of operating profit). This carries the current tax rate, interest cost and share count forward unchanged.",
    "Pitch check: 2027E sales 16.94, EPS 8.53; 2028E sales 20.33, EPS 10.65; P/E 37.9x / 29.7x / 23.8x.",
]
for i, t in enumerate(notes):
    ws.cell(row=14 + i, column=1, value=t).font = F_NOTE

# =====================================================================  VALUATION
ws = sheet("Valuation", "VALUATION: 28x 2028E EPS, TWELVE MONTHS OUT", [52, 16, 70])
head(ws, 4, ["Line", "Value", "Working / note"])
val = [
    ("eps28", "2028E adjusted EPS, US$", "=Financials!F10", USD2, "Financials tab. The market will be pricing 2028 by October 2027."),
    ("mult", "Target multiple, x", "=" + A["pe"], MULT0, "Assumptions tab."),
    ("base", "Base-case value, US$", "=B5*B6", USD2, "EPS x multiple."),
    ("tp", "TARGET PRICE, US$", "=ROUND(B7/" + A["rnd"] + ",0)*" + A["rnd"], USD2, "Base value rounded to the nearest US$10."),
    ("px", "Reference price, US$", "=" + A["price"], USD2, "NYSE close, 5 October 2026."),
    ("up", "Upside to target", "=B8/B9-1", PCT1, "Published as +18%."),
    ("upb", "Upside to unrounded base value", "=B7/B9-1", PCT1, ""),
    ("call", "CALL", '=IF(B10>0,"LONG","NO CALL")', None, "LONG / NO CALL is the site's convention. Conviction Medium (judgement, not formula)."),
    (None, "", None, None, ""),
    ("mv", "Market value, US$bn", "=" + A["shares"] + "*" + A["price"] + "/1000", BN2, "Diluted shares x price."),
    ("nc", "Net cash, US$m", "=" + A["cash"] + "+" + A["stinv"] + "-" + A["debt"], "#,##0.0", "Cash + short-term investments - debt, 30 June 2026."),
    ("ev", "Enterprise value, US$bn", "=B14-B15/1000", BN2, "Market value less net cash."),
    (None, "", None, None, ""),
    ("pe26", "P/E on 2026 guide midpoint", "=" + A["price"] + "/Financials!D10", MULT, "Published as 37.9x."),
    ("pec27", "P/E on 2027 consensus", "=" + A["price"] + "/" + A["cons27"], MULT, "Published as 28.4x (MarketBeat US$8.93)."),
    ("pe27", "P/E on my 2027E", "=" + A["price"] + "/Financials!E10", MULT, "Published as 29.7x."),
    ("pe28", "P/E on my 2028E", "=" + A["price"] + "/Financials!F10", MULT, "Published as 23.8x."),
    ("vc", "My 2027E EPS against consensus", "=Financials!E10/" + A["cons27"] + "-1", PCT1, "Published as about 4.5% below."),
    ("pepk", "P/E at the May peak on that day's 2026 guide", "=" + A["peak"] + "/" + A["gmay"], MULT, "Published as about 60x (intraday peak over the US$6.30 to 6.40 guide midpoint)."),
    ("petp", "Target price on 2027E EPS", "=B8/Financials!E10", MULT, "Cross-check: the target on next year's earnings."),
    ("prem", "Vertiv forward P/E premium to Eaton / nVent / Trane (28.75x)", "=" + A["fwdpe"] + "/Peers!C6-1", PCT1, "stockanalysis.com forward multiples, 6 October 2026."),
]
V = {}
r = 5
for key, label, f, fmt, note in val:
    ws.cell(row=r, column=1, value=label)
    if f is not None:
        c = ws.cell(row=r, column=2, value=f)
        if fmt: c.number_format = fmt
        if key in ("tp", "call"):
            ws.cell(row=r, column=1).font = F_B; c.font = F_B; c.fill = FILL_KEY
    ws.cell(row=r, column=3, value=note).alignment = WRAP
    if key: V[key] = f"Valuation!$B${r}"
    r += 1

# =====================================================================  SCENARIOS
ws = sheet("Scenarios", "THREE CASES, TWELVE MONTHS OUT", [12, 12, 12, 12, 12, 12, 12, 12, 12, 13, 12])
head(ws, 4, ["Case", "2027 growth", "2028 growth", "2028 margin", "2027 sales", "2028 sales", "2028 adj OP", "2028 EPS", "Multiple", "Value, US$", "vs price"])
sc_rows = {}
r = 5
for name, g1, g2, m, mult, p in SCEN:
    ws.cell(row=r, column=1, value=name).font = F_B
    if name == "Base":
        vals = ["=" + A["g27"], "=" + A["g28"], "=" + A["m28"]]
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=2 + j, value=v); c.font = F_LINK; c.number_format = PCT1
        c = ws.cell(row=r, column=9, value="=" + A["pe"]); c.font = F_LINK
    else:
        for j, v in enumerate([g1, g2, m]):
            c = ws.cell(row=r, column=2 + j, value=v); c.font = F_INB; c.fill = FILL_ASSUME; c.number_format = PCT1
        c = ws.cell(row=r, column=9, value=mult); c.font = F_INB; c.fill = FILL_ASSUME
    ws.cell(row=r, column=9).number_format = MULT0
    ws.cell(row=r, column=5, value=f"=Financials!$D$5*(1+B{r})").number_format = BN2
    ws.cell(row=r, column=6, value=f"=E{r}*(1+C{r})").number_format = BN2
    ws.cell(row=r, column=7, value=f"=F{r}*D{r}").number_format = BN2
    ws.cell(row=r, column=8, value=f"=G{r}*Financials!$E$9").number_format = USD2
    ws.cell(row=r, column=10, value=f"=H{r}*I{r}").number_format = USD2
    ws.cell(row=r, column=11, value=f"=J{r}/{A['price']}-1").number_format = PCT0
    sc_rows[name] = r
    r += 1
ws["A9"] = "Probability"; ws["A9"].font = F_B
for j, (name, *_rest) in enumerate(SCEN):
    ws.cell(row=8, column=2 + j, value=name).font = F_B
    c = ws.cell(row=9, column=2 + j, value=SCEN[j][5]); c.font = F_INB; c.fill = FILL_ASSUME; c.number_format = PCT0
ws["E9"] = "=SUM(B9:D9)"; ws["E9"].number_format = PCT0; ws["F9"] = "must be 100%"; ws["F9"].font = F_NOTE
ws["A11"] = "Probability-weighted value, US$"; ws["A11"].font = F_B
ws["E11"] = "=B9*J5+C9*J6+D9*J7"; ws["E11"].number_format = USD2; ws["E11"].font = F_B; ws["E11"].fill = FILL_KEY
ws["F11"] = f"=E11/{A['price']}-1"; ws["F11"].number_format = PCT1
for i, t in enumerate([
    "Published: bear US$8.17 x 22 = ~US$180 (down 29%); base US$10.65 x 28 = ~US$298 (up 18%); bull US$11.52 x 34 = ~US$392 (up 54%);",
    "weighted 25/50/25 on the unrounded cases (179.70 / 298.20 / 391.80) = about US$292, about 15% above the price.",
    "Bear growth of 11% and bull growth of 23% are the midpoints of the",
    "published ranges (10 to 12%, 22 to 24%); they reproduce the published EPS. All cases start from the 2026 guide midpoint of US$14.0bn",
    "and use the same EPS-per-operating-profit ratio as the base case.",
]):
    ws.cell(row=14 + i, column=1, value=t).font = F_NOTE

# =====================================================================  SENSITIVITY
ws = sheet("Sensitivity", "VALUE PER SHARE: 2028E EPS AGAINST MULTIPLE (US$)", [22, 14, 14, 14])
head(ws, 4, ["2028 EPS \\ multiple", "", "", ""])
for j, (pe, link) in enumerate(zip(SENS_PE, [False, True, False])):
    c = ws.cell(row=4, column=2 + j, value=("=" + A["pe"]) if link else pe)
    c.number_format = MULT0; c.font = Font(bold=True, color="FFFFFF"); c.fill = FILL_HDR
    c.alignment = Alignment(horizontal="right")
for i, (eps, link) in enumerate(zip(SENS_EPS, [False, True, False])):
    rr = 5 + i
    c = ws.cell(row=rr, column=1, value="=Financials!F10" if link else eps)
    c.number_format = USD2; c.font = F_LINK if link else F_IN
    for j in range(3):
        col = get_column_letter(2 + j)
        cc = ws.cell(row=rr, column=2 + j, value=f"=$A{rr}*{col}$4"); cc.number_format = USD0
ws["B6"].font = F_B; ws["C6"].fill = FILL_KEY
ws["A9"] = "Cells above the reference price"; ws["A9"].font = F_B
ws["B9"] = f'=COUNTIF(B5:D7,">"&{A["price"]})'; ws["C9"] = "of 9"
for i, t in enumerate([
    "Middle row links to the model's 2028E EPS and the middle column to the base multiple, so the grid moves with the assumptions.",
    "Published grid: 228 / 266 / 304; 256 / 298 / 341; 276 / 322 / 368.",
]):
    ws.cell(row=11 + i, column=1, value=t).font = F_NOTE

# =====================================================================  QUARTERLY
ws = sheet("Quarterly", "QUARTERLY HISTORY, Q1 2024 TO Q2 2026", [46] + [10] * 10)
head(ws, 4, ["US$ billion unless stated"] + Q)
qr = {}
r = 5


def qrow(label, vals, fmt, key, font=None, bold=False):
    global r
    ws.cell(row=r, column=1, value=label).font = F_B if bold else Font()
    for j, v in enumerate(vals):
        c = ws.cell(row=r, column=2 + j, value=v if v is not None else "")
        c.number_format = fmt; c.alignment = Alignment(horizontal="right")
        if font and not (isinstance(v, str) and v.startswith("=")):
            c.font = font
    qr[key] = r; r += 1


cols = [get_column_letter(2 + j) for j in range(10)]
qrow("Net sales", Q_SALES, BN2, "sales", F_IN, True)
qrow("Reported sales growth y/y (derived)", [None] * 4 + [f"={cols[j]}5/{cols[j-4]}5-1" for j in range(4, 10)], PCT1, "g")
qrow("Adjusted operating margin", [m / 100 for m in Q_MARGIN], PCT1, "m", F_IN, True)
qrow("Margin change y/y, points (derived)", [None] * 4 + [f"=({cols[j]}7-{cols[j-4]}7)*100" for j in range(4, 10)], '+0.0;-0.0', "dm")
qrow("Adjusted operating profit (derived, sales x margin)", [f"={c}5*{c}7" for c in cols], "0.000", "op")
qrow("Organic growth, products", [v / 100 for v in Q_PROD], PCT1, "pr", F_IN)
qrow("Organic growth, service and spares", [v / 100 for v in Q_SERV], PCT1, "sv", F_IN)
qrow("Service minus products, points (derived)", [f"=({c}11-{c}10)*100" for c in cols], '+0.0;-0.0', "gap")
qrow("Service grew faster than products? (derived)", [f'=IF({c}11>{c}10,"Yes","No")' for c in cols], "General", "yes")
r += 1
ws.cell(row=r, column=1, value="Annual checks").font = Font(bold=True, color=NAVY); r += 1
ws.cell(row=r, column=1, value="Sum of quarterly sales, 2024 / 2025")
ws.cell(row=r, column=2, value="=SUM(B5:E5)").number_format = BN2
ws.cell(row=r, column=3, value="=SUM(F5:I5)").number_format = BN2
ws.cell(row=r, column=4, value="Annual inputs 8.01 / 10.23; 2025 differs by rounding of quarters.").font = F_NOTE; r += 1
ws.cell(row=r, column=1, value="Quarters since Q3 25 beating prior year margin")
ws.cell(row=r, column=2, value='=COUNTIF(H8:K8,">0")'); ws.cell(row=r, column=3, value="of 4"); r += 1
ws.cell(row=r, column=1, value="Quarters where service grew more slowly than products")
ws.cell(row=r, column=2, value='=COUNTIF(B13:K13,"No")'); ws.cell(row=r, column=3, value="of 10"); r += 2
ws.cell(row=r, column=1, value="Adjusted EPS, orders and backlog").font = Font(bold=True, color=NAVY); r += 1
head(ws, r, ["Item", "Value", "Note"]); r += 1
for item, v, fmt, note in [
    ("Adjusted EPS, 2024 (sum of quarters)", "=" + A["e24"], USD2, "Quarterly adjusted EPS is not compiled in the source pitch; annual figures only."),
    ("Adjusted EPS, 2025 (sum of quarters)", "=" + A["e25"], USD2, ""),
    ("Organic orders growth, Q4 2025", 2.52, "0%", "About 252%. Q4 2025 results release, 11 February 2026."),
    ("Book-to-bill, Q4 2025", 2.9, "0.0", "About 2.9. Q4 2025 results release, cited in the journal piece of 25 September 2026."),
    ("Backlog, 31 December 2025, US$bn", 15.0, BN2, "Up 109% on a year earlier. The 2026 Form 10-K (due February 2027) gives the next figure."),
    ("Backlog growth y/y, 2025", 1.09, "0%", ""),
    ("Orders and backlog, Q1 2026 onwards", "n.d.", "General", "Quarterly disclosure ended; announced on the Q4 2025 call, 11 February 2026. Backlog now annual only."),
    ("Earlier quarterly orders and book-to-bill", "n.c.", "General", "Not compiled in the source pitch; not reproduced here."),
    ("EMEA organic growth, Q2 2026", -0.024, PCT1, "Vertiv Form 10-Q / Q2 2026 release. A weak spot in the evidence."),
]:
    ws.cell(row=r, column=1, value=item)
    c = ws.cell(row=r, column=2, value=v); c.number_format = fmt
    c.font = F_LINK if isinstance(v, str) and v.startswith("=") else F_IN
    c.alignment = Alignment(horizontal="right")
    ws.cell(row=r, column=3, value=note).font = F_NOTE
    r += 1
ws.cell(row=r + 1, column=1, value="Source: Vertiv quarterly results releases on Form 8-K, Exhibit 99.1, 24 April 2024 to 29 July 2026, including the organic growth reconciliation tables.").font = F_NOTE

# =====================================================================  PEERS
ws = sheet("Peers", "PEER MULTIPLES (FORWARD P/E, stockanalysis.com, 6 OCTOBER 2026)", [26, 14, 14, 80])
head(ws, 4, ["Company", "Ticker", "Forward P/E", "What they make"])
peer_vals = [32.1, 28.75, 28.75, 28.75, 24.0]
for i, ((co, tk, pe_txt, what), v) in enumerate(zip(PEERS, peer_vals)):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=co).font = F_B if co == "Vertiv" else Font()
    ws.cell(row=rr, column=2, value=tk)
    c = ws.cell(row=rr, column=3, value=("=" + A["fwdpe"]) if co == "Vertiv" else v)
    c.number_format = MULT; c.font = F_LINK if co == "Vertiv" else F_IN
    ws.cell(row=rr, column=4, value=what).alignment = WRAP
# row 6 used by Valuation premium formula (Eaton)
ws["A11"] = "Vertiv on the 2026 guide midpoint"; ws["C11"] = "=Valuation!B18"; ws["C11"].number_format = MULT
ws["A12"] = "Vertiv on my 2028E (base multiple)"; ws["C12"] = "=Valuation!B6"; ws["C12"].number_format = MULT
for i, t in enumerate([
    "Eaton, nVent and Trane are published as 'about 28.5 to 29x'; the 28.75x midpoint is used here for arithmetic only.",
    "Multiples only, to set the base multiple. No view on the peers' shares is expressed.",
]):
    ws.cell(row=14 + i, column=1, value=t).font = F_NOTE

# =====================================================================  SOURCES
ws = sheet("Sources", "SOURCES", [6, 120])
head(ws, 4, ["No.", "Source and what it supports"])
srcs = [
    "Vertiv quarterly results releases on Form 8-K, Exhibit 99.1, 24 April 2024 to 29 July 2026: sales, organic growth by products and service, adjusted operating profit and margin, adjusted EPS, orders and backlog.",
    "Vertiv Form 10-Q for the quarter to 30 June 2026, filed 29 July 2026: cash, short-term investments, debt, diluted share count, regional sales.",
    "Vertiv 2025 Form 10-K, filed 13 February 2026.",
    "Vertiv 2026 Investor Conference presentation, 19 May 2026: 2030 targets (adjusted operating margin of about 27% or more; 20 to 22% organic growth a year).",
    "Vertiv Form 8-K of 2 September 2026: agreement to buy UtilityInnovation Group for about US$1.45bn in cash plus up to US$1.15bn of earn-out.",
    "Vertiv fourth quarter 2025 earnings call, 11 February 2026: end of quarterly orders and backlog disclosure.",
    "MarketBeat, retrieved 6 October 2026: 2027 consensus adjusted EPS of US$8.93.",
    "stockanalysis.com, retrieved 6 October 2026: forward P/E for Vertiv, Eaton, nVent, Trane Technologies and Schneider Electric.",
    "Yahoo Finance, retrieved 6 October 2026: month-end closes, the 5 October 2026 close and the 14 May 2026 intraday peak.",
    "The Physical Layer, 'Vertiv is paid per megawatt, and AI has made every megawatt harder to build', 25 September 2026: https://thephysicallayer.fyi/journal/every-megawatt-got-harder/",
    "The published pitch and call, 6 October 2026: https://thephysicallayer.fyi/research/",
    "Estimates for 2027 and 2028 are the author's own. Personal research, not investment advice.",
]
for i, s in enumerate(srcs):
    ws.cell(row=5 + i, column=1, value=i + 1)
    ws.cell(row=5 + i, column=2, value=s).alignment = WRAP

# =====================================================================  COVER (filled last)
ws = wb["Cover"]
head(ws, 4, ["Item", "Value", "Note"])
cover = [
    ("Call", "=" + V["call"], None, "Initiate at LONG. Published 6 October 2026, 19:29 HKT."),
    ("Target price, US$", "=" + V["tp"], USD2, "28x 2028E EPS, rounded to the nearest US$10."),
    ("Reference price, US$", "=" + A["price"], USD2, "NYSE close, 5 October 2026."),
    ("Upside to target", "=" + V["up"], PCT1, ""),
    ("Conviction", "Medium", None, "Bear case is a 29% fall and Q3 results land within weeks."),
    ("Horizon", "12 months, to October 2027", None, ""),
    ("2027E / 2028E adjusted EPS, US$", "=TEXT(Financials!E10,\"0.00\")&\" / \"&TEXT(Financials!F10,\"0.00\")", None, "Own estimates."),
    ("Probability-weighted value, US$", "=Scenarios!E11", USD2, "25 / 50 / 25 bear / base / bull."),
    ("Wrong if", "See note", None, "Adjusted operating margin falls year on year in any quarter to mid 2027 while sales are still growing; or full year 2026 adjusted EPS lands below the US$6.65 bottom of Vertiv's own guide; or the 2026 annual report shows backlog below the US$15.0bn of a year earlier."),
]
for i, (lab, v, fmt, note) in enumerate(cover):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=lab).font = F_B
    c = ws.cell(row=rr, column=2, value=v)
    if fmt: c.number_format = fmt
    c.alignment = Alignment(horizontal="right")
    c.font = F_LINK if isinstance(v, str) and v.startswith("=") else F_IN
    ws.cell(row=rr, column=3, value=note).alignment = WRAP
ws.row_dimensions[13].height = 44
ws["A16"] = "Tabs"; ws["A16"].font = Font(bold=True, color=NAVY)
for i, (t, d) in enumerate([
    ("Assumptions", "Every input, with its source. Change the yellow cells to move the target."),
    ("Financials", "2024A to 2028E: sales, margin, operating profit, the EPS bridge, P/E."),
    ("Valuation", "Target price, market value, net cash, multiple cross-checks."),
    ("Scenarios", "Bear / base / bull and the probability-weighted value."),
    ("Sensitivity", "2028E EPS against multiple."),
    ("Quarterly", "Q1 2024 to Q2 2026: sales, margin, organic products and service growth, orders and backlog."),
    ("Peers", "Forward multiples used to set the base multiple."),
    ("Sources", "Every source used."),
]):
    ws.cell(row=17 + i, column=1, value=t)
    ws.cell(row=17 + i, column=2, value=d)
ws["A26"] = "Colour key: blue = hardcoded input; yellow fill = my own assumption; green = link to another tab; black = formula."
ws["A26"].font = F_NOTE
ws["A27"] = "Companion note: 2026-10-06_Vertiv_Initiation.pdf. Personal research, not investment advice."
ws["A27"].font = F_NOTE

for w in wb.worksheets:
    w.freeze_panes = "A5"
wb.active = 0
wb.save(XLSX)
print("model written", XLSX)
