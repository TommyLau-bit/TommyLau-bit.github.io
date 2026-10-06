"""Excel model for the Nebius initiation. Live formulas off blue-font input cells.

Output: public/research/2026-10-06_Nebius_Model.xlsx
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
import nebius_data as D

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(os.path.join(HERE, "..", "..", "public", "research", "2026-10-06_Nebius_Model.xlsx"))

NAVY = "1F3864"
BLUE = "0000FF"
AMBER = "FFF2CC"
LIGHT = "D9E2F3"
hdr = Font(bold=True, color="FFFFFF", size=10)
navyfill = PatternFill("solid", fgColor=NAVY)
ttl = Font(bold=True, color=NAVY, size=12)
bold = Font(bold=True)
inp_font = Font(color=BLUE)
inp_fill = PatternFill("solid", fgColor=AMBER)
out_fill = PatternFill("solid", fgColor=LIGHT)
muted = Font(italic=True, color="595959", size=9)
thin = Side(style="thin", color="BFBFBF")
navyside = Side(style="thin", color=NAVY)

wb = openpyxl.Workbook()
REF = {}

F_USD2 = '"US$"#,##0.00'
F_USD0 = '"US$"#,##0'
F_M1 = '#,##0.0'
F_M2 = '#,##0.00'
F_INT = '#,##0'
F_PCT0 = '0%'
F_PCT1 = '0.0%'
F_X1 = '0.0"x"'


def sheet(name, title, widths, sub=None):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = ttl
    if sub:
        ws["A2"] = sub
        ws["A2"].font = muted
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def head(ws, row, cells, col=1):
    for i, c in enumerate(cells):
        x = ws.cell(row=row, column=col + i, value=c)
        x.font = hdr
        x.fill = navyfill
        x.alignment = Alignment(horizontal="left" if i == 0 else "right", vertical="center", wrap_text=True)


def put(ws, r, c, v, fmt=None, inp=False, b=False, wrap=False):
    x = ws.cell(row=r, column=c, value=v)
    if fmt:
        x.number_format = fmt
    if inp:
        x.font = Font(color=BLUE, bold=b)
        x.fill = inp_fill
    elif b:
        x.font = bold
    if isinstance(v, str) and v.startswith("=") and not inp:
        x.fill = out_fill if b else PatternFill()
    if wrap:
        x.alignment = Alignment(wrap_text=True, vertical="top")
    x.border = Border(bottom=thin)
    return x


def total_row(ws, r, c1, c2):
    for c in range(c1, c2 + 1):
        x = ws.cell(row=r, column=c)
        x.font = Font(bold=True, color=x.font.color.rgb if x.font and x.font.color and x.font.color.rgb == "FF0000FF" else None)
        x.border = Border(top=navyside, bottom=thin)


# ======================= COVER =======================
ws = wb.active
ws.title = "Cover"
ws["A1"] = "NEBIUS GROUP N.V. (Nasdaq: NBIS)  INITIATION MODEL, 6 OCTOBER 2026"
ws["A1"].font = ttl
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 96
ws.sheet_view.showGridLines = False
rows = [
    ("Analyst", "Tommy Lau"),
    ("Call", "INITIATE AT NO CALL (no position at this price). Conviction: Medium. Horizon: 12 months, to October 2027."),
    ("Reference price", "US$232.57, Nasdaq close 5 October 2026"),
    ("Target price", "None. The note states what would turn it into a call."),
    ("Re-look", "A price below about US$185, or Q4 2026 results (early 2027) showing at least 800 MW connected, year-end ARR inside US$7 to 9bn and new deals still at US$20m per MW or more."),
    ("Wrong if", "Connected power at end 2026 below 800 MW, or year-end ARR below US$7bn, or prepayments in fewer than half of new deals."),
    ("Companion note", "2026-10-06_Nebius_Initiation.pdf"),
    ("Published pitch", "https://thephysicallayer.fyi/research/"),
    ("Journal piece", "https://thephysicallayer.fyi/journal/nebius-paid-for-what-is-switched-on/"),
    ("", ""),
    ("Colour key", "Blue font on amber = input you can change. Black font = live formula. Light blue fill = key output."),
    ("Units", "US$ millions unless stated. Shares in millions in the bridge. Per MW figures in US$ millions."),
    ("", ""),
    ("Tabs", ""),
    ("Inputs", "Sourced figures and my assumptions, each with its source."),
    ("Quarterly", "Q1 2025 to Q2 2026 history, 2026 guide and consensus 2027."),
    ("ARR_Recon", "Why US$20 to 25m per MW and a US$7 to 9bn run-rate are both true: the connected vs active lag."),
    ("UnitEcon", "What a megawatt earns at the old and new contract prices."),
    ("CapStructure", "Convertible notes, if-converted share bridge, equity value and enterprise value."),
    ("Valuation", "EV multiples on consensus 2027, peer multiples, and the 5 GW build against market value."),
    ("Scenarios", "Bear, base and bull for 2028, valued in October 2027, and the probability-weighted value."),
    ("Sensitivity", "2028 EBITDA against the multiple, live grid."),
    ("Sources", "Every source used."),
    ("", ""),
    ("Disclaimer", "Personal research, not investment advice."),
]
for i, (a, b) in enumerate(rows, 3):
    ws.cell(row=i, column=1, value=a).font = bold
    c = ws.cell(row=i, column=2, value=b)
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws["B4"].font = Font(bold=True, color=NAVY)

# ======================= INPUTS =======================
ws = sheet("Inputs", "INPUTS", [34, 16, 14, 86],
           "Blue font on amber = input. Change any of these and every other tab recalculates.")
head(ws, 4, ["Input", "Value", "Unit", "Source / basis"])
inputs = [
    ("PRICING", None, None, None),
    ("price", "Share price, reference", D.PRICE, "US$", "Nasdaq close, 5 Oct 2026. Yahoo Finance, retrieved 6 Oct 2026.", F_USD2),
    ("relook", "Re-look price", D.RELOOK, "US$", "My level from the published pitch: base case offers about a quarter of upside here.", F_USD2),
    ("GUIDANCE AND CONSENSUS", None, None, None),
    ("arr_lo", "Year-end 2026 ARR guide, low", D.ARR_GUIDE[0], "US$bn", "Nebius Q2 2026 shareholder letter, 12 Aug 2026.", F_M2),
    ("arr_hi", "Year-end 2026 ARR guide, high", D.ARR_GUIDE[1], "US$bn", "As above.", F_M2),
    ("conn_lo", "Connected power guide end 2026, low", D.CONNECTED_GUIDE[0], "MW", "As above. Connected = wired into fully built and equipped data centres.", F_INT),
    ("conn_hi", "Connected power guide end 2026, high", D.CONNECTED_GUIDE[1], "MW", "As above.", F_INT),
    ("capex_lo", "2026 capital spending guide, low", D.CAPEX_GUIDE[0], "US$bn", "As above.", F_M1),
    ("capex_hi", "2026 capital spending guide, high", D.CAPEX_GUIDE[1], "US$bn", "As above.", F_M1),
    ("cons26", "Consensus revenue 2026", D.CONS_REV_26, "US$bn", "Yahoo Finance consensus, retrieved 6 Oct 2026.", F_M2),
    ("cons27", "Consensus revenue 2027", D.CONS_REV_27, "US$bn", "As above.", F_M2),
    ("acv_base", "Annual contract value per MW, 2026 fleet base", D.ACV_BASE, "US$m/MW", "Nebius Q2 2026 letter, 'ACV per MW is stepping up'. Approximate, revenue recognition basis, excl. prepayments.", F_M1),
    ("acv_new_lo", "ACV per MW, Q2 2026 large deals, low", D.ACV_NEW[0], "US$m/MW", "As above.", F_M1),
    ("acv_new_hi", "ACV per MW, Q2 2026 large deals, high", D.ACV_NEW[1], "US$m/MW", "As above.", F_M1),
    ("acv_q3", "ACV per MW, Q3 2026 short-term deals (floor)", 40.0, "US$m/MW", "As above. Given as above US$40m; treated as a floor.", F_M1),
    ("arr_ye25", "ARR at end 2025", D.ARR_YE25, "US$bn", "Nebius Q4 2025 shareholder letter.", F_M2),
    ("active_ye25", "Active power at end 2025, approx", D.ACTIVE_YE25, "MW", "Nebius shareholder letters; approximate.", F_INT),
    ("pipeline", "Power pipeline", D.PIPELINE_GW, "GW", "Nebius, as described in the 30 Sep 2026 journal piece and the pitch.", F_M1),
    ("UNIT ECONOMICS (MY ASSUMPTIONS)", None, None, None),
    ("capex_mw", "Capital cost per MW", D.CAPEX_PER_MW, "US$m", "MY ESTIMATE: 2026 capex of US$20 to 25bn against ~600 to 800 MW of new connected power, some of it 2027 spend.", F_M1),
    ("chip_share", "Share of cost that is chips and network", D.CHIP_SHARE, "%", "MY ASSUMPTION.", F_PCT0),
    ("chip_life", "Chip and network life", D.CHIP_LIFE, "years", "MY ASSUMPTION. Five years, because a hall of older chips earns less each year.", F_INT),
    ("bldg_life", "Building life", D.BLDG_LIFE, "years", "MY ASSUMPTION.", F_INT),
    ("margin", "EBITDA margin on contract value", D.EBITDA_MARGIN, "%", "MY ASSUMPTION, close to the 49.7% AI cloud adj. EBITDA margin reported for Q2 2026.", F_PCT1),
    ("SHARES, CASH AND DEBT (FILINGS)", None, None, None),
    ("sh_a", "Class A shares, 30 Jun 2026", D.SH_A / 1e6, "m", "Nebius interim financial statements to 30 Jun 2026: 238,400,165.", '#,##0.000'),
    ("sh_b", "Class B shares, 30 Jun 2026", D.SH_B / 1e6, "m", "As above: 33,455,053.", '#,##0.000'),
    ("nvda", "Nvidia pre-funded warrants", D.NVDA_WARRANTS / 1e6, "m shares", "As above: 21,065,936 shares.", '#,##0.000'),
    ("aug_exch", "Shares issued ~24 Aug 2026 for US$800m of 2029/2031 notes", D.AUG_EXCHANGE_SH / 1e6, "m", "Nebius release, 24 Aug 2026. Approximate (~15.8m).", '#,##0.0'),
    ("options", "Options outstanding", D.OPTIONS / 1e6, "m", "Interim statements: 6,740,600.", '#,##0.000'),
    ("waep", "Options weighted average exercise price", D.OPT_WAEP, "US$", "As above.", F_USD2),
    ("rsus", "Restricted stock units", D.RSUS / 1e6, "m", "As above: 6,234,091.", '#,##0.000'),
    ("cash", "Cash and equivalents, 30 Jun 2026", D.CASH, "US$m", "As above.", F_M1),
    ("restricted", "Restricted cash, 30 Jun 2026 (excluded from net debt)", D.RESTRICTED, "US$m", "As above.", F_M1),
    ("secured", "Secured facility, SOFR + 2.50% to Oct 2030", D.SECURED, "US$m", "Nebius release, 10 Jul 2026. Approximate.", F_M1),
    ("cash_pf", "Pro forma cash before Q3 spending", D.CASH_PF, "US$m", "MY ESTIMATE, from the pitch: 30 Jun cash plus the Aug 2026 notes and the secured facility, before fees. See CapStructure cross-check.", F_M1),
    ("def_cur", "Deferred revenue, current", D.DEF_REV_CUR, "US$m", "Interim statements, 30 Jun 2026. Mostly customer prepayments.", F_M1),
    ("def_nc", "Deferred revenue, non-current", D.DEF_REV_NC, "US$m", "As above.", F_M1),
    ("rpo", "Remaining performance obligations", D.RPO, "US$bn", "As above.", F_M2),
    ("atm_sh", "ATM shares sold May to Jun 2026", D.ATM_SH / 1e6, "m", "Nebius filings: 12,729,493 shares.", '#,##0.000'),
    ("atm_avg", "ATM average price", D.ATM_AVG, "US$", "As above.", F_USD2),
    ("atm_net", "ATM net proceeds", D.ATM_NET, "US$bn", "As above.", F_M2),
    ("MARKET AND CUSTOMERS", None, None, None),
    ("short_sh", "Short interest, 15 Sep 2026", D.SHORT_SH / 1e6, "m shares", "Yahoo Finance, retrieved 6 Oct 2026.", F_M2),
    ("short_pct", "Short interest, share of float", D.SHORT_PCT / 100, "%", "As above.", F_PCT1),
    ("cust1", "Largest customer, share of Q2 2026 revenue", 0.24, "%", "Nebius Q2 2026 results, Form 6-K. Customers unnamed.", F_PCT0),
    ("cust2", "Second customer", 0.21, "%", "As above.", F_PCT0),
    ("cust3", "Third customer", 0.14, "%", "As above.", F_PCT0),
]
r = 5
for row in inputs:
    if row[1] is None:
        c = ws.cell(row=r, column=1, value=row[0])
        c.font = Font(bold=True, color=NAVY)
        r += 1
        continue
    key, lbl, val, unit, src, fmt = row
    ws.cell(row=r, column=1, value=lbl).border = Border(bottom=thin)
    put(ws, r, 2, val, fmt, inp=True)
    ws.cell(row=r, column=3, value=unit).border = Border(bottom=thin)
    s = ws.cell(row=r, column=4, value=src)
    s.alignment = Alignment(wrap_text=True, vertical="top")
    s.border = Border(bottom=thin)
    REF[key] = f"Inputs!$B${r}"
    r += 1
cust_total_r = r
ws.cell(row=r, column=1, value="Top three customers, combined").font = bold
put(ws, r, 2, f"={REF['cust1'].split('!')[1]}+{REF['cust2'].split('!')[1]}+{REF['cust3'].split('!')[1]}", F_PCT0, b=True)
REF["cust_total"] = f"Inputs!$B${r}"
ws.freeze_panes = "A5"

# ======================= QUARTERLY =======================
ws = sheet("Quarterly", "QUARTERLY HISTORY, 2026 GUIDE AND CONSENSUS 2027", [40, 11, 11, 11, 11, 11, 11, 15, 15],
           "US$ millions unless stated. History is input (blue). Guide and consensus columns are flagged and not my forecasts.")
head(ws, 4, ["", "Q1 25", "Q2 25", "Q3 25", "Q4 25", "Q1 26", "Q2 26", "2026 guide / cons.", "2027 consensus"])
put(ws, 5, 1, "Revenue")
for i, v in enumerate(D.REVENUE):
    put(ws, 5, 2 + i, v, F_M1, inp=True)
put(ws, 5, 8, f"={REF['cons26']}*1000", F_INT)
put(ws, 5, 9, f"={REF['cons27']}*1000", F_INT)
put(ws, 6, 1, "Revenue growth, quarter on quarter")
for i in range(1, 6):
    col = get_column_letter(2 + i)
    prev = get_column_letter(1 + i)
    put(ws, 6, 2 + i, f"={col}5/{prev}5-1", F_PCT0)
put(ws, 6, 9, "=I5/H5-1", F_PCT0)
put(ws, 7, 1, "ARR at quarter end, US$bn")
for i, v in enumerate(D.ARR):
    put(ws, 7, 2 + i, v, F_M2, inp=True)
put(ws, 7, 8, "7.00 to 9.00")
ws.cell(row=7, column=8).alignment = Alignment(horizontal="right")
put(ws, 8, 1, "ARR growth, quarter on quarter")
for i in range(1, 6):
    col = get_column_letter(2 + i)
    prev = get_column_letter(1 + i)
    put(ws, 8, 2 + i, f"={col}7/{prev}7-1", F_PCT0)
put(ws, 9, 1, "Capital spending")
for i, v in enumerate(D.CAPEX):
    put(ws, 9, 2 + i, v, F_M1, inp=True)
put(ws, 9, 8, "20,000 to 25,000")
ws.cell(row=9, column=8).alignment = Alignment(horizontal="right")
put(ws, 10, 1, "Capital spending per US$1 of revenue")
for i in range(6):
    col = get_column_letter(2 + i)
    put(ws, 10, 2 + i, f"={col}9/{col}5", '0.0"x"')
put(ws, 11, 1, "AI cloud adj. EBITDA margin")
for i, v in enumerate(D.AI_MARGIN):
    if v is None:
        put(ws, 11, 2 + i, "n.d.")
        ws.cell(row=11, column=2 + i).alignment = Alignment(horizontal="right")
    else:
        put(ws, 11, 2 + i, v / 100, F_PCT1, inp=True)
put(ws, 12, 1, "Connected power, end of period, MW")
put(ws, 12, 8, "800 to 1,000")
ws.cell(row=12, column=8).alignment = Alignment(horizontal="right")
put(ws, 13, 1, "Active power, end of period, MW (approx)")
put(ws, 13, 5, f"={REF['active_ye25']}", F_INT)
put(ws, 15, 1, "Second half 2026: ARR multiple needed to reach guide low", b=True)
put(ws, 15, 2, f"={REF['arr_lo']}/G7", '0.00"x"', b=True)
put(ws, 16, 1, "Second half 2026: ARR multiple needed to reach guide high")
put(ws, 16, 2, f"={REF['arr_hi']}/G7", '0.00"x"')
put(ws, 17, 1, "H1 2026 capital spending")
put(ws, 17, 2, "=F9+G9", F_INT)
put(ws, 18, 1, "Implied H2 2026 capital spending, low")
put(ws, 18, 2, f"={REF['capex_lo']}*1000-B17", F_INT)
put(ws, 19, 1, "Implied H2 2026 capital spending, high")
put(ws, 19, 2, f"={REF['capex_hi']}*1000-B17", F_INT)
put(ws, 20, 1, "Consensus 2027 revenue / ARR guide high")
put(ws, 20, 2, f"={REF['cons27']}/{REF['arr_hi']}", '0.00"x"')
notes = [
    "Sources: Nebius quarterly shareholder letters and results on Form 6-K, Q2 2025 to Q2 2026. Before 2026, ARR covers the core AI infrastructure business.",
    "Q1 2025 capital spending is derived from the first half less the second quarter. Q4 2025, Q1 2026 and Q2 2026 capital spending are approximate, as reported.",
    "2026 guide (ARR, connected power, capital spending): Q2 2026 letter, 12 Aug 2026. 2026 and 2027 revenue: Yahoo Finance consensus, 6 Oct 2026. Neither is my forecast.",
    "n.d. = not disclosed on a comparable basis. Implied H2 capex of US$12 to 17bn matches the pitch (guide less H1 of ~US$8.2bn).",
]
for i, t in enumerate(notes):
    c = ws.cell(row=22 + i, column=1, value=t)
    c.font = muted
ws.freeze_panes = "B5"

# ======================= ARR_RECON =======================
ws = sheet("ARR_Recon", "THE RECONCILIATION: CONTRACT VALUE PER MW AGAINST THE RUN-RATE GUIDE", [62, 16, 16, 50],
           "If new deals pay US$20 to 25m per MW, why is the year-end run-rate only US$7 to 9bn? Because much connected power is not yet billing.")
head(ws, 4, ["Step", "Low", "High", "Working"])
put(ws, 5, 1, "Connected power guide, end 2026, MW")
put(ws, 5, 2, f"={REF['conn_lo']}", F_INT); put(ws, 5, 3, f"={REF['conn_lo']}", F_INT)
put(ws, 5, 4, "Bottom of the 800 MW to 1 GW range, used for both columns")
put(ws, 6, 1, "New-deal contract value per MW, US$m")
put(ws, 6, 2, f"={REF['acv_new_lo']}", F_M1); put(ws, 6, 3, f"={REF['acv_new_hi']}", F_M1)
put(ws, 6, 4, "Q2 2026 large deals")
put(ws, 7, 1, "Naive annual revenue if all connected MW billed at new prices, US$bn", b=True)
put(ws, 7, 2, "=B5*B6/1000", F_M1, b=True); put(ws, 7, 3, "=C5*C6/1000", F_M1, b=True)
put(ws, 7, 4, "800 x US$20 to 25m = US$16 to 20bn")
put(ws, 8, 1, "Year-end 2026 ARR guide, US$bn")
put(ws, 8, 2, f"={REF['arr_lo']}", F_M1); put(ws, 8, 3, f"={REF['arr_hi']}", F_M1)
put(ws, 8, 4, "ARR = December revenue x 12, so only what bills that month")
put(ws, 9, 1, "Gap, US$bn")
put(ws, 9, 2, "=B7-B8", F_M1); put(ws, 9, 3, "=C7-C8", F_M1)
put(ws, 11, 1, "CLOSING THE GAP", b=True)
put(ws, 12, 1, "Fleet base contract value per MW, US$m")
put(ws, 12, 2, f"={REF['acv_base']}", F_M1); put(ws, 12, 3, f"={REF['acv_base']}", F_M1)
put(ws, 12, 4, "Most new deals earn mainly in 2027, so December 2026 bills at the fleet base")
put(ws, 13, 1, "MW billing in December 2026 implied by the guide", b=True)
put(ws, 13, 2, "=B8*1000/B12", F_INT, b=True); put(ws, 13, 3, "=C8*1000/C12", F_INT, b=True)
put(ws, 13, 4, "US$7 to 9bn / US$12m = ~580 to 750 MW")
put(ws, 14, 1, "Connected but not yet billing, MW (vs 800 MW connected)")
put(ws, 14, 2, "=B5-B13", F_INT); put(ws, 14, 3, "=C5-C13", F_INT)
put(ws, 14, 4, "The lag of several months between connected and active power")
put(ws, 15, 1, "Billing MW as share of 800 MW connected")
put(ws, 15, 2, "=B13/B5", F_PCT0); put(ws, 15, 3, "=C13/C5", F_PCT0)
put(ws, 17, 1, "CROSS-CHECK: END 2025", b=True)
put(ws, 18, 1, "ARR at end 2025, US$bn")
put(ws, 18, 2, f"={REF['arr_ye25']}", F_M2)
put(ws, 19, 1, "Active power at end 2025, MW (approx)")
put(ws, 19, 2, f"={REF['active_ye25']}", F_INT)
put(ws, 20, 1, "ARR per active MW, US$m", b=True)
put(ws, 20, 2, "=B18*1000/B19", F_M2, b=True)
put(ws, 20, 4, "~US$7.4m, a similar ratio to the 2026 guide on a connected basis")
put(ws, 22, 1, "Sources: Nebius Q2 2026 shareholder letter, 12 Aug 2026 (guide, ACV per MW, definitions of connected and active power); Q2 2026 call (deployment sequence). Working is mine.").font = muted

# ======================= UNIT ECON =======================
ws = sheet("UnitEcon", "UNIT ECONOMICS: WHAT A MEGAWATT EARNS", [44, 16, 16, 16, 40],
           "My estimates, built on Nebius's figures. US$ millions per MW per year.")
head(ws, 4, ["Cost build", "Value", "", "", "Working"])
put(ws, 5, 1, "Capital cost per MW"); put(ws, 5, 2, f"={REF['capex_mw']}", F_M1)
put(ws, 6, 1, "Chips and network share"); put(ws, 6, 2, f"={REF['chip_share']}", F_PCT0)
put(ws, 7, 1, "Annual wear on chips and network"); put(ws, 7, 2, f"=B5*B6/{REF['chip_life']}", F_M2)
put(ws, 7, 5, "US$24m over 5 years")
put(ws, 8, 1, "Annual wear on building"); put(ws, 8, 2, f"=B5*(1-B6)/{REF['bldg_life']}", F_M2)
put(ws, 8, 5, "US$6m over 20 years")
put(ws, 9, 1, "Total annual wear (depreciation) per MW", b=True); put(ws, 9, 2, "=B7+B8", F_M2, b=True)
put(ws, 10, 1, "EBITDA margin"); put(ws, 10, 2, f"={REF['margin']}", F_PCT1)
head(ws, 12, ["Contract value per MW", "ACV", "EBITDA per MW", "Profit after wear", "Pre-tax return on capital"])
cases = [("US$12m, the 2026 fleet", "acv_base"), ("US$20m, new deals, low", "acv_new_lo"), ("US$25m, new deals, high", "acv_new_hi")]
for i, (lbl, k) in enumerate(cases):
    rr = 13 + i
    put(ws, rr, 1, lbl)
    put(ws, rr, 2, f"={REF[k]}", F_M1)
    put(ws, rr, 3, f"=B{rr}*$B$10", F_M2)
    put(ws, rr, 4, f"=C{rr}-$B$9", F_M2, b=True)
    put(ws, rr, 5, f"=D{rr}/$B$5", F_PCT0, b=True)
put(ws, 17, 1, "Pitch figures: EBITDA 6.0 / 10.0 / 12.5; profit after wear 0.9 / 4.9 / 7.4; return 3% / 16% / 25%.").font = muted
put(ws, 18, 1, "At the old price a megawatt barely earns back its cost before the chips age out. At the new price it earns a good return.").font = muted
put(ws, 20, 1, "BREAK-EVEN AND THE 5 GW BUILD", b=True)
put(ws, 21, 1, "Contract value per MW for zero profit after wear"); put(ws, 21, 2, "=B9/B10", F_M2)
put(ws, 22, 1, "Contract value per MW for a 10% pre-tax return"); put(ws, 22, 2, "=(B9+0.1*B5)/B10", F_M2)
put(ws, 23, 1, "Power pipeline, MW"); put(ws, 23, 2, f"={REF['pipeline']}*1000", F_INT)
put(ws, 24, 1, "Capital to build the pipeline, US$bn", b=True); put(ws, 24, 2, "=B23*B5/1000", F_M1, b=True)
put(ws, 25, 1, "Equity value, if-converted, US$bn"); put(ws, 25, 2, "=CapStructure!C33/1000", F_M1)
put(ws, 26, 1, "Build cost / equity value"); put(ws, 26, 2, "=B24/B25", '0.0"x"')

# ======================= CAP STRUCTURE =======================
ws = sheet("CapStructure", "CAPITAL STRUCTURE, IF-CONVERTED SHARE BRIDGE AND ENTERPRISE VALUE", [30, 14, 10, 14, 12, 12, 16, 30],
           "Notes in the money at the reference price are counted as shares; the rest stay as debt. US$ millions; shares in millions.")
head(ws, 4, ["Convertible notes", "Principal", "Coupon", "Conversion price", "Maturity", "In the money?", "Shares if converted", "Treatment"])
r0 = 5
for i, (name, p, c, cp, mat, note) in enumerate(D.CONVERTS):
    rr = r0 + i
    put(ws, rr, 1, name + (f" ({note})" if note else ""))
    put(ws, rr, 2, p, '#,##0.00', inp=True)
    put(ws, rr, 3, c, '0.000%', inp=True)
    put(ws, rr, 4, cp, F_USD2, inp=True)
    put(ws, rr, 5, mat, inp=True)
    put(ws, rr, 6, f'=IF(D{rr}<{REF["price"]},"Yes","No")')
    put(ws, rr, 7, f"=B{rr}/D{rr}", '#,##0.000')
    put(ws, rr, 8, f'=IF(F{rr}="Yes","Counted as shares","Stays as debt")')
rl = r0 + len(D.CONVERTS) - 1
rt = rl + 1
put(ws, rt, 1, "Total convertible notes", b=True); put(ws, rt, 2, f"=SUM(B{r0}:B{rl})", '#,##0.00', b=True)
put(ws, rt, 7, f'=SUMIF(F{r0}:F{rl},"Yes",G{r0}:G{rl})', '#,##0.000', b=True)
put(ws, rt, 8, "Shares from in-the-money notes")
total_row(ws, rt, 1, 8)
# bridge
head(ws, 16, ["Share bridge", "", "Shares, m", "", "", "", "", "Working"])
bridge = [
    (17, "Class A shares, 30 Jun 2026", f"={REF['sh_a']}", ""),
    (18, "Class B shares, 30 Jun 2026", f"={REF['sh_b']}", ""),
    (19, "Basic shares, 30 Jun 2026", "=C17+C18", "271,855,218"),
    (20, "Nvidia pre-funded warrants", f"={REF['nvda']}", "Near-zero exercise price, counted in full"),
    (21, "Shares issued ~24 Aug 2026 for 2029/2031 notes", f"={REF['aug_exch']}", "US$800m of notes exchanged"),
    (22, "Options, treasury method", f"={REF['options']}*MAX(0,1-{REF['waep']}/{REF['price']})", "6.74m at WAEP US$88.65"),
    (23, "Restricted stock units", f"={REF['rsus']}", "Counted in full"),
    (24, "In-the-money convertible notes", f"=G{rt}", "Six notes below the price"),
]
for rr, lbl, f, w in bridge:
    put(ws, rr, 1, lbl); put(ws, rr, 3, f, '#,##0.0'); put(ws, rr, 8, w)
put(ws, 25, 1, "If-converted shares", b=True); put(ws, 25, 3, "=SUM(C19:C24)", '#,##0.0', b=True)
put(ws, 25, 8, "Pitch: ~370m (369.6m)")
total_row(ws, 25, 1, 8)
# EV
head(ws, 27, ["Enterprise value", "", "US$m", "", "", "", "", "Working"])
put(ws, 28, 1, "Reference price, US$"); put(ws, 28, 3, f"={REF['price']}", F_USD2)
put(ws, 29, 1, "If-converted shares, m"); put(ws, 29, 3, "=C25", '#,##0.0')
put(ws, 30, 1, "Notes that stay as debt"); put(ws, 30, 3, f'=SUMIF(F{r0}:F{rl},"No",B{r0}:B{rl})', F_INT)
put(ws, 30, 8, "The two Aug 2026 notes, US$5.75bn")
put(ws, 31, 1, "Secured facility"); put(ws, 31, 3, f"={REF['secured']}", F_INT)
put(ws, 32, 1, "Debt staying as debt"); put(ws, 32, 3, "=C30+C31", F_INT); put(ws, 32, 8, "Pitch: ~US$6.5bn")
put(ws, 33, 1, "Equity value, if-converted", b=True); put(ws, 33, 3, "=C28*C29", F_INT, b=True)
put(ws, 33, 8, "Pitch: ~US$86bn")
put(ws, 34, 1, "Less: pro forma cash"); put(ws, 34, 3, f"=-{REF['cash_pf']}", F_INT); put(ws, 34, 8, "My estimate, ~US$14.5bn")
put(ws, 35, 1, "Plus: debt staying as debt"); put(ws, 35, 3, "=C32", F_INT)
put(ws, 36, 1, "Enterprise value", b=True); put(ws, 36, 3, "=C33+C34+C35", F_INT, b=True); put(ws, 36, 8, "Pitch: ~US$78bn")
total_row(ws, 36, 1, 8)
head(ws, 38, ["Pro forma cash cross-check", "", "US$m", "", "", "", "", "Working"])
put(ws, 39, 1, "Cash and equivalents, 30 Jun 2026"); put(ws, 39, 3, f"={REF['cash']}", F_INT)
put(ws, 40, 1, "Aug 2026 notes, principal"); put(ws, 40, 3, "=C30", F_INT)
put(ws, 41, 1, "Secured facility"); put(ws, 41, 3, f"={REF['secured']}", F_INT)
put(ws, 42, 1, "Sum before fees"); put(ws, 42, 3, "=SUM(C39:C41)", F_INT); put(ws, 42, 8, "Consistent with ~US$14.5bn after fees")
put(ws, 43, 1, "Restricted cash (excluded)"); put(ws, 43, 3, f"={REF['restricted']}", F_INT)
head(ws, 45, ["Other balance sheet items", "", "Value", "", "", "", "", "Working"])
put(ws, 46, 1, "Deferred revenue, total, US$m"); put(ws, 46, 3, f"={REF['def_cur']}+{REF['def_nc']}", F_INT)
put(ws, 46, 8, "Customer prepayments; funds part of the build")
put(ws, 47, 1, "Remaining performance obligations, US$bn"); put(ws, 47, 3, f"={REF['rpo']}", F_M2)
put(ws, 48, 1, "ATM May to Jun 2026, gross at average price, US$bn"); put(ws, 48, 3, f"={REF['atm_sh']}*{REF['atm_avg']}/1000", F_M2)
put(ws, 48, 8, "Net proceeds US$2.81bn")
put(ws, 49, 1, "Short interest as share of if-converted count"); put(ws, 49, 3, f"={REF['short_sh']}/C25", F_PCT1)
put(ws, 49, 8, "~19.8% of float on Yahoo's float measure")
put(ws, 51, 1, "Sources: Nebius interim financial statements to 30 Jun 2026; Nebius releases of 10 Jul 2026 (secured loan) and 19 and 24 Aug 2026 (convertible notes). Bridge, cash estimate and treatment are mine.").font = muted

# ======================= VALUATION =======================
ws = sheet("Valuation", "WHAT THE MARKET PRICES IN: EV MULTIPLES ON CONSENSUS 2027", [56, 16, 16, 48])
head(ws, 4, ["Item", "Value", "", "Working"])
put(ws, 5, 1, "Enterprise value, US$bn"); put(ws, 5, 2, "=CapStructure!C36/1000", F_M1)
put(ws, 6, 1, "Consensus revenue 2027, US$bn"); put(ws, 6, 2, f"={REF['cons27']}", F_M2)
put(ws, 7, 1, "EBITDA margin applied"); put(ws, 7, 2, f"={REF['margin']}", F_PCT0)
put(ws, 8, 1, "2027 EBITDA, US$bn, unrounded"); put(ws, 8, 2, "=B6*B7", F_M2)
put(ws, 9, 1, "2027 EBITDA, US$bn, as stated in the pitch (rounded)"); put(ws, 9, 2, "=ROUND(B8,1)", F_M1)
put(ws, 10, 1, "EV / 2027 EBITDA, as in the pitch", b=True); put(ws, 10, 2, "=B5/B9", F_X1, b=True)
put(ws, 10, 4, "Pitch: about 12.6x on ~US$6.2bn")
put(ws, 11, 1, "EV / 2027 EBITDA, on unrounded EBITDA"); put(ws, 11, 2, "=B5/B8", F_X1)
put(ws, 11, 4, "Same thing before rounding; shown for transparency")
put(ws, 12, 1, "EV / 2027 revenue", b=True); put(ws, 12, 2, "=B5/B6", F_X1, b=True)
put(ws, 12, 4, "Pitch: 6.3x")
put(ws, 13, 1, "EV / 2026 consensus revenue"); put(ws, 13, 2, f"=B5/{REF['cons26']}", F_X1)
put(ws, 14, 1, "Consensus 2027 revenue / end-2026 ARR guide, low to high")
put(ws, 14, 2, f"={REF['cons27']}/{REF['arr_lo']}", '0.00"x"'); put(ws, 14, 3, f"={REF['cons27']}/{REF['arr_hi']}", '0.00"x"')
put(ws, 14, 4, "Consensus assumes the run-rate roughly doubles again through 2027")
head(ws, 16, ["Peer landscape", "Multiple", "Premium of NBIS", "Basis (Yahoo Finance, 6 Oct 2026)"])
peer_rows = [("Nebius (my EV)", "=B12", "EV / 2027 revenue"),
             ("CoreWeave", 3.6, "EV / 2027 revenue"),
             ("Oracle", 4.3, "Next financial year revenue"),
             ("IREN", 2.4, "Next financial year revenue")]
for i, (n, v, basis) in enumerate(peer_rows):
    rr = 17 + i
    put(ws, rr, 1, n)
    put(ws, rr, 2, v, F_X1, inp=not str(v).startswith("="))
    put(ws, rr, 3, "" if i == 0 else f"=$B$17/B{rr}-1", F_PCT0)
    put(ws, rr, 4, basis)
put(ws, 22, 1, "Peer multiples are on different bases and fiscal years; they frame the premium, they are not a valuation input.").font = muted

# ======================= SCENARIOS =======================
ws = sheet("Scenarios", "THREE CASES, TWELVE MONTHS OUT: 2028 EBITDA AT A MULTIPLE, LESS END-2027 NET DEBT", [10, 46, 12, 10, 10, 12, 10, 10, 12, 12, 12],
           "By October 2027 the market will be pricing 2028. Inputs in blue. Net debt is after another year of heavy building; share counts assume some new shares are sold.")
head(ws, 4, ["Case", "What happens", "2028 revenue, US$bn", "Margin", "Multiple", "Net debt end 2027, US$bn", "Shares, m",
             "Weight", "2028 EBITDA, US$bn", "Value per share, US$", "Change vs price"])
for i, (n, what, rev, m, x, nd, sh, wt) in enumerate(D.SCEN):
    rr = 5 + i
    put(ws, rr, 1, n, b=True)
    put(ws, rr, 2, what)
    put(ws, rr, 3, rev, F_M1, inp=True)
    put(ws, rr, 4, m, F_PCT0, inp=True)
    put(ws, rr, 5, x, F_X1, inp=True)
    put(ws, rr, 6, nd, F_M1, inp=True)
    put(ws, rr, 7, sh, F_INT, inp=True)
    put(ws, rr, 8, wt, F_PCT0, inp=True)
    put(ws, rr, 9, f"=C{rr}*D{rr}", F_M2)
    put(ws, rr, 10, f"=(I{rr}*E{rr}-F{rr})/G{rr}*1000", F_USD2, b=True)
    put(ws, rr, 11, f"=J{rr}/{REF['price']}-1", F_PCT0, b=True)
put(ws, 8, 1, "Weighted", b=True); put(ws, 8, 2, "Probability-weighted value")
put(ws, 8, 8, "=SUM(H5:H7)", F_PCT0)
put(ws, 8, 10, "=SUMPRODUCT(H5:H7,J5:J7)", F_USD2, b=True)
put(ws, 8, 11, f"=J8/{REF['price']}-1", F_PCT0, b=True)
total_row(ws, 8, 1, 11)
put(ws, 10, 1, "Check", b=True)
put(ws, 10, 2, "Weights sum to 100%"); put(ws, 10, 3, '=IF(ABS(H8-1)<0.0001,"OK","CHECK")')
put(ws, 12, 1, "RE-LOOK", b=True)
put(ws, 13, 2, "Re-look price, US$"); put(ws, 13, 3, f"={REF['relook']}", F_USD2)
put(ws, 14, 2, "Base case upside from the re-look price"); put(ws, 14, 3, "=J6/C13-1", F_PCT0)
put(ws, 15, 2, "Bear case downside from the re-look price"); put(ws, 15, 3, "=J5/C13-1", F_PCT0)
put(ws, 16, 2, "Price at which the base case offers 25% upside"); put(ws, 16, 3, "=J6/1.25", F_USD2)
put(ws, 18, 1, "Pitch: bear US$78 (down 66%), base US$229 (down 1%), bull US$430 (up 85%), weighted ~US$242 (up 4%). Base multiple of 11x sits between CoreWeave and where Nebius trades today.").font = muted

# ======================= SENSITIVITY =======================
ws = sheet("Sensitivity", "SENSITIVITY: VALUE PER SHARE, US$, BY 2028 EBITDA AND MULTIPLE", [30, 14, 14, 14, 30])
put(ws, 3, 1, "Net debt end 2027, US$bn"); put(ws, 3, 2, D.SENS_ND, F_M1, inp=True)
put(ws, 4, 1, "Shares, m"); put(ws, 4, 2, D.SENS_SH, F_INT, inp=True)
head(ws, 6, ["2028 EBITDA, US$bn  /  multiple"] + [None] * 3)
for j, x in enumerate(D.SENS_MULT):
    c = put(ws, 6, 2 + j, x, F_X1, inp=True)
    c.font = Font(color=BLUE, bold=True)
for i, e in enumerate(D.SENS_EBITDA):
    rr = 7 + i
    put(ws, rr, 1, e, '"US$"0.0"bn"', inp=True)
    for j in range(3):
        col = get_column_letter(2 + j)
        put(ws, rr, 2 + j, f"=($A{rr}*{col}$6-$B$3)/$B$4*1000", F_USD0)
put(ws, 11, 1, "Pitch grid: 135 / 174 / 213; 182 / 231 / 280; 228 / 287 / 347. The range is wide and centred close to today's price.").font = muted
put(ws, 12, 1, "Reference price, US$"); put(ws, 12, 2, f"={REF['price']}", F_USD2)
put(ws, 13, 1, "2028 EBITDA the price implies at 11x, US$bn"); put(ws, 13, 2, f"=({REF['price']}*B4/1000+B3)/11", F_M2)

# ======================= SOURCES =======================
ws = sheet("Sources", "SOURCES", [120])
srcs = [
    "Nebius quarterly shareholder letters and results on Form 6-K, Q2 2025 to Q2 2026: revenue, annualised run-rate, margins, capital spending, power and customer concentration.",
    "Nebius Q2 2026 shareholder letter, 12 August 2026: guidance, contract value per megawatt ('ACV per MW is stepping up'), and definitions of connected and active power.",
    "Nebius Q2 2026 earnings call, 12 August 2026: the deployment sequence after power is connected.",
    "Nebius interim financial statements to 30 June 2026: shares (Class A and B), cash, restricted cash, deferred revenue, RPO, notes, options, RSUs and the Nvidia pre-funded warrants.",
    "Nebius releases of 19 and 24 August 2026 on the convertible notes (new 2030 and 2034 notes; exchange of US$800m of 2029/2031 notes for ~15.8m shares), and of 10 July 2026 on the secured loan.",
    "Nebius Form 6-K on the Microsoft agreement, 8 September 2025, and announcement of the second Meta agreement, 16 March 2026.",
    "Bloomberg News on Meta's cloud plans, 1 July 2026.",
    "Consensus revenue, peer multiples, short interest (15 Sep 2026) and share prices: Yahoo Finance, retrieved 6 October 2026.",
    "On-demand GPU price changes as reported by Spheron and Coin Republic, September 2026 (not confirmed on a Nebius page).",
    "Unit economics, pro forma cash, net debt, share count, 2028 scenarios and the sensitivity grid are my own estimates.",
    "Published pitch: https://thephysicallayer.fyi/research/   Journal piece, 30 Sep 2026: https://thephysicallayer.fyi/journal/nebius-paid-for-what-is-switched-on/",
    "Personal research, not investment advice.",
]
for i, s in enumerate(srcs, 3):
    c = ws.cell(row=i, column=1, value=s)
    c.alignment = Alignment(wrap_text=True, vertical="top")

wb.calculation = CalcProperties(fullCalcOnLoad=True)
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("Saved", OUT)

# ---- recalculate with LibreOffice so the saved file carries cached values (formulas are kept) ----
import shutil, subprocess, tempfile
soffice = shutil.which("soffice") or "/opt/homebrew/bin/soffice"
if os.path.exists(soffice):
    tmp = tempfile.mkdtemp(prefix="nbis_recalc_")
    subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, OUT],
                   check=True, capture_output=True)
    rec = os.path.join(tmp, os.path.basename(OUT))
    r = subprocess.run(["python3", os.path.join(HERE, "verify_model.py"), rec], capture_output=True, text=True)
    print(r.stdout[-900:])
    if r.returncode == 0:
        shutil.copyfile(rec, OUT)
        print("Recalculated copy verified and saved over", OUT)
    else:
        print("VERIFY FAILED; left the uncalculated file in place")
