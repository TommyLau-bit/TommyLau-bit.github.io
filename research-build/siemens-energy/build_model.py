"""Siemens Energy AG (Xetra: ENR) initiation model: sum of the parts. Live formulas throughout.

House style follows the Vertiv and Hitachi models: navy header rows, navy bold titles, blue font for hardcoded inputs,
yellow fill on my own assumptions, green font for cross-sheet links, black for formulas.
"""
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from data import *

NAVY = "1F3864"; MUTED = "595959"
F_TTL = Font(bold=True, color=NAVY, size=12)
F_SUB = Font(color=MUTED, size=9)
F_HDR = Font(bold=True, color="FFFFFF", size=10)
F_IN = Font(color="0000FF")
F_INB = Font(color="0000FF", bold=True)
F_LINK = Font(color="008000")
F_B = Font(bold=True)
F_NOTE = Font(color=MUTED, size=9, italic=True)
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_ASSUME = PatternFill("solid", fgColor="FFFF00")
FILL_KEY = PatternFill("solid", fgColor="D9E2F3")
WRAP = Alignment(wrap_text=True, vertical="top")
RIGHT = Alignment(horizontal="right")

EUR0 = '"€"#,##0'
EUR2 = '"€"#,##0.00'
N0 = '#,##0'
BN2 = '#,##0.00'
PCT1 = '0.0%'
PCT0 = '0%'
MULTF = '0.0"x"'
MULT0 = '0"x"'
DLONG = datetime.date.fromisoformat(DATE).strftime("%-d %B %Y")
SUB = f"Tommy Lau | Siemens Energy AG (Xetra: ENR) | {DLONG} | Personal research, not investment advice."

wb = openpyxl.Workbook()


def sheet(name, title, widths):
    ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet(name)
    ws.title = name
    ws["A1"] = title; ws["A1"].font = F_TTL
    ws["A2"] = SUB; ws["A2"].font = F_SUB
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def head(ws, row, cells, start=1):
    for i, c in enumerate(cells, start):
        x = ws.cell(row=row, column=i, value=c)
        x.font = F_HDR; x.fill = FILL_HDR
        x.alignment = Alignment(horizontal="left" if i == start else "right", vertical="center", wrap_text=True)


def note(ws, row, texts):
    for i, t in enumerate(texts):
        ws.cell(row=row + i, column=1, value=t).font = F_NOTE


# =====================================================================  COVER (filled last)
ws = sheet("Cover", "SIEMENS ENERGY AG (XETRA: ENR)  INITIATION MODEL", [40, 24, 90])

# =====================================================================  ASSUMPTIONS
ws = sheet("Assumptions", "ASSUMPTIONS AND INPUTS", [58, 14, 10, 100])
ws["A3"] = ("Blue = hardcoded input. Yellow fill = my own assumption (change these to move the value). "
            "Green = link to another tab. Black = formula. Euro million unless stated. Fiscal year ends 30 September.")
ws["A3"].font = F_NOTE
head(ws, 5, ["Input", "Value", "Unit", "Source / note"])
A = {}
MINE = set()
Q3 = "Earnings Release Q3 FY2026, 5 August 2026"
rows = [
    ("MARKET", None, None, None, None),
    ("price", "Reference share price", PRICE, "€", "Xetra close, 27 October 2026. Yahoo Finance ENR.DE, retrieved 28 October 2026."),
    ("hiclose", "Record close, 24 April 2026", HI_CLOSE, "€", "Yahoo Finance ENR.DE; intraday high €191.66 the same day."),
    ("cons27", "FY2027 consensus EPS (year to September 2027)", CONS_EPS27, "€", "stockanalysis.com (S&P Global data), retrieved 28 October 2026. Cross-check only."),
    ("constp", "Average analyst target", CONS_TP, "€", "stockanalysis.com, 28 analysts; range €100 to €260. Cross-check only."),
    ("SHARES (MILLIONS)", None, None, None, None),
    ("issued", "Issued shares", ISSUED, "m", "861,104,914 (voting rights notices, April to September 2026)."),
    ("treas0", "Treasury shares, 30 September 2025", TREAS_SEP25, "m", "Half-year Financial Report 2026, Note 4."),
    ("bb1", "Bought, first tranche (March to May 2026)", BB_T1, "m", "Share buyback page: 12,618,469 shares for €2.0bn."),
    ("bb2", "Bought, second tranche (June to August 2026)", BB_T2, "m", "Share buyback page: 6,467,098 shares for €1.0bn."),
    ("bb3", "Bought, third tranche (from 24 September 2026, to date)", BB_T3, "m", "Share buyback page, read 28 October 2026: 4,513,232 shares for €623m."),
    ("emp", "Treasury shares issued to employee plans, H1 FY2026", EMP_H1, "m", "Half-year Financial Report 2026, Note 4. H2 issuance not yet disclosed, so not deducted."),
    ("BALANCE SHEET, 30 JUNE 2026 (€ BILLION)", None, None, None, None),
    ("anc", "Adjusted net cash (cash less debt less pensions)", ADJ_NET_CASH, "€bn", Q3 + ": cash 11.088, debt 2.959, pensions 0.410."),
    ("bbcash", "Buybacks paid after 30 June 2026", BB_CASH_AFTER_JUN, "€bn", "Buyback page: July 0.460, August 0.181, September 0.117, October 0.506."),
    ("eqm", "Investments accounted for using the equity method", EQUITY_METHOD, "€bn", Q3 + ". Taken at book; includes the Siemens Energy India stake."),
    ("ncini", "Non-controlling interests' share of net income, FY2026", NCI_NI, "€bn", "Q3 FY2026 analyst presentation: about €0.3bn."),
    ("ncix", "Multiple on minorities' earnings", NCI_MULT, "x", "MY ASSUMPTION. Minorities valued on earnings, not the €0.473bn book value."),
    ("contract", "Contract liabilities (mostly customer advances)", CONTRACT_LIAB, "€bn", Q3 + ". Shown for the risk, not used in the value."),
]
for k in SEGS:
    nm = SEG_NAMES[k]
    rows += [
        (nm.upper(), None, None, None, None),
        (f"r25_{k}", f"{nm} revenue, FY2025", REV25[k], "€m", "Shareholder Letter Q4 FY2025, 14 November 2025."),
        (f"m25_{k}", f"{nm} margin before SI, FY2025", M25[k], "%", "Shareholder Letter Q4 FY2025."),
        (f"r9_{k}", f"{nm} revenue, 9M FY2026", REV9M[k], "€m", Q3 + "."),
        (f"p9_{k}", f"{nm} profit before SI, 9M FY2026", P9M[k], "€m", Q3 + "."),
        (f"q4r_{k}", f"{nm} revenue, Q4 FY2026", Q4_REV[k], "€m", "MY ASSUMPTION, inside the FY2026 outlook; results due 11 November 2026."),
        (f"q4m_{k}", f"{nm} margin before SI, Q4 FY2026", Q4_M[k], "%", "MY ASSUMPTION."),
        (f"g27_{k}", f"{nm} revenue growth, FY2027", G27[k], "%", "MY ASSUMPTION. FY2028 target: " + TGT28[k][0] + " revenue CAGR."),
        (f"m27_{k}", f"{nm} margin before SI, FY2027", M27[k], "%", "MY ASSUMPTION. FY2028 target margin: " + TGT28[k][1] + "."),
    ]
    MINE |= {f"q4r_{k}", f"q4m_{k}", f"g27_{k}", f"m27_{k}"}
rows += [
    ("GROUP", None, None, None, None),
    ("rev25", "Group revenue, FY2025", GROUP_REV25, "€m", "Shareholder Letter Q4 FY2025: €39.1bn."),
    ("pb25", "Group profit before SI, FY2025", GROUP_PBSI25, "€m", "Shareholder Letter Q4 FY2025."),
    ("rec9", "Reconciliation line, 9M FY2026", RECON9M, "€m", Q3 + "."),
    ("rec26", "Reconciliation line, FY2026", RECON26, "€m", "MY ASSUMPTION, Q3 analyst presentation: about minus €0.3bn."),
    ("elim26", "Revenue eliminations, FY2026", ELIM_REV26, "€m", "MY ASSUMPTION, about the 9M pace."),
    ("rec27", "Reconciliation line (corporate costs), FY2027", RECON27, "€m", "MY ASSUMPTION."),
    ("elim27", "Revenue eliminations, FY2027", ELIM_REV27, "€m", "MY ASSUMPTION."),
    ("si27", "Special items, FY2027", SI27, "€m", "MY ASSUMPTION, for the EPS cross-check only."),
    ("ppa27", "Amortisation of acquired intangibles, FY2027", PPA27, "€m", "MY ASSUMPTION, about the 9M FY2026 pace."),
    ("fin27", "Financial result, FY2027", FIN27, "€m", "MY ASSUMPTION."),
    ("tax27", "Tax rate, FY2027", TAX27, "%", "MY ASSUMPTION. Siemens Energy guides below 20% for FY2026."),
    ("nci27", "Non-controlling interests, FY2027", NCI27, "€m", "MY ASSUMPTION."),
    ("MULTIPLES (EV / FY2027E PROFIT BEFORE SI)", None, None, None, None),
    ("x_GT", "Grid Technologies", MULT["GT"], "x", "MY ASSUMPTION. Grid equipment peers' median trailing EV/EBIT 20.8x (Peers tab); 22.4x without Hitachi."),
    ("x_GS", "Gas Services", MULT["GS"], "x", "MY ASSUMPTION. Mitsubishi Heavy 21.5x, Baker Hughes 20.9x; lower for cycle-high new-unit pricing."),
    ("x_TI", "Transformation of Industry", MULT["TI"], "x", "MY ASSUMPTION."),
    ("x_SG", "Siemens Gamesa", MULT["SG"], "x", "MY ASSUMPTION. Vestas 15.5x."),
    ("sgburn", "Siemens Gamesa cash still to burn, FY2027 to FY2028", SG_BURN, "€bn", "MY ASSUMPTION. 9M FY2026 free cash flow pre tax minus €1.7bn."),
    ("x_corp", "Corporate costs (capitalised)", MULT["Corp"], "x", "MY ASSUMPTION."),
]
MINE |= {"ncix", "rec26", "elim26", "rec27", "elim27", "si27", "ppa27", "fin27", "tax27", "nci27", "x_GT", "x_GS", "x_TI", "x_SG", "sgburn", "x_corp"}
r = 6
for key, label, val, unit, src in rows:
    if label is None:
        ws.cell(row=r, column=1, value=key).font = Font(bold=True, color=NAVY)
        r += 1; continue
    ws.cell(row=r, column=1, value=label)
    c = ws.cell(row=r, column=2, value=val)
    c.font = F_INB if key in MINE else F_IN
    if key in MINE: c.fill = FILL_ASSUME
    c.number_format = PCT1 if unit == "%" else (MULT0 if unit == "x" else ("#,##0.000" if unit == "m" else (EUR2 if unit == "€" else (BN2 if unit == "€bn" else N0))))
    ws.cell(row=r, column=3, value=unit)
    ws.cell(row=r, column=4, value=src).alignment = WRAP
    A[key] = f"Assumptions!$B${r}"
    r += 1
L = lambda k: "=" + A[k]

# =====================================================================  SEGMENTS
ws = sheet("Segments", "SEGMENTS: REVENUE AND PROFIT BEFORE SPECIAL ITEMS, FY2025A TO FY2027E (€ MILLION)", [34] + [12] * 11)
head(ws, 4, ["Segment", "FY25 revenue", "FY25 margin", "9M revenue", "9M profit", "Q4E revenue*", "Q4E profit*", "FY26E revenue*",
             "FY26E profit*", "FY26E margin*", "FY27E revenue*", "FY27E profit*"])
for i, k in enumerate(SEGS):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=SEG_NAMES[k])
    vals = [L(f"r25_{k}"), L(f"m25_{k}"), L(f"r9_{k}"), L(f"p9_{k}"), L(f"q4r_{k}"), f"=F{rr}*{A[f'q4m_{k}']}",
            f"=D{rr}+F{rr}", f"=E{rr}+G{rr}", f"=I{rr}/H{rr}", f"=H{rr}*(1+{A[f'g27_{k}']})", f"=K{rr}*{A[f'm27_{k}']}"]
    for j, v in enumerate(vals):
        c = ws.cell(row=rr, column=2 + j, value=v)
        c.number_format = PCT1 if j in (1, 8) else N0; c.alignment = RIGHT
        if v.startswith("=Assumptions"): c.font = F_LINK
ws["A9"] = "Reconciliation and eliminations"
ws["E9"] = L("rec9"); ws["I9"] = L("rec26"); ws["H9"] = L("elim26"); ws["K9"] = L("elim27"); ws["L9"] = L("rec27")
for c in ("E9", "I9", "H9", "K9", "L9"): ws[c].font = F_LINK; ws[c].number_format = N0
ws["A10"] = "Siemens Energy"; ws["A10"].font = F_B
for col in "HIKL":
    ws[f"{col}10"] = f"=SUM({col}5:{col}9)"; ws[f"{col}10"].number_format = N0; ws[f"{col}10"].font = F_B
ws["J10"] = "=I10/H10"; ws["J10"].number_format = PCT1
ws["A11"] = "Group margin before SI, FY2027E"; ws["L11"] = "=L10/K10"; ws["L11"].number_format = PCT1
ws["A13"] = "EPS cross-check, FY2027E (€m)"; ws["A13"].font = Font(bold=True, color=NAVY)
eps = [("Profit before SI", "=L10"), ("Special items", L("si27")), ("Amortisation of acquired intangibles", L("ppa27")),
       ("Financial result", L("fin27")), ("Income before tax", "=SUM(B14:B17)"), ("Tax", f"=-B18*{A['tax27']}"),
       ("Non-controlling interests", f"=-{A['nci27']}"), ("Net income to shareholders", "=SUM(B18:B20)"),
       ("Shares outstanding, m", "=Valuation!D20"), ("EPS FY2027E, €", "=B21/B22"), ("FY2027 consensus EPS, €", L("cons27")),
       ("Mine against consensus", "=B23/B24-1")]
for i, (lab, f) in enumerate(eps):
    rr = 14 + i
    ws.cell(row=rr, column=1, value=lab)
    c = ws.cell(row=rr, column=2, value=f)
    c.number_format = EUR2 if "€" in lab and "€m" not in lab else (PCT1 if lab.startswith("Mine") else ("#,##0.0" if "m" == lab[-1] else N0))
ws["A23"].font = F_B; ws["B23"].font = F_B
note(ws, 27, [
    "*Q4 FY2026, FY2026E and FY2027E are my own estimates, not Siemens Energy's. FY2026 results are due on 11 November 2026.",
    "FY2026 outlook (comparable revenue growth, margin before SI): Gas Services 16 to 18%, 14 to 16%; Grid 25 to 27%, 18 to 20%; TI 5 to 7%, 11 to 13%;",
    "Siemens Gamesa 3 to 5%, break-even; group 14 to 16%, 10 to 12% (raised 12 May 2026, confirmed 5 August and 30 September 2026).",
])

# =====================================================================  GRID: THE CLAIM'S TESTS
ws = sheet("Grid", "GRID TECHNOLOGIES: ORDERS, REVENUE, BACKLOG AND THE CLAIM'S TESTS", [44] + [12] * 7)
head(ws, 4, ["€ million"] + GT_Q_LABELS)
ws["A5"] = "Orders"; ws["A6"] = "Revenue"; ws["A7"] = "Orders / revenue"
for j, (o, rv) in enumerate(zip(GT_Q_ORD, GT_Q_REV)):
    col = get_column_letter(2 + j)
    c = ws[f"{col}5"]; c.value = o; c.font = F_IN; c.number_format = N0
    c = ws[f"{col}6"]; c.value = rv; c.font = F_IN; c.number_format = N0
    ws[f"{col}7"] = f"={col}5/{col}6"; ws[f"{col}7"].number_format = '0.00"x"'
ws["A9"] = "Orders, 9M FY2026"; ws["B9"] = "=SUM(F5:H5)"; ws["B9"].number_format = N0
ws["A10"] = "Revenue, 9M FY2026"; ws["B10"] = "=SUM(F6:H6)"; ws["B10"].number_format = N0
ws["A11"] = "Orders / revenue, 9M FY2026"; ws["B11"] = "=B9/B10"; ws["B11"].number_format = '0.00"x"'
ws["A12"] = "Revenue, four quarters to June 2026"; ws["B12"] = "=SUM(E6:H6)"; ws["B12"].number_format = N0
ws["A13"] = "My FY2026E revenue"; ws["B13"] = "=Segments!H6"; ws["B13"].font = F_LINK; ws["B13"].number_format = N0
ws["A14"] = "9M orders less my full-year revenue (positive = test passed already)"; ws["B14"] = "=B9-B13"; ws["B14"].number_format = N0
head(ws, 16, ["Full years", "Orders", "Revenue", "Orders / revenue"])
for i, (y, o, rv) in enumerate([("FY2024", GT_ORD["FY24"], GT_REV_HIST[2] * 1000), ("FY2025", GT_ORD["FY25"], REV25["GT"])]):
    rr = 17 + i
    ws.cell(row=rr, column=1, value=y)
    for j, v in enumerate([o, rv]):
        c = ws.cell(row=rr, column=2 + j, value=v); c.font = F_IN; c.number_format = N0
    ws.cell(row=rr, column=4, value=f"=B{rr}/C{rr}").number_format = '0.00"x"'
head(ws, 20, ["Backlog, € billion (rounded as reported)"] + GT_BL_LABELS)
ws["A21"] = "Grid order backlog"
for j, v in enumerate(GT_BL):
    c = ws.cell(row=21, column=2 + j, value=v); c.font = F_IN; c.number_format = "0"
ws["A22"] = "Change on previous reading"
for j in range(1, len(GT_BL)):
    col, pc = get_column_letter(2 + j), get_column_letter(1 + j)
    ws[f"{col}22"] = f"={col}21-{pc}21"; ws[f"{col}22"].number_format = "+0;-0;0"
ws["A23"] = "Growth, September 2024 to June 2026"; ws["B23"] = "=F21/B21-1"; ws["B23"].number_format = PCT1
ws["A24"] = "Backlog / revenue over four quarters to June 2026"; ws["B24"] = "=F21*1000/B12"; ws["B24"].number_format = '0.00"x"'
head(ws, 26, ["Grid history", "FY22", "FY23", "FY24", "FY25"])
ws["A27"] = "Revenue, €bn"; ws["A28"] = "Margin before SI"
for j, (v, m) in enumerate(zip(GT_REV_HIST, GT_M_HIST)):
    c = ws.cell(row=27, column=2 + j, value=v); c.font = F_IN; c.number_format = BN2
    c = ws.cell(row=28, column=2 + j, value=m); c.font = F_IN; c.number_format = PCT1
note(ws, 30, [
    "Sources: Earnings Releases Q1 to Q3 FY2026 (11 February, 12 May, 5 August 2026), with prior-year quarters; Shareholder Letters Q4 FY2025 and Q4 FY2024;",
    "Capital Market Day Grid Technologies presentation, 20 November 2025 (FY22 and FY23 revenue and margins). Backlog is rounded to €1bn as reported.",
    "The claim breaks if Grid's orders fall below revenue in a full fiscal year, or its backlog in euros falls for two quarters in a row, before 30 September 2030.",
])

# =====================================================================  VALUATION (SUM OF THE PARTS)
ws = sheet("Valuation", "VALUATION: SUM OF THE PARTS ON FY2027E PROFIT BEFORE SPECIAL ITEMS", [60, 14, 10, 16, 12, 60])
head(ws, 4, ["Part", "FY2027E profit, €bn", "Multiple", "Value, €bn", "Share", "Working / note"])
parts = [("Grid Technologies", "=Segments!L6/1000", "x_GT"), ("Gas Services", "=Segments!L5/1000", "x_GS"),
         ("Transformation of Industry", "=Segments!L7/1000", "x_TI"), ("Siemens Gamesa", "=Segments!L8/1000", "x_SG"),
         ("Corporate costs", "=Segments!L9/1000", "x_corp")]
for i, (lab, f, mk) in enumerate(parts):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=lab)
    c = ws.cell(row=rr, column=2, value=f); c.font = F_LINK; c.number_format = BN2
    c = ws.cell(row=rr, column=3, value="=" + A[mk]); c.font = F_LINK; c.number_format = MULT0
    ws.cell(row=rr, column=4, value=f"=B{rr}*C{rr}" + (f"-{A['sgburn']}" if mk == "x_SG" else "")).number_format = BN2
    ws.cell(row=rr, column=5, value=f"=D{rr}/$D$19").number_format = PCT0
ws["F8"] = "Profit at the multiple, less the cash still to burn."
vrows = [
    (12, "Enterprise value of the parts", "=SUM(D5:D9)", BN2, ""),
    (13, "Adjusted net cash, 30 June 2026", L("anc"), BN2, "Cash less debt less pensions, as Siemens Energy reports it."),
    (14, "Less buybacks paid since 30 June", "=-" + A["bbcash"], BN2, "The share count below is after these buybacks."),
    (15, "Add equity-method investments at book", L("eqm"), BN2, ""),
    (16, "Less non-controlling interests at a multiple of their earnings", f"=-{A['ncini']}*{A['ncix']}", BN2, ""),
    (18, "Bridge from enterprise value to equity value", "=SUM(D13:D16)", BN2, ""),
    (19, "Equity value, €bn", "=D12+D18", BN2, ""),
    (20, "Shares outstanding, m", f"={A['issued']}-({A['treas0']}+{A['bb1']}+{A['bb2']}+{A['bb3']}-{A['emp']})", "#,##0.0",
     "Issued less treasury shares, derived from the buyback reports. Cancelling treasury shares does not change this."),
    (21, "BASE VALUE PER SHARE, €", "=D19/D20*1000", EUR2, ""),
    (22, "Reference price, €", L("price"), EUR2, "Xetra close, 27 October 2026."),
    (23, "Base value against price", "=D21/D22-1", PCT1, "Band: LONG at +15% or more, SHORT at -25% or worse, NO CALL between."),
    (24, "VIEW ON THE BAND", '=IF(D23>=0.15,"LONG",IF(D23<=-0.25,"SHORT","NO CALL"))', None, "Tommy Lau's call, 28 October 2026: LONG, target €160, Medium conviction."),
    (26, "Market value, €bn", "=D20*D22/1000", BN2, ""),
    (27, "Adjusted enterprise value at the price, €bn", "=D26-D18", BN2, "Market value less the bridge items."),
    (28, "EV / FY2026E group profit before SI", "=D27/(Segments!I10/1000)", MULTF, ""),
    (29, "EV / FY2027E group profit before SI", "=D27/(Segments!L10/1000)", MULTF, ""),
    (30, "P/E on FY2027 consensus EPS", "=D22/" + A["cons27"], MULTF, "Cross-check only."),
    (31, "P/E on my FY2027E EPS", "=D22/Segments!B23", MULTF, ""),
    (32, "Value the price leaves for Grid, €bn", "=D27-SUM(D6:D9)", BN2, "Rest of the group at my base multiples."),
    (33, "That value / Grid FY2027E profit before SI", "=D32/B5", MULTF, ""),
    (34, "Gas Services multiple the price implies, with Grid at base", "=(D27-D5-SUM(D7:D9))/B6", MULTF, ""),
    (35, "Price at which the base value is 15% above it, €", "=D21/1.15", EUR2, "Above this, the band gives NO CALL."),
    (36, "Price at which the base value is 25% below it, €", "=D21/0.75", EUR2, "Above this, the band gives SHORT."),
    (37, "Grid share of my equity value", "=D5/D19", PCT1, ""),
    (38, "Gas Services share of my equity value", "=D6/D19", PCT1, ""),
    (39, "Price against the record close", "=D22/" + A["hiclose"] + "-1", PCT1, "€187.62 on 24 April 2026."),
]
for rr, lab, f, fmt, nt in vrows:
    ws.cell(row=rr, column=1, value=lab)
    c = ws.cell(row=rr, column=4, value=f)
    if fmt: c.number_format = fmt
    if rr in (21, 24):
        ws.cell(row=rr, column=1).font = F_B; c.font = F_B; c.fill = FILL_KEY
    if isinstance(f, str) and f.startswith("=Assumptions"): c.font = F_LINK
    ws.cell(row=rr, column=6, value=nt).alignment = WRAP
ws["E18"] = "=D18/$D$19"; ws["E18"].number_format = PCT0
ws["A12"].font = F_B; ws["A19"].font = F_B

# =====================================================================  SCENARIOS
ws = sheet("Scenarios", "THREE CASES, TWELVE MONTHS OUT", [12, 11, 11, 10, 11, 11, 10, 9, 11, 10, 11, 13, 12, 11])
head(ws, 4, ["Case", "Grid growth", "Grid margin", "Grid x", "Gas growth", "Gas margin", "Gas x", "TI x", "Gamesa margin",
             "Gamesa x", "Gamesa burn €bn", "Grid profit €bn", "Value, €", "vs price"])
r = 5
for name, gt, gs, tix, sg, p in SCEN:
    ws.cell(row=r, column=1, value=name).font = F_B
    if name == "Base":
        links = ["g27_GT", "m27_GT", "x_GT", "g27_GS", "m27_GS", "x_GS", "x_TI", "m27_SG", "x_SG", "sgburn"]
        for j, k in enumerate(links):
            c = ws.cell(row=r, column=2 + j, value="=" + A[k]); c.font = F_LINK
    else:
        for j, v in enumerate([gt[0], gt[1], gt[2], gs[0], gs[1], gs[2], tix, sg[0], sg[1], sg[2]]):
            c = ws.cell(row=r, column=2 + j, value=v); c.font = F_INB; c.fill = FILL_ASSUME
    for col in (2, 3, 5, 6, 9): ws.cell(row=r, column=col).number_format = PCT1
    for col in (4, 7, 8, 10): ws.cell(row=r, column=col).number_format = MULT0
    ws.cell(row=r, column=11).number_format = BN2
    ws.cell(row=r, column=12, value=f"=Segments!$H$6*(1+B{r})*C{r}/1000").number_format = BN2
    ev = (f"L{r}*D{r}+Segments!$H$5*(1+E{r})*F{r}/1000*G{r}+Valuation!$B$7*H{r}"
          f"+Segments!$K$8*I{r}/1000*J{r}-K{r}+Valuation!$D$9")
    ws.cell(row=r, column=13, value=f"=({ev}+Valuation!$D$18)/Valuation!$D$20*1000").number_format = EUR2
    ws.cell(row=r, column=14, value=f"=M{r}/{A['price']}-1").number_format = PCT0
    r += 1
ws["A9"] = "Probability"; ws["A9"].font = F_B
for j, s in enumerate(SCEN):
    ws.cell(row=8, column=2 + j, value=s[0]).font = F_B
    c = ws.cell(row=9, column=2 + j, value=s[5]); c.font = F_INB; c.fill = FILL_ASSUME; c.number_format = PCT0
ws["E9"] = "=SUM(B9:D9)"; ws["E9"].number_format = PCT0; ws["F9"] = "must be 100%"; ws["F9"].font = F_NOTE
ws["A11"] = "Probability-weighted value, €"; ws["A11"].font = F_B
ws["E11"] = "=B9*M5+C9*M6+D9*M7"; ws["E11"].number_format = EUR2; ws["E11"].font = F_B; ws["E11"].fill = FILL_KEY
ws["F11"] = f"=E11/{A['price']}-1"; ws["F11"].number_format = PCT1
note(ws, 13, [
    "Grid and Gas Services grow from my FY2026 estimates. Transformation of Industry keeps my FY2027E profit; only its multiple moves.",
    "Siemens Gamesa: FY2027E revenue at the case margin and multiple, less the cash still to burn. Corporate costs, the bridge and shares are the same in every case.",
])

# =====================================================================  SENSITIVITY
ws = sheet("Sensitivity", "VALUE PER SHARE: GAS SERVICES MULTIPLE AGAINST THE GRID MULTIPLE (€)", [30, 14, 14, 14])
head(ws, 4, ["Gas x (rows) \\ Grid x (columns)", "", "", ""])
for j, (m, link) in enumerate(zip(SENS_GT_X, [False, True, False])):
    c = ws.cell(row=4, column=2 + j, value=("=" + A["x_GT"]) if link else m)
    c.number_format = MULT0; c.font = Font(bold=True, color="FFFFFF"); c.fill = FILL_HDR; c.alignment = RIGHT
for i, (g, link) in enumerate(zip(SENS_GS_X, [False, True, False])):
    rr = 5 + i
    c = ws.cell(row=rr, column=1, value=("=" + A["x_GS"]) if link else g)
    c.number_format = MULT0; c.font = F_LINK if link else F_IN
    for j in range(3):
        col = get_column_letter(2 + j)
        ws.cell(row=rr, column=2 + j,
                value=f"=(Valuation!$B$5*{col}$4+Valuation!$B$6*$A{rr}+SUM(Valuation!$D$7:$D$9)+Valuation!$D$18)/Valuation!$D$20*1000").number_format = EUR0
ws["C6"].fill = FILL_KEY
ws["A9"] = "Cells above the reference price"; ws["A9"].font = F_B
ws["B9"] = f'=COUNTIF(B5:D7,">"&{A["price"]})'; ws["C9"] = "of 9"
ws["A10"] = "Cells 15% or more above the price"; ws["A10"].font = F_B
ws["B10"] = f'=COUNTIF(B5:D7,">="&{A["price"]}*1.15)'; ws["C10"] = "of 9"
note(ws, 12, ["My FY2027E profits throughout; everything else at base. The middle row and column link to the base multiples."])

# =====================================================================  PEERS
ws = sheet("Peers", "PEER MULTIPLES (stockanalysis.com, RETRIEVED 28 OCTOBER 2026)", [22, 14, 13, 18, 70])
head(ws, 4, ["Company", "Ticker", "Forward P/E", "EV/EBIT (trailing)", "What they make"])
for i, (co, tk, pe, eve, what, grp) in enumerate(PEERS):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=co).font = F_B if grp == "self" else Font()
    ws.cell(row=rr, column=2, value=tk)
    for j, v in enumerate([pe, eve]):
        c = ws.cell(row=rr, column=3 + j, value=v); c.number_format = MULTF; c.font = F_IN
    ws.cell(row=rr, column=5, value=what).alignment = WRAP
ws["A16"] = "Grid equipment median (rows 6 to 10)"; ws["A16"].font = F_B
ws["C16"] = "=MEDIAN(C6:C10)"; ws["D16"] = "=MEDIAN(D6:D10)"
ws["A17"] = "Grid equipment median without Hitachi (rows 6 to 9)"; ws["A17"].font = F_B
ws["C17"] = "=MEDIAN(C6:C9)"; ws["D17"] = "=MEDIAN(D6:D9)"
for c in ("C16", "D16", "C17", "D17"): ws[c].number_format = MULTF
note(ws, 19, ["Multiples only, to set the range of segment multiples. No view on the peers' shares is expressed. GE Vernova's trailing",
              "EV/EBIT of 84x reflects earnings still recovering and is not used for any multiple."])

# =====================================================================  SOURCES
ws = sheet("Sources", "SOURCES", [6, 140])
head(ws, 4, ["No.", "Source and what it supports"])
srcs = [
    "Siemens Energy, Earnings Release Q3 FY2026, 5 August 2026: segment orders, revenue and profit before SI (Q3 and 9M), balance sheet, adjusted net cash, FY2026 outlook, external revenue split.",
    "Siemens Energy, Earnings Releases Q1 FY2026 (11 February 2026) and Q2 FY2026 (12 May 2026): quarterly Grid figures with prior-year quarters; the raised FY2026 outlook.",
    "Siemens Energy, Shareholder Letter Q4 FY2025, 14 November 2025: FY2025 segment results, FY2028 group targets, the up to €6bn buyback and up to €10bn shareholder returns to FY2028.",
    "Siemens Energy, Shareholder Letter Q4 FY2024, 13 November 2024: FY2024 Grid orders €20,901m, revenue €9,280m, backlog €33bn.",
    "Siemens Energy, Capital Market Day, 20 November 2025: CEO, CFO, Grid Technologies, Gas Services, Siemens Gamesa and Transformation of Industry presentations (FY2028 targets; Grid FY22 to FY25; the transformer market chart).",
    "Siemens Energy, Half-year Financial Report 2026, 12 May 2026, Note 4: treasury shares.",
    "Siemens Energy, Q3 FY2026 analyst presentation, 5 August 2026: FY2026 financial considerations (reconciliation line, minorities, tax rate).",
    "Siemens Energy, Q4 FY2026 Pre-Close Call transcript, 30 September 2026; investor relations financial calendar: Q4 FY2026 results on 11 November 2026.",
    "Siemens Energy share buyback page (siemens-energy.com/sharebuyback), read 28 October 2026: shares and amounts by tranche and month.",
    "Siemens AG voting rights notice, 2 April 2026 (via Deutsche Boerse): 861,104,914 voting rights; Siemens AG at 5.54%.",
    "Yahoo Finance, ENR.DE, retrieved 28 October 2026: daily and month-end closes, highs and lows.",
    "stockanalysis.com (S&P Global data), retrieved 28 October 2026: consensus EPS and average target; peer forward P/E and trailing EV/EBIT.",
    "The Physical Layer, 'Siemens Energy has measured the transformer shortage, and its own chart runs to 2030', 28 October 2026: https://thephysicallayer.fyi/journal/siemens-energy-measures-the-shortage/",
    "Estimates for FY2026 and FY2027 are the author's own. Personal research, not investment advice.",
]
for i, s in enumerate(srcs):
    ws.cell(row=5 + i, column=1, value=i + 1)
    ws.cell(row=5 + i, column=2, value=s).alignment = WRAP

# =====================================================================  COVER
ws = wb["Cover"]
head(ws, 4, ["Item", "Value", "Note"])
cover = [
    ("Call (Tommy Lau, 28 October 2026)", "=Valuation!D24", None, "LONG at +15% or more, SHORT at -25% or worse, NO CALL between."),
    ("Base value per share, €", "=Valuation!D21", EUR2, "Sum of the parts on FY2027E profit before SI, plus the bridge, over shares outstanding."),
    ("Reference price, €", "=" + A["price"], EUR2, "Xetra close, 27 October 2026."),
    ("Base value against price", "=Valuation!D23", PCT1, ""),
    ("Probability-weighted value, €", "=Scenarios!E11", EUR2, "25 / 50 / 25 bear / base / bull."),
    ("Grid share of my equity value", "=Valuation!D37", PCT1, ""),
    ("Target, €", CALL["target"], EUR0, "Probability-weighted value, rounded."),
    ("Conviction", CALL["conviction"], None, "Judgement, not formula."),
    ("Horizon", "12 months, to October 2027", None, ""),
    ("Wrong if", "See note", None, WRONG_IF),
]
for i, (lab, v, fmt, nt) in enumerate(cover):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=lab).font = F_B
    c = ws.cell(row=rr, column=2, value=v)
    if fmt: c.number_format = fmt
    c.alignment = RIGHT
    c.font = F_LINK if isinstance(v, str) and v.startswith("=") else F_IN
    ws.cell(row=rr, column=3, value=nt).alignment = WRAP
ws.row_dimensions[14].height = 44
ws["A17"] = "Tabs"; ws["A17"].font = Font(bold=True, color=NAVY)
for i, (t, d) in enumerate([
    ("Assumptions", "Every input, with its source. Change the yellow cells to move the value."),
    ("Segments", "Revenue and profit before SI by segment: FY2025A, 9M FY2026A, Q4 FY2026E, FY2026E, FY2027E; EPS cross-check."),
    ("Grid", "Grid orders, revenue and backlog by quarter: the claim's tests."),
    ("Valuation", "Sum of the parts, the bridge, base value, the band, what the price implies for Grid."),
    ("Scenarios", "Bear / base / bull and the probability-weighted value."),
    ("Sensitivity", "Gas Services multiple against the Grid multiple."),
    ("Peers", "Forward P/E and trailing EV/EBIT used to set the multiples."),
    ("Sources", "Every source used."),
]):
    ws.cell(row=18 + i, column=1, value=t)
    ws.cell(row=18 + i, column=2, value=d)
ws["A28"] = "Colour key: blue = hardcoded input; yellow fill = my own assumption; green = link to another tab; black = formula."
ws["A28"].font = F_NOTE
ws["A29"] = f"Companion note: {DATE}_{COMPANY}_Initiation.pdf. Personal research, not investment advice."
ws["A29"].font = F_NOTE

for w in wb.worksheets:
    w.freeze_panes = "A5"
wb.active = 0
wb.save(XLSX)
print("model written", XLSX)
